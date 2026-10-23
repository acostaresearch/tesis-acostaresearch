#!/usr/bin/env python3
"""Ajustes de estructura del manuscrito para la guía de una revista (Fase 8).

Complementa a articulo_docx.py (que agrega secciones) con lo que este no hace:
renombrar, mover, quitar, unir secciones o cambiar su nivel, llevar tablas y figuras al final
(si la revista lo pide; el _RESPALDO las conserva en el texto) y limpiar los metadatos del
archivo para la revisión doble ciego. Solo necesita python-docx.

Una "sección" es su título (estilo Heading 1, 2 o 3) más todo lo que sigue
hasta el siguiente título del mismo nivel o superior: párrafos, tablas,
figuras y subtítulos. El título se compara sin tildes ni mayúsculas.

  python3 ajustar_docx.py renombrar  --docx A_v8.docx --de "Métodos" --a "Materiales y métodos"
  python3 ajustar_docx.py mover      --docx A_v8.docx --seccion "Conclusiones" --antes-de "Discusión"
  python3 ajustar_docx.py mover      --docx A_v8.docx --seccion "Abstract" --antes-de "Resumen"
  python3 ajustar_docx.py quitar     --docx A_v8.docx --seccion "Agradecimientos"
  python3 ajustar_docx.py anonimizar --docx A_v8.docx
  python3 ajustar_docx.py al-final   --docx A_v8.docx --tipo tablas|figuras|todo [--marcar]
                                     [--antes-de-referencias]
      (--marcar deja «[Tabla N aquí]» donde estaba; por defecto van tras las
      Referencias, cada una en su página. Correr `verificar` ANTES: el motor
      lee como referencia todo lo que sigue al título Referencias.)
  python3 ajustar_docx.py reemplazar --docx A_v8.docx --cambios parrafos.json
      (párrafos enteros, sin marcas: humanización completa y recortes; se corre
      igual sobre el _RESPALDO. Conserva el marcador de cada cita por clave:
      la cita debe reaparecer tal cual en el texto nuevo. [POR CONFIRMAR] en rojo)
  python3 ajustar_docx.py nivel      --docx A_v8.docx --seccion "Revisión de la literatura" --nivel 2
      (la baja a subsección con sus subtítulos; si no queda dentro de la
      Introducción, después `mover --antes-de "Método"`)
  python3 ajustar_docx.py unir       --docx A_v8.docx --primera "Resultados" --segunda "Discusión" \
                                     --titulo "Resultados y discusión" [--sin-subtitulo]
  python3 ajustar_docx.py secciones  --docx A_v8.docx

Escribe sobre el mismo archivo salvo que se pase --out. Se niega a tocar una
versión _v7 o anterior: se trabaja siempre sobre la copia _v8.
"""
import argparse
import copy
import re
import sys
import unicodedata

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    sys.exit("Falta python-docx: pip install python-docx")


def _norm(s):
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def _texto(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _nivel(el):
    if el.tag != qn("w:p"):
        return None
    st = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    if st is None:
        return None
    m = re.search(r"(?:heading|ttulo|t[ií]tulo)\s*(\d)", st.get(qn("w:val")) or "", re.I)
    return int(m.group(1)) if m else None


def _hijos(doc):
    return [e for e in doc.element.body.iterchildren() if e.tag != qn("w:sectPr")]


def _bloque(doc, nombre):
    hijos = _hijos(doc)
    for i, e in enumerate(hijos):
        n = _nivel(e)
        if n and _norm(_texto(e)) == _norm(nombre):
            fin = i + 1
            while fin < len(hijos):
                n2 = _nivel(hijos[fin])
                if n2 and n2 <= n:
                    break
                fin += 1
            return hijos[i:fin]
    disponibles = [_texto(e) for e in hijos if _nivel(e)]
    sys.exit(f"No encontré la sección «{nombre}». Títulos del documento: {disponibles}")


def _guardar(doc, a):
    out = a.out or a.docx
    m = re.search(r"_v(\d+)", out)
    if m and int(m.group(1)) < 8:
        sys.exit("No modifico una versión anterior a la _v8: copia primero a ARTICULO_[clave]_v8.docx.")
    doc.save(out)
    print(f"ok: {out}")


def cmd_secciones(a):
    doc = Document(a.docx)
    for e in _hijos(doc):
        n = _nivel(e)
        if n:
            print("  " * (n - 1) + _texto(e))


def cmd_renombrar(a):
    doc = Document(a.docx)
    titulo = _bloque(doc, a.de)[0]
    ts = list(titulo.iter(qn("w:t")))
    ts[0].text = a.a
    for t in ts[1:]:
        t.text = ""
    _guardar(doc, a)


def _poner_nivel(el, nivel):
    ppr = el.find(qn("w:pPr"))
    st = ppr.find(qn("w:pStyle")) if ppr is not None else None
    if st is not None:
        st.set(qn("w:val"), f"Heading{nivel}")


def cmd_nivel(a):
    """Cambia el nivel de un título y baja o sube con él sus subtítulos."""
    doc = Document(a.docx)
    bloque = _bloque(doc, a.seccion)
    delta = int(a.nivel) - _nivel(bloque[0])
    for e in bloque:
        n = _nivel(e)
        if n:
            _poner_nivel(e, max(1, min(3, n + delta)))
    _guardar(doc, a)


def cmd_unir(a):
    """Une dos secciones contiguas o no en una: «Resultados» + «Discusión» →
    «Resultados y discusión». La segunda pasa a subsección (o pierde su título
    con --sin-subtitulo) y se coloca al final de la primera."""
    doc = Document(a.docx)
    b1 = _bloque(doc, a.primera)
    b2 = _bloque(doc, a.segunda)
    ts = list(b1[0].iter(qn("w:t")))
    ts[0].text = a.titulo
    for t in ts[1:]:
        t.text = ""
    n1 = _nivel(b1[0])
    if a.sin_subtitulo:
        b2[0].getparent().remove(b2[0])
        b2 = b2[1:]
    for e in b2:
        n = _nivel(e)
        if n:
            _poner_nivel(e, min(3, n + 1))
    ancla = b1[-1]
    for e in b2:
        ancla.addnext(e)
        ancla = e
    _guardar(doc, a)


def cmd_mover(a):
    doc = Document(a.docx)
    bloque = _bloque(doc, a.seccion)
    ancla = _bloque(doc, a.antes_de)[0]
    if ancla in bloque:
        sys.exit("La sección ancla está dentro de la que se mueve.")
    for e in bloque:
        ancla.addprevious(e)
    _guardar(doc, a)


def cmd_quitar(a):
    doc = Document(a.docx)
    bloque = _bloque(doc, a.seccion)
    for e in bloque:
        e.getparent().remove(e)
    print(f"quitada: {a.seccion} ({len(bloque)} elementos)")
    _guardar(doc, a)


RE_CAP = re.compile(r"^(Tabla|Table|Figura|Figure)\s+(\d+)$")
RE_NOTA = re.compile(r"^(Nota|Note)\.")


def cmd_al_final(a):
    """Lleva tablas y/o figuras (rótulo, título, objeto y nota) al final del
    documento, cada una en página propia, y deja la llamada del texto intacta."""
    doc = Document(a.docx)
    quiere = {"tablas": ("tabla", "table"), "figuras": ("figura", "figure"),
              "todo": ("tabla", "table", "figura", "figure")}[a.tipo]
    hijos = _hijos(doc)
    bloques, i = [], 0
    while i < len(hijos):
        m = RE_CAP.match(_texto(hijos[i]).strip()) if hijos[i].tag == qn("w:p") else None
        if m and m.group(1).lower() in quiere and i + 2 < len(hijos):
            fin = i + 3                       # rótulo, título, tabla o párrafo con la imagen
            if fin < len(hijos) and RE_NOTA.match(_texto(hijos[fin]).strip()):
                fin += 1
            bloques.append((f"{m.group(1)} {m.group(2)}", hijos[i:fin]))
            i = fin
        else:
            i += 1
    if not bloques:
        sys.exit("No encontré tablas ni figuras con rótulo «Tabla N» / «Figura N».")
    sect = doc.element.body.find(qn("w:sectPr"))
    if a.antes_de_referencias:
        refs = [e for e in hijos if _nivel(e) == 1 and _norm(_texto(e)) in
                ("referencias", "references", "referencias bibliograficas", "bibliografia")]
        if not refs:
            sys.exit("No encontré el título de Referencias.")
        sect = refs[0]

    def salto():
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        r.append(br)
        p.append(r)
        sect.addprevious(p)

    # Título propio para que el motor no lea las tablas como referencias
    modelo = next((e for e in hijos if _nivel(e) == 1), None)
    if modelo is not None:
        en = bloques[0][0].startswith(("Table", "Figure"))
        rotulo = {"tablas": ("Tablas", "Tables"), "figuras": ("Figuras", "Figures"),
                  "todo": ("Tablas y figuras", "Tables and figures")}[a.tipo][en]
        tit = copy.deepcopy(modelo)
        ts = list(tit.iter(qn("w:t")))
        ts[0].text = rotulo
        for t in ts[1:]:
            t.text = ""
        salto()
        sect.addprevious(tit)
    for k, (nombre, bloque) in enumerate(bloques):
        primero = bloque[0]
        if a.marcar:                          # «[Tabla 2 aquí]» donde estaba
            marca = copy.deepcopy(primero)
            ts = list(marca.iter(qn("w:t")))
            ts[0].text = f"[{nombre} here]" if nombre.startswith(("Table", "Figure")) else f"[{nombre} aquí]"
            for t in ts[1:]:
                t.text = ""
            primero.addprevious(marca)
        if k:
            salto()
        for e in bloque:
            sect.addprevious(e)
    print(f"al final: {', '.join(n for n, _ in bloques)}")
    _guardar(doc, a)


RE_MD = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
RE_PEND = re.compile(r"(\[POR CONFIRMAR[^\]]*\])")
PREFIJO_CITA = "_AcCita"   # marcadores que pone articulo_docx.py alrededor de cada cita por clave


def _citas_del_parrafo(p):
    """[(bookmarkStart, bookmarkEnd, texto visible)] de las citas por clave del párrafo."""
    out, hijos = [], list(p)
    for i, e in enumerate(hijos):
        if e.tag == qn("w:bookmarkStart") and (e.get(qn("w:name")) or "").startswith(PREFIJO_CITA):
            bid, txt = e.get(qn("w:id")), []
            for f in hijos[i + 1:]:
                if f.tag == qn("w:bookmarkEnd") and f.get(qn("w:id")) == bid:
                    out.append((e, f, "".join(txt)))
                    break
                txt.append(_texto(f))
    return out


def _run(texto, base, negrita=False, cursiva=False, rojo=False):
    r = OxmlElement("w:r")
    rpr = copy.deepcopy(base) if base is not None else OxmlElement("w:rPr")
    for tag in ("w:b", "w:i", "w:color"):
        for e in rpr.findall(qn(tag)):
            rpr.remove(e)
    if negrita or cursiva or rojo:
        pos = 1 if rpr.find(qn("w:rFonts")) is not None else 0
        if negrita or rojo:
            rpr.insert(pos, OxmlElement("w:b"))
        if cursiva:
            rpr.insert(pos, OxmlElement("w:i"))
        if rojo:
            col = OxmlElement("w:color")
            col.set(qn("w:val"), "FF0000")
            rpr.append(col)
    r.append(rpr)
    tt = OxmlElement("w:t")
    tt.set(qn("xml:space"), "preserve")
    tt.text = texto
    r.append(tt)
    return r


def _runs_md(texto, base):
    """*cursiva*, **negrita** y [POR CONFIRMAR: …] en rojo y negrita, como el motor."""
    out = []
    for trozo in RE_MD.split(texto):
        if not trozo:
            continue
        negrita, cursiva = trozo.startswith("**"), trozo.startswith("*") and not trozo.startswith("**")
        limpio = trozo.strip("*") if (negrita or cursiva) else trozo
        for sub in RE_PEND.split(limpio):
            if sub:
                out.append(_run(sub, base, negrita, cursiva, bool(RE_PEND.fullmatch(sub))))
    return out


def cmd_reemplazar(a):
    """Reemplaza párrafos enteros (sin control de cambios) desde un JSON
    [{"buscar": "fragmento único del párrafo", "reemplazar": "párrafo nuevo completo"}].
    Para la humanización completa, los recortes y las conversiones a mano.
    Conserva el formato del párrafo, pone en rojo los [POR CONFIRMAR] nuevos y
    **conserva el marcador de cada cita por clave** (el que permite a
    `articulo_docx.py normalizar` reconvertirla en un reenvío): cada cita del
    párrafo tiene que reaparecer, tal cual, en «reemplazar». Si falta alguna,
    el cambio no se aplica, salvo que lleve "quitar_citas": true (recorte que
    elimina la cita a propósito; después, `normalizar --fuentes` rehace la lista)."""
    import json
    doc = Document(a.docx)
    cambios = json.load(open(a.cambios, encoding="utf-8"))
    parrafos = [p for p in doc.element.body.iter(qn("w:p"))]
    hechos, fallos = 0, []
    for k, c in enumerate(cambios, 1):
        b = _norm(c["buscar"])
        hall = [p for p in parrafos if b and b in _norm(_texto(p))]
        if len(hall) != 1:
            fallos.append(f"✗ {k}: «{c['buscar'][:60]}» aparece {len(hall)} veces (debe ser 1)")
            continue
        p = hall[0]
        nuevo = c.get("reemplazar", "")
        citas = _citas_del_parrafo(p)
        ubicadas, perdidas, desde = [], [], 0
        for ini, fin, txt in citas:
            i = nuevo.find(txt, desde) if txt else -1
            if i < 0:
                i = nuevo.find(txt) if txt else -1
                if i >= 0 and any(i < j + len(t) and j < i + len(txt) for j, t, _, _ in ubicadas):
                    i = -1
            if i < 0:
                perdidas.append(txt)
                continue
            ubicadas.append((i, txt, ini, fin))
            desde = i + len(txt)
        if perdidas and not c.get("quitar_citas"):
            fallos.append(f"✗ {k}: en «reemplazar» faltan, tal cual, las citas "
                          + "; ".join(f"«{t}»" for t in perdidas)
                          + ". Cópialas sin cambios (o \"quitar_citas\": true si se quitan a propósito)")
            continue
        runs = [r for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None]
        base = runs[0].find(qn("w:rPr")) if runs else None
        for r in list(p):
            if r.tag != qn("w:pPr"):
                p.remove(r)
        pos = 0
        for i, txt, ini, fin in sorted(ubicadas, key=lambda x: x[0]):
            for r in _runs_md(nuevo[pos:i], base):
                p.append(r)
            p.append(ini)
            p.append(_run(txt, base))
            p.append(fin)
            pos = i + len(txt)
        for r in _runs_md(nuevo[pos:], base):
            p.append(r)
        if perdidas:
            print(f"  · {k}: citas quitadas a propósito: {'; '.join(perdidas)}")
        hechos += 1
    print(f"reemplazados {hechos} de {len(cambios)} párrafos")
    for f in fallos:
        print(f)
    _guardar(doc, a)
    if fallos:
        sys.exit("Corrige esos «buscar» (copiados de `articulo_docx.py texto`) y repite solo esos.")


def cmd_anonimizar(a):
    doc = Document(a.docx)
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""
    cp.title = cp.title or ""
    print("metadatos de autor vaciados (autor, última modificación, comentarios)")
    _guardar(doc, a)
    print("Ojo: si el autor abre y guarda el archivo en Word, Word vuelve a poner su nombre. "
          "En Word: Revisar > Proteger > 'Quitar información personal de este archivo al guardar' "
          "(Mac: Word > Preferencias > Seguridad).")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for nombre, f, extra in [
        ("secciones", cmd_secciones, []),
        ("renombrar", cmd_renombrar, [("--de", True), ("--a", True)]),
        ("mover", cmd_mover, [("--seccion", True), ("--antes-de", True)]),
        ("quitar", cmd_quitar, [("--seccion", True)]),
        ("nivel", cmd_nivel, [("--seccion", True), ("--nivel", True)]),
        ("unir", cmd_unir, [("--primera", True), ("--segunda", True), ("--titulo", True),
                            ("--sin-subtitulo", "flag")]),
        ("anonimizar", cmd_anonimizar, []),
        ("reemplazar", cmd_reemplazar, [("--cambios", True)]),
        ("al-final", cmd_al_final, [("--tipo", False), ("--marcar", "flag"),
                                         ("--antes-de-referencias", "flag")]),
    ]:
        p = sub.add_parser(nombre)
        p.add_argument("--docx", required=True)
        p.add_argument("--out")
        for arg, req in extra:
            if req == "flag":
                p.add_argument(arg, action="store_true")
            elif arg == "--tipo":
                p.add_argument(arg, choices=["tablas", "figuras", "todo"], default="tablas")
            else:
                p.add_argument(arg, required=req)
        p.set_defaults(f=f)
    a = ap.parse_args()
    a.f(a)


if __name__ == "__main__":
    main()
