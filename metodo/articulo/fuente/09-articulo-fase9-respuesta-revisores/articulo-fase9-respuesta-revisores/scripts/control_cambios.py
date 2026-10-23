#!/usr/bin/env python3
"""Control de cambios real de Word (w:ins / w:del) sobre un .docx existente.

Ruta del Artículo Empírico, Fase 9 (Acosta | IA & Research). Autocontenido:
solo necesita python-docx (que ya trae lxml). No usa Node, LibreOffice ni
scripts externos, y funciona igual en claude.ai, Claude Code, la API u otro
cliente con Python.

Subcomandos
-----------
  aplicar      Marca los cambios de un JSON como revisiones de Word a nombre
               del autor. No regenera el documento: edita el XML en su sitio y
               conserva formato, citas, tablas y numeración.
  validar      Comprueba que el resultado abre, cuenta inserciones y
               eliminaciones por autor, que no hay texto cambiado sin marcar
               (con --original) y que, aceptando todo, aparecen los
               "reemplazar" y desaparecen los "buscar" (con --cambios).
  aceptar      Genera la versión limpia (todos los cambios aceptados) para
               revisarla o pasarle `articulo_docx.py verificar`.
  respaldo     Aplica el mismo JSON al _RESPALDO y lo deja aceptado (sin
               marcas): el respaldo neutro queda con el texto revisado.
  comentarios  Lista los comentarios que ya trae el Word (los del revisor),
               con el texto al que están anclados.

Formato del JSON de cambios (lista)
-----------------------------------
  [
    {"buscar": "texto exacto del manuscrito", "reemplazar": "texto nuevo",
     "comentario": "opcional: nota visible en el margen de Word"},
    {"buscar": "frase que sale", "reemplazar": ""},             # eliminación
    {"insertar_despues": "fragmento único del párrafo ancla",
     "texto": "Párrafo nuevo completo, con *cursiva* si hace falta."}
  ]

- "buscar" se localiza dentro de UN párrafo aunque Word lo tenga repartido en
  varios fragmentos con formatos distintos. Espacios dobles y espacios duros
  cuentan como un espacio. Si aparece más de una vez se cambia la primera y se
  avisa (agregar contexto o "todas": true).
- Valor repetido (cifras de tablas): {"buscar": ".81", "reemplazar": ".83",
  "tabla": 3, "fila": 4, "columna": 2} busca solo en esa celda (desde 1; la
  fila de encabezado es la 1; la tabla es la N-ésima del documento), o
  "ocurrencia": n toma la n-ésima aparición, contada sobre el documento tal
  como lo dejaron los cambios anteriores del mismo JSON.
- Solo se marca lo que de verdad cambia (diferencia por palabras): lo que se
  repite entre "buscar" y "reemplazar" queda intacto, con su formato.
- En "reemplazar" y "texto", *cursiva* y **negrita** fuerzan ese formato; lo
  demás hereda el formato del texto vecino.
- Si "reemplazar" es "" y "buscar" es el párrafo entero, se elimina también el
  párrafo (su marca), como hace Word.

Uso
---
  python3 control_cambios.py aplicar --docx enviado.docx --cambios cambios.json \\
          --autor "Nombre Apellido" --out revisado.docx
  python3 control_cambios.py validar --docx revisado.docx --original enviado.docx \\
          --cambios cambios.json --autor "Nombre Apellido"
  python3 control_cambios.py aceptar --docx revisado.docx --out limpio.docx
  python3 control_cambios.py respaldo --docx enviado_RESPALDO.docx --cambios cambios.json \\
          --out revisado_RESPALDO.docx
  python3 control_cambios.py comentarios --docx enviado.docx
"""
import argparse
import copy
import datetime as dt
import difflib
import json
import re
import sys
import zipfile

try:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    sys.exit("Falta python-docx: pip install python-docx")
from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
T, R, P = qn("w:t"), qn("w:r"), qn("w:p")
INS, DEL = qn("w:ins"), qn("w:del")
RPR, PPR = qn("w:rPr"), qn("w:pPr")
CT_COMMENTS = "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"

# Orden de hijos que exige el esquema (Word no abre si se altera)
ORDEN_RPR = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
             "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
             "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
             "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
             "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def w(tag):
    return qn("w:" + tag)


# ------------------------------------------------------------------ utilidades
def _local(el):
    return etree.QName(el).localname


def texto_atomo(r):
    """Texto visible de un run (t, tab, br)."""
    out = []
    for c in r:
        n = _local(c)
        if n == "t":
            out.append(c.text or "")
        elif n == "tab":
            out.append("\t")
        elif n in ("br", "cr"):
            out.append("\n")
        elif n == "noBreakHyphen":
            out.append("-")
    return "".join(out)


def _ancestros_hasta_p(el):
    a = el.getparent()
    while a is not None and a.tag != P:
        yield a
        a = a.getparent()


def _visible(r):
    """Un run es visible si no está dentro de una eliminación."""
    for a in _ancestros_hasta_p(r):
        if a.tag in (DEL, w("moveFrom")):
            return False
    return True


def runs_de(p):
    """Runs del párrafo en orden de documento (no los de párrafos anidados)."""
    out = []
    for r in p.iter(R):
        q = r.getparent()
        while q is not None and q.tag != P:
            q = q.getparent()
        if q is p:
            out.append(r)
    return out


def _texto_r_todo(r):
    out = []
    for c in r:
        n = _local(c)
        if n in ("t", "delText"):
            out.append(c.text or "")
        elif n == "tab":
            out.append("\t")
        elif n in ("br", "cr"):
            out.append("\n")
    return "".join(out)


def normalizar(s):
    return re.sub(r"\s+", " ", s.replace(" ", " ")).strip()


def _norm_con_mapa(s):
    """Texto normalizado (espacios colapsados) y mapa índice-normalizado → índice-original."""
    out, mapa, prev_esp = [], [], False
    for i, ch in enumerate(s):
        esp = ch.isspace() or ch == " "
        if esp:
            if prev_esp:
                continue
            out.append(" ")
        else:
            out.append(ch)
        mapa.append(i)
        prev_esp = esp
    return "".join(out), mapa


def markdown_plano(s):
    """'a *b* **c**' -> ('a b c', [(ini, fin, italica, negrita), ...])."""
    plano, marcas, pos = [], [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*|([^*]+|\*)", s):
        if m.group(1) is not None:
            frag, it, bo = m.group(1), False, True
        elif m.group(2) is not None:
            frag, it, bo = m.group(2), True, False
        else:
            frag, it, bo = m.group(3), False, False
        if it or bo:
            marcas.append((pos, pos + len(frag), it, bo))
        plano.append(frag)
        pos += len(frag)
    return "".join(plano), marcas


def _tokens(s):
    return re.findall(r"\s+|\w+|[^\w\s]", s)


# ----------------------------------------------------------- manipulación XML
def _rpr_poner(rpr, nombre):
    """Asegura <w:nombre/> en el rPr respetando el orden del esquema."""
    if rpr.find(w(nombre)) is not None:
        return
    nuevo = etree.Element(w(nombre))
    idx = ORDEN_RPR.index(nombre)
    for i, c in enumerate(rpr):
        n = _local(c)
        if n in ORDEN_RPR and ORDEN_RPR.index(n) > idx:
            c.addprevious(nuevo)
            return
    rpr.append(nuevo)


def _rpr_copia(r):
    rpr = r.find(RPR) if r is not None else None
    rpr = copy.deepcopy(rpr) if rpr is not None else etree.Element(RPR)
    for c in list(rpr):
        if _local(c) in ("rPrChange", "ins", "del"):
            rpr.remove(c)
    return rpr


def atomizar(p):
    """Parte cada run en runs de un solo elemento de contenido (mismo rPr)."""
    for r in runs_de(p):
        rpr = r.find(RPR)
        contenido = [c for c in r if c.tag != RPR]
        ultimo = r
        for c in contenido[1:]:
            nr = etree.Element(R)
            for k, v in r.attrib.items():
                nr.set(k, v)
            if rpr is not None:
                nr.append(copy.deepcopy(rpr))
            nr.append(c)
            ultimo.addnext(nr)
            ultimo = nr


class Parrafo:
    """Vista de un párrafo como lista de átomos con offsets de texto visible."""

    def __init__(self, p):
        self.p = p
        self.recalcular()

    def recalcular(self):
        self.atomos = []  # (ini, fin, run)
        pos = 0
        for r in runs_de(self.p):
            if not _visible(r):
                continue
            t = texto_atomo(r)
            self.atomos.append((pos, pos + len(t), r))
            pos += len(t)
        self.texto = "".join(texto_atomo(a[2]) for a in self.atomos)

    def partir_en(self, off):
        """Parte el átomo que contiene `off` en su interior."""
        for ini, fin, r in self.atomos:
            if ini < off < fin:
                t = r.find(T)
                if t is None:
                    return
                corte = off - ini
                izq, der = t.text[:corte], t.text[corte:]
                nr = copy.deepcopy(r)
                t.text = izq
                t.set(XML_SPACE, "preserve")
                nt = nr.find(T)
                nt.text = der
                nt.set(XML_SPACE, "preserve")
                r.addnext(nr)
                self.recalcular()
                return


def _nuevo_run(texto, rpr):
    r = etree.Element(R)
    r.append(copy.deepcopy(rpr))
    partes = re.split(r"(\t|\n)", texto)
    for parte in partes:
        if parte == "":
            continue
        if parte == "\t":
            r.append(etree.Element(w("tab")))
        elif parte == "\n":
            r.append(etree.Element(w("br")))
        else:
            t = etree.SubElement(r, T)
            t.text = parte
            t.set(XML_SPACE, "preserve")
    return r


def _runs_formateados(texto, marcas, base_rpr, desfase=0):
    """Runs para `texto` aplicando cursiva/negrita de las marcas markdown."""
    if not texto:
        return []
    cortes = {0, len(texto)}
    for a, b, _, _ in marcas:
        for x in (a - desfase, b - desfase):
            if 0 < x < len(texto):
                cortes.add(x)
    cortes = sorted(cortes)
    out = []
    for a, b in zip(cortes, cortes[1:]):
        rpr = copy.deepcopy(base_rpr)
        for ma, mb, it, bo in marcas:
            if ma - desfase <= a and b <= mb - desfase:
                if it:
                    _rpr_poner(rpr, "i")
                if bo:
                    _rpr_poner(rpr, "b")
        out.append(_nuevo_run(texto[a:b], rpr))
    return out


class Marcador:
    def __init__(self, doc, autor, fecha):
        self.autor, self.fecha = autor, fecha
        ids = [int(v) for v in etree.XPath("//@w:id", namespaces={"w": W_NS})(doc.element)
               if str(v).lstrip("-").isdigit()]
        self.sig = max(ids + [999]) + 1

    def nuevo_id(self):
        self.sig += 1
        return str(self.sig - 1)

    def envoltura(self, tag):
        e = etree.Element(tag)
        e.set(w("id"), self.nuevo_id())
        e.set(w("author"), self.autor)
        e.set(w("date"), self.fecha)
        return e


def _envolver_del(runs, m):
    """Envuelve runs consecutivos en <w:del> (t -> delText). Devuelve los w:del.
    Lo que el mismo autor insertó en esta sesión simplemente se quita."""
    grupos, actual, propios = [], [], []
    for r in runs:
        padre = r.getparent()
        if padre.tag == INS and padre.get(w("author")) == m.autor:
            propios.append(r)
            continue
        if actual and actual[-1].getnext() is r:
            actual.append(r)
        else:
            if actual:
                grupos.append(actual)
            actual = [r]
    if actual:
        grupos.append(actual)
    dels = []
    for g in grupos:
        d = m.envoltura(DEL)
        g[0].addprevious(d)
        for r in g:
            for t in r.findall(T):
                t.tag = w("delText")
                t.set(XML_SPACE, "preserve")
            d.append(r)
        dels.append(d)
    for r in propios:
        padre = r.getparent()
        if dels or len(padre) > 1:
            padre.remove(r)
        else:  # deja el contenedor como ancla para lo que se inserte
            dels.append(padre)
            padre.remove(r)
    return dels


def _colocar_despues(ancla, nuevo, m):
    """Coloca `nuevo` justo después de `ancla` sin anidar w:ins dentro de w:ins."""
    padre = ancla.getparent()
    if padre is not None and padre.tag == INS:
        # partir la inserción ajena: lo que sigue a `ancla` pasa a un w:ins gemelo
        resto = []
        s = ancla.getnext()
        while s is not None:
            resto.append(s)
            s = s.getnext()
        padre.addnext(nuevo)
        if resto:
            gemelo = etree.Element(INS)
            for k, v in padre.attrib.items():
                gemelo.set(k, v)
            gemelo.set(w("id"), m.nuevo_id())
            for s in resto:
                gemelo.append(s)
            nuevo.addnext(gemelo)
        return
    ancla.addnext(nuevo)


def _colocar_antes(ancla, nuevo, m):
    padre = ancla.getparent()
    if padre is not None and padre.tag == INS:
        previo = ancla.getprevious()
        if previo is None:
            padre.addprevious(nuevo)
        else:
            _colocar_despues(previo, nuevo, m)
        return
    ancla.addprevious(nuevo)


# ------------------------------------------------------------------ comentarios
def _parte_comentarios(doc):
    for rel in doc.part.rels.values():
        if rel.reltype == RT.COMMENTS:
            return rel.target_part
    return None


def _xml_comentarios(parte):
    if hasattr(parte, "_element"):
        return parte._element, None
    raiz = etree.fromstring(parte.blob)
    return raiz, raiz


def agregar_comentario(doc, inicio, fin, texto, m):
    parte = _parte_comentarios(doc)
    if parte is None:
        xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               f'<w:comments xmlns:w="{W_NS}"/>').encode("utf-8")
        parte = Part(PackURI("/word/comments.xml"), CT_COMMENTS, xml, doc.part.package)
        doc.part.relate_to(parte, RT.COMMENTS)
    raiz, suelto = _xml_comentarios(parte)
    cid = m.nuevo_id()
    c = etree.SubElement(raiz, w("comment"))
    c.set(w("id"), cid)
    c.set(w("author"), m.autor)
    c.set(w("date"), m.fecha)
    c.set(w("initials"), "".join(x[0] for x in m.autor.split() if x)[:4].upper())
    cp = etree.SubElement(c, P)
    cr = etree.SubElement(cp, R)
    ct = etree.SubElement(cr, T)
    ct.text = texto
    ct.set(XML_SPACE, "preserve")
    if suelto is not None:
        parte._blob = etree.tostring(suelto, xml_declaration=True, encoding="UTF-8", standalone=True)
    ini = etree.Element(w("commentRangeStart"))
    ini.set(w("id"), cid)
    end = etree.Element(w("commentRangeEnd"))
    end.set(w("id"), cid)
    ref = etree.Element(R)
    etree.SubElement(ref, w("commentReference")).set(w("id"), cid)
    _subir_a_p(inicio).addprevious(ini)
    f = _subir_a_p(fin)
    f.addnext(end)
    end.addnext(ref)


def _subir_a_p(el):
    """Sube hasta el hijo directo del párrafo (los marcadores van ahí)."""
    while el.getparent() is not None and el.getparent().tag != P:
        el = el.getparent()
    return el


# ---------------------------------------------------------------- aplicar
def _todos_parrafos(doc):
    return list(doc.element.body.iter(P))


def _buscar(parrafos, buscar):
    """[(Parrafo, ini, fin)] de todas las apariciones (texto visible, espacios normalizados)."""
    objetivo = normalizar(buscar)
    halladas = []
    if not objetivo:
        return halladas
    for p in parrafos:
        vista = Parrafo(p)
        n, mapa = _norm_con_mapa(vista.texto)
        desde = 0
        while True:
            i = n.find(objetivo, desde)
            if i < 0:
                break
            fin_n = i + len(objetivo) - 1
            halladas.append((vista, mapa[i], mapa[fin_n] + 1))
            desde = i + 1
    return halladas


def _pista(parrafos, buscar):
    """Hasta dónde coincide `buscar`, para ubicar la diferencia."""
    palabras = normalizar(buscar).split(" ")
    todo = "\n".join(normalizar(Parrafo(p).texto) for p in parrafos)
    mejor = ""
    for k in range(len(palabras), 0, -1):
        pref = " ".join(palabras[:k])
        if pref in todo:
            mejor = pref
            break
    if not mejor:
        return "ninguna palabra inicial coincide"
    sig = palabras[len(mejor.split(' '))] if len(mejor.split(" ")) < len(palabras) else ""
    i = todo.find(mejor)
    real = todo[i + len(mejor): i + len(mejor) + 40].split("\n")[0]
    return f"coincide hasta «…{mejor[-50:]}»; ahí esperabas «{sig}» y el documento dice «{real.strip()[:40]}»"


def aplicar_reemplazo(vista, ini, fin, nuevo_md, m, borrar_parrafo=False):
    plano, marcas = markdown_plano(nuevo_md)
    original = vista.texto[ini:fin]
    a_tok, b_tok = _tokens(original), _tokens(plano)
    clave = lambda ts: [" " if t.isspace() else t for t in ts]
    sm = difflib.SequenceMatcher(None, clave(a_tok), clave(b_tok), autojunk=False)
    # offsets por token
    a_off, x = [0], 0
    for t in a_tok:
        x += len(t)
        a_off.append(x)
    b_off, x = [0], 0
    for t in b_tok:
        x += len(t)
        b_off.append(x)
    # une cambios separados solo por espacios o un signo ("et al." -> una sola marca)
    ops = []
    for op in sm.get_opcodes():
        if op[0] == "equal":
            continue
        if ops:
            hueco = "".join(a_tok[ops[-1][2]:op[1]])
            if len(hueco.strip()) <= 1 and op[1] - ops[-1][2] <= 2:
                ops[-1] = ("replace", ops[-1][1], op[2], ops[-1][3], op[4])
                continue
        ops.append(op)
    # partir átomos en todos los bordes
    for tag, i1, i2, j1, j2 in ops:
        for off in (ini + a_off[i1], ini + a_off[i2]):
            vista.partir_en(off)
    atomos = list(vista.atomos)
    creados = []
    n_del = n_ins = 0
    for tag, i1, i2, j1, j2 in ops:
        a0, a1 = ini + a_off[i1], ini + a_off[i2]
        borrados = [r for s, e, r in atomos if a0 <= s and e <= a1 and e > s]
        dels = _envolver_del(borrados, m) if borrados else []
        n_del += len(dels)
        creados += dels
        if j2 > j1:
            texto = plano[b_off[j1]:b_off[j2]]
            if borrados:
                base = _rpr_copia(borrados[0])
            else:
                antes = [r for s, e, r in atomos if e == a0 and e > s]
                despues = [r for s, e, r in atomos if s == a0 and e > s]
                base = _rpr_copia(antes[-1] if antes else (despues[0] if despues else None))
            ins = m.envoltura(INS)
            for r in _runs_formateados(texto, marcas, base, desfase=b_off[j1]):
                ins.append(r)
            if dels:
                _colocar_despues(dels[-1], ins, m)
            else:
                antes = [r for s, e, r in atomos if e == a0 and e > s]
                despues = [r for s, e, r in atomos if s == a0 and e > s]
                if antes:
                    _colocar_despues(antes[-1], ins, m)
                elif despues:
                    _colocar_antes(despues[0], ins, m)
                else:
                    vista.p.append(ins)
            n_ins += 1
            creados.append(ins)
    if borrar_parrafo:
        ppr = vista.p.find(PPR)
        if ppr is None:
            ppr = etree.Element(PPR)
            vista.p.insert(0, ppr)
        prpr = ppr.find(RPR)
        if prpr is None:
            prpr = etree.Element(RPR)
            ancla = ppr.find(w("sectPr"))
            if ancla is None:
                ancla = ppr.find(w("pPrChange"))
            ancla.addprevious(prpr) if ancla is not None else ppr.append(prpr)
        prpr.insert(0, m.envoltura(DEL))
    return creados, n_ins, n_del


def _es_titulo(p):
    est = p.find(PPR + "/" + w("pStyle")) if p.find(PPR) is not None else None
    v = est.get(w("val")) if est is not None else ""
    return bool(re.search(r"heading|t[ií]tulo|title", v or "", re.I))


def insertar_parrafo(vista_ancla, texto_md, m, formato_de="auto"):
    ancla = vista_ancla.p
    modelo = ancla
    if formato_de == "siguiente" or (formato_de == "auto" and _es_titulo(ancla)):
        s = ancla.getnext()
        while s is not None and (s.tag != P or _es_titulo(s) or not Parrafo(s).texto.strip()):
            s = s.getnext()
        if s is not None:
            modelo = s
    nuevo = etree.Element(P)
    ppr = copy.deepcopy(modelo.find(PPR)) if modelo.find(PPR) is not None else etree.Element(PPR)
    for c in list(ppr):
        if _local(c) in ("pPrChange", "sectPr"):
            ppr.remove(c)
    prpr = ppr.find(RPR)
    if prpr is None:
        prpr = etree.SubElement(ppr, RPR)
    for c in list(prpr):
        if _local(c) in ("ins", "del", "moveFrom", "moveTo", "rPrChange"):
            prpr.remove(c)
    prpr.insert(0, m.envoltura(INS))
    nuevo.append(ppr)
    vm = Parrafo(modelo)
    base = _rpr_copia(vm.atomos[0][2] if vm.atomos else None)
    plano, marcas = markdown_plano(texto_md)
    ins = m.envoltura(INS)
    for r in _runs_formateados(plano, marcas, base):
        ins.append(r)
    nuevo.append(ins)
    ancla.addnext(nuevo)
    return [ins]


def cmd_aplicar(a):
    doc = Document(a.docx)
    cambios = json.load(open(a.cambios, encoding="utf-8"))
    fecha = a.fecha or dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    m = Marcador(doc, a.autor, fecha)
    parrafos = _todos_parrafos(doc)
    ok, fallos, total_ins, total_del, n_com = 0, [], 0, 0, 0
    for k, c in enumerate(cambios, 1):
        try:
            if "insertar_despues" in c:
                hall = _buscar(parrafos, c["insertar_despues"])
                if not hall:
                    fallos.append((k, c["insertar_despues"], _pista(parrafos, c["insertar_despues"])))
                    continue
                vista = hall[0][0]
                atomizar(vista.p)
                creados = insertar_parrafo(Parrafo(vista.p), c.get("texto", ""), m,
                                           c.get("formato_de", "auto"))
                parrafos = _todos_parrafos(doc)
                total_ins += 1
                aviso = f" (el ancla aparece {len(hall)} veces: se usó la primera)" if len(hall) > 1 else ""
                print(f"✓ {k}: párrafo nuevo insertado{aviso}")
            else:
                buscar, nuevo = c["buscar"], c.get("reemplazar", "")
                donde = parrafos
                if "tabla" in c:   # apuntar a una celda: {"tabla": 3, "fila": 2, "columna": 4} (desde 1)
                    tablas = list(doc.element.body.iter(w("tbl")))
                    nt, nf, nc = int(c["tabla"]), int(c.get("fila", 0)), int(c.get("columna", 0))
                    if not 1 <= nt <= len(tablas):
                        fallos.append((k, buscar, f"el documento tiene {len(tablas)} tablas; no hay tabla {nt}"))
                        continue
                    filas = list(tablas[nt - 1].iter(w("tr")))
                    if not 1 <= nf <= len(filas):
                        fallos.append((k, buscar, f"la tabla {nt} tiene {len(filas)} filas; no hay fila {nf}"))
                        continue
                    celdas = [x for x in filas[nf - 1] if x.tag == w("tc")]
                    if not 1 <= nc <= len(celdas):
                        fallos.append((k, buscar, f"la fila {nf} de la tabla {nt} tiene {len(celdas)} "
                                                  f"columnas; no hay columna {nc}"))
                        continue
                    donde = list(celdas[nc - 1].iter(P))
                elif "ocurrencia" in c:   # la n-ésima aparición del texto (desde 1)
                    donde = parrafos
                hall = _buscar(donde, buscar)
                if "ocurrencia" in c and hall:
                    n = int(c["ocurrencia"])
                    hall = hall[n - 1:n] if n <= len(hall) else []
                if not hall:
                    fallos.append((k, buscar, _pista(parrafos, buscar)))
                    continue
                objetivos = hall if c.get("todas") else hall[:1]
                creados = []
                # de atrás hacia adelante para no mover offsets del mismo párrafo
                for vista, ini, fin in reversed(objetivos):
                    atomizar(vista.p)
                    v2 = Parrafo(vista.p)
                    entero = normalizar(v2.texto) == normalizar(buscar) and normalizar(nuevo) == ""
                    cr, ni, nd = aplicar_reemplazo(v2, ini, fin, nuevo, m, borrar_parrafo=entero)
                    creados = cr + creados
                    total_ins += ni
                    total_del += nd
                if not creados:
                    print(f"· {k}: «buscar» y «reemplazar» son iguales; nada que marcar")
                    continue
                extra = ""
                if len(hall) > 1 and not c.get("todas"):
                    extra = f" (aparece {len(hall)} veces: se cambió la primera; agrega contexto o \"todas\": true)"
                print(f"✓ {k}: marcado{extra}")
            if c.get("comentario") and creados:
                if a.sin_comentarios:
                    print(f"  · comentario de {k} omitido (--sin-comentarios)")
                else:
                    agregar_comentario(doc, creados[0], creados[-1], c["comentario"], m)
                    n_com += 1
            ok += 1
        except Exception as e:  # un cambio raro no tumba los demás
            fallos.append((k, c.get("buscar") or c.get("insertar_despues", "?"), f"error: {e}"))
    doc.save(a.out)
    print(f"\nGuardado: {a.out}  ·  autor: {a.autor}  ·  fecha: {fecha}")
    print(f"Aplicados {ok} de {len(cambios)}  ·  inserciones {total_ins}  ·  eliminaciones {total_del}"
          f"  ·  comentarios {n_com}")
    if fallos:
        print(f"\nNO ENCONTRADOS ({len(fallos)}): el texto de «buscar» no es idéntico al del Word.")
        for k, b, pista in fallos:
            print(f"✗ {k}: «{b[:70]}{'…' if len(b) > 70 else ''}»\n    {pista}")
        print("Corrige esos «buscar» copiando el texto desde `articulo_docx.py texto` y vuelve a "
              "aplicar SOLO esos cambios sobre el archivo de salida.")
    return len(fallos)


# ---------------------------------------------------------------- vistas de texto
def _incluir_factory(modo, autor=None):
    """modo 'aceptar': acepta todo; 'rechazar_autor': deshace solo los de `autor`
    (los ajenos se toman aceptados); 'rechazar': deshace todo."""
    def incluir(r):
        for anc in _ancestros_hasta_p(r):
            if anc.tag in (INS, w("moveTo")):
                propio = autor is None or anc.get(w("author")) == autor
                if modo == "rechazar" or (modo == "rechazar_autor" and propio):
                    return False
            elif anc.tag in (DEL, w("moveFrom")):
                propio = autor is None or anc.get(w("author")) == autor
                if modo == "aceptar" or (modo == "rechazar_autor" and not propio):
                    return False
        return True
    return incluir


def texto_documento(doc, modo, autor=None):
    inc = _incluir_factory(modo, autor)
    partes = []
    for p in doc.element.body.iter(P):
        partes.append("".join(_texto_r_todo(r) for r in runs_de(p) if inc(r)))
    return normalizar(" ".join(partes))


def cmd_validar(a):
    errores = 0
    try:
        with zipfile.ZipFile(a.docx) as z:
            malo = z.testzip()
            if malo:
                print(f"✗ ZIP dañado en {malo}")
                errores += 1
            for n in z.namelist():
                if n.endswith(".xml") or n.endswith(".rels"):
                    etree.fromstring(z.read(n))
        doc = Document(a.docx)
        print("✓ El documento abre (ZIP íntegro, XML bien formado, python-docx lo carga)")
    except Exception as e:
        print(f"✗ El documento no abre: {e}")
        sys.exit(1)
    body = doc.element.body
    conteo = {}
    for el in body.iter(INS, DEL):
        au = el.get(w("author"), "(sin autor)")
        d = conteo.setdefault(au, {"ins": 0, "del": 0})
        d["ins" if el.tag == INS else "del"] += 1
    if not conteo:
        print("✗ No hay ninguna marca de revisión")
        errores += 1
    for au, d in conteo.items():
        print(f"  {au}: {d['ins']} inserciones, {d['del']} eliminaciones")
    if a.autor and a.autor not in conteo:
        print(f"✗ No hay marcas a nombre de «{a.autor}»")
        errores += 1
    parte = _parte_comentarios(doc)
    if parte is not None:
        raiz, _ = _xml_comentarios(parte)
        print(f"  comentarios en el archivo: {len(raiz.findall(w('comment')))}")
    ids = [e.get(w("id")) for e in body.iter(INS, DEL)]
    if len(ids) != len(set(ids)):
        print("✗ Hay w:id de revisión repetidos")
        errores += 1
    aceptado = texto_documento(doc, "aceptar")
    rechazado = texto_documento(doc, "rechazar_autor", a.autor)
    if a.original:
        orig = texto_documento(Document(a.original), "aceptar")
        if orig == rechazado:
            print("✓ Rechazando los cambios del autor se recupera el original: no hay cambios sin marcar")
        else:
            errores += 1
            print("✗ Hay texto cambiado SIN MARCAR (difiere del original):")
            sa, sb = orig.split(" "), rechazado.split(" ")
            for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, sa, sb, autojunk=False).get_opcodes():
                if tag != "equal":
                    ctx = " ".join(sa[max(0, i1 - 6):i1])
                    print(f"    …{ctx} [-{' '.join(sa[i1:i2])}-]{{+{' '.join(sb[j1:j2])}+}}")
    if a.cambios:
        cambios = json.load(open(a.cambios, encoding="utf-8"))
        for k, c in enumerate(cambios, 1):
            if "insertar_despues" in c:
                t = normalizar(markdown_plano(c.get("texto", ""))[0])
                bien = t in aceptado
                print(f"{'✓' if bien else '✗'} {k}: párrafo nuevo {'presente' if bien else 'NO aparece'} al aceptar")
                errores += 0 if bien else 1
                continue
            b = normalizar(c["buscar"])
            r = normalizar(markdown_plano(c.get("reemplazar", ""))[0])
            if b not in rechazado:
                print(f"✗ {k}: no se aplicó: «buscar» no está en el documento")
                errores += 1
                continue
            esta_r = (r in aceptado) if r else True
            if b and r and b in r:
                se_fue = True
            else:
                se_fue = aceptado.count(b) < rechazado.count(b) if b in rechazado else True
                if c.get("todas"):
                    se_fue = b not in aceptado
            bien = esta_r and se_fue
            errores += 0 if bien else 1
            det = []
            if not esta_r:
                det.append("el texto nuevo no aparece")
            if not se_fue:
                det.append("el texto viejo sigue igual de veces")
            print(f"{'✓' if bien else '✗'} {k}: {'correcto al aceptar' if bien else '; '.join(det)}")
    print("\nRESULTADO:", "VALIDADO" if errores == 0 else f"{errores} problema(s)")
    sys.exit(0 if errores == 0 else 1)


# ---------------------------------------------------------------- aceptar
def cmd_aceptar(a):
    doc = Document(a.docx)
    body = doc.element.body

    def propio(el):
        return a.autor is None or el.get(w("author")) == a.autor

    # párrafos cuya marca se eliminó: se funden con el siguiente
    for p in list(body.iter(P)):
        mk = p.find(PPR + "/" + RPR + "/" + w("del")) if p.find(PPR) is not None else None
        if mk is not None and propio(mk):
            sig = p.getnext()
            contenido = [c for c in p if c.tag != PPR]
            if sig is not None and sig.tag == P:
                ancla = sig.find(PPR)
                for c in reversed(contenido):
                    (ancla.addnext(c) if ancla is not None else sig.insert(0, c))
            p.getparent().remove(p)
    for el in list(body.iter(DEL, w("moveFrom"))):
        if el.getparent() is not None and el.getparent().tag == RPR:
            continue
        if propio(el):
            el.getparent().remove(el)
    for el in list(body.iter(INS, w("moveTo"))):
        if not propio(el):
            continue
        padre = el.getparent()
        if padre.tag == RPR:
            padre.remove(el)
            continue
        i = padre.index(el)
        for c in list(el):
            padre.insert(i, c)
            i += 1
        padre.remove(el)
    for el in list(body.iter(w("rPrChange"), w("pPrChange"))):
        el.getparent().remove(el)
    doc.save(a.out)
    print(f"ok: {a.out} (cambios aceptados{'' if a.autor is None else ' de ' + a.autor})")


# ---------------------------------------------------------------- respaldo
def cmd_respaldo(a):
    """Aplica el mismo JSON al _RESPALDO (formato neutro) y deja el resultado
    con los cambios ya aceptados: mismo texto que la versión de la revista,
    sin marcas, listo por si hay que reenviar a otra revista."""
    import os
    import tempfile
    fd, tmp = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        fallos = cmd_aplicar(argparse.Namespace(docx=a.docx, cambios=a.cambios, autor="respaldo",
                                                out=tmp, fecha=None, sin_comentarios=True))
        cmd_aceptar(argparse.Namespace(docx=tmp, out=a.out, autor=None))
    finally:
        os.remove(tmp)
    limpio = texto_documento(Document(a.out), "aceptar")
    faltan = 0
    for k, c in enumerate(json.load(open(a.cambios, encoding="utf-8")), 1):
        nuevo = c.get("texto") if "insertar_despues" in c else c.get("reemplazar", "")
        t = normalizar(markdown_plano(nuevo or "")[0])
        if t and t not in limpio:
            faltan += 1
            print(f"✗ {k}: el texto nuevo no aparece en el respaldo")
    print("RESPALDO:", "AL DÍA" if not (fallos or faltan) else
          f"{fallos} no encontrados, {faltan} sin aplicar: corrige y repite sobre el respaldo anterior")
    sys.exit(1 if (fallos or faltan) else 0)


# ---------------------------------------------------------------- comentarios
def cmd_comentarios(a):
    doc = Document(a.docx)
    parte = _parte_comentarios(doc)
    if parte is None:
        print("El documento no trae comentarios.")
        return
    raiz, _ = _xml_comentarios(parte)
    anclado, abiertos = {}, set()
    for el in doc.element.body.iter():
        if el.tag == w("commentRangeStart"):
            abiertos.add(el.get(w("id")))
        elif el.tag == w("commentRangeEnd"):
            abiertos.discard(el.get(w("id")))
        elif el.tag == T and abiertos:
            for i in abiertos:
                anclado[i] = anclado.get(i, "") + (el.text or "")
    for n, c in enumerate(raiz.findall(w("comment")), 1):
        cid = c.get(w("id"))
        txt = " ".join(normalizar("".join(t.text or "" for t in pp.iter(T))) for pp in c.iter(P))
        print(f"[{n}] {c.get(w('author'), '?')}: {txt.strip()}")
        if anclado.get(cid):
            ref = normalizar(anclado[cid])
            print(f"     sobre: «{ref[:160]}{'…' if len(ref) > 160 else ''}»")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("aplicar")
    p.add_argument("--docx", required=True)
    p.add_argument("--cambios", required=True, help="JSON con la lista de cambios")
    p.add_argument("--autor", required=True, help='Nombre exacto del autor: "Nombre Apellido"')
    p.add_argument("--out", required=True)
    p.add_argument("--fecha", help="ISO 8601, p. ej. 2026-10-03T10:00:00Z (por defecto, ahora)")
    p.add_argument("--sin-comentarios", action="store_true", help="ignora los campos 'comentario'")
    p.set_defaults(f=cmd_aplicar)
    p = sub.add_parser("validar")
    p.add_argument("--docx", required=True)
    p.add_argument("--original", help="el manuscrito enviado, para detectar cambios sin marcar")
    p.add_argument("--cambios", help="el mismo JSON usado en aplicar")
    p.add_argument("--autor")
    p.set_defaults(f=cmd_validar)
    p = sub.add_parser("aceptar")
    p.add_argument("--docx", required=True)
    p.add_argument("--out", required=True)
    p.add_argument("--autor", help="aceptar solo los cambios de este autor")
    p.set_defaults(f=cmd_aceptar)
    p = sub.add_parser("respaldo", help="mismo JSON sobre el _RESPALDO, ya aceptado (sin marcas)")
    p.add_argument("--docx", required=True, help="ARTICULO_[clave]_v8_RESPALDO.docx")
    p.add_argument("--cambios", required=True)
    p.add_argument("--out", required=True, help="ARTICULO_[clave]_v9_RESPALDO.docx")
    p.set_defaults(f=cmd_respaldo)
    p = sub.add_parser("comentarios")
    p.add_argument("--docx", required=True)
    p.set_defaults(f=cmd_comentarios)
    a = ap.parse_args()
    a.f(a)


if __name__ == "__main__":
    main()
