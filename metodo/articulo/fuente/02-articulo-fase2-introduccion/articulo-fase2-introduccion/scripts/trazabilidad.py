#!/usr/bin/env python3
"""
Matriz de trazabilidad de la Ruta del Artículo — empírico y de revisión (Acosta | IA & Research).

Un solo archivo, TRAZABILIDAD_[clave].xlsx, con siete columnas fijas. La crea la
Fase 2; las Fases 3 y 6 le agregan filas. Los ID de fuente (F001...) son los
del inventario.json y nunca se renumeran. Requiere openpyxl (pip install openpyxl).

Uso
---
  python3 trazabilidad.py crear     --out TRAZABILIDAD_clave.xlsx --filas filas.json [--inventario inventario.json]
  python3 trazabilidad.py agregar   --base TRAZABILIDAD_clave.xlsx --filas filas.json [--inventario inventario.json]
  python3 trazabilidad.py verificar TRAZABILIDAD_clave.xlsx [--inventario inventario.json]

filas.json: lista de objetos con estas claves (las mismas siete columnas):
  [{"cita": "García y López (2022)", "id": "F001", "ubicacion": "Introducción, bloque 2",
    "afirmacion": "...", "frase": "texto literal del abstract", "tipo": "contenido",
    "direccion": "no aplica"}]
  tipo: contenido | titulo | cifra      direccion: si | no | no informa | no aplica
  Para una cifra estadística: id "WEB", cita "INEI (2024)", frase = texto literal
  de la fuente, y en afirmacion el organismo, el año y el ámbito.

verificar (y también crear/agregar con --inventario) comprueba fila por fila:
  - que el ID exista en el inventario;
  - que una fila de tipo contenido tenga abstract y que la frase aparezca
    LITERALMENTE en él (sin distinguir mayúsculas ni espacios);
  - que una fila de tipo titulo no tenga frase de abstract;
  - que ninguna fila diga Dirección = no.
Marca ✗ lo que obliga a corregir y ! lo que se revisa a mano. Código 1 si hay ✗.
"""

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

COLUMNAS = ["Cita en el texto", "ID", "Ubicación", "Afirmación que sostiene",
            "Frase textual del abstract o de la fuente", "Tipo", "Dirección coincide"]
CLAVES = ["cita", "id", "ubicacion", "afirmacion", "frase", "tipo", "direccion"]
ANCHOS = [24, 8, 22, 50, 60, 11, 12]
TIPOS = {"contenido", "titulo", "cifra"}
DIRECCIONES = {"si", "no", "no aplica", "no informa"}


def norm(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(re.sub(r"[^a-z0-9%.,;:()\- ]+", " ", t).split())


def cargar_filas(ruta):
    datos = json.loads(Path(ruta).read_text(encoding="utf-8"))
    if isinstance(datos, dict):
        datos = datos.get("filas", [datos])
    filas = []
    for d in datos:
        d = {norm(k).replace(" ", "_"): v for k, v in d.items()}
        alias = {"cita_en_el_texto": "cita", "afirmacion_que_sostiene": "afirmacion",
                 "direccion_coincide": "direccion", "seccion_o_hipotesis": "ubicacion",
                 "bloque": "ubicacion"}
        d = {alias.get(k, k): v for k, v in d.items()}
        fila = [str(d.get(k, "") or "").strip() for k in CLAVES]
        fila[5] = norm(fila[5]) or "contenido"
        fila[6] = norm(fila[6]) or "no aplica"
        filas.append(fila)
    return filas


def cargar_inventario(ruta):
    if not ruta:
        return None
    datos = json.loads(Path(ruta).read_text(encoding="utf-8"))
    datos = datos.get("fuentes", datos) if isinstance(datos, dict) else datos
    return {r.get("id"): r for r in datos}


def revisar(filas, inventario):
    errores, avisos = [], []
    for n, f in enumerate(filas, 2):  # fila 1 = encabezado
        cita, fid, _, afirm, frase, tipo, dirc = f
        tipo, dirc = norm(tipo), norm(dirc)
        etq = f"fila {n} ({cita or fid})"
        if tipo not in TIPOS:
            errores.append(f"{etq}: Tipo '{tipo}' no es contenido, titulo ni cifra")
        if dirc not in DIRECCIONES:
            errores.append(f"{etq}: Dirección '{dirc}' no es si, no, no informa o no aplica")
        if dirc == "no":
            errores.append(f"{etq}: Dirección = no. Corregir el texto antes de entregar")
        if not afirm:
            errores.append(f"{etq}: falta la afirmación que sostiene")
        if tipo == "contenido" and not frase:
            errores.append(f"{etq}: tipo contenido sin frase textual")
        if tipo == "titulo" and frase:
            avisos.append(f"{etq}: tipo titulo con frase; solo se le atribuye lo que dice el título")
        if tipo == "cifra":
            if not frase:
                errores.append(f"{etq}: cifra sin frase textual de la fuente")
            if not re.search(r"(19|20)\d{2}", cita + afirm):
                avisos.append(f"{etq}: cifra sin año visible")
            continue
        if inventario is None:
            continue
        r = inventario.get(fid)
        if r is None:
            errores.append(f"{etq}: el ID {fid or '(vacío)'} no está en el inventario")
            continue
        if tipo == "contenido":
            ab = r.get("abstract") or ""
            if tipo == "contenido" and r.get("resumen_parcial"):
                errores.append(f"{etq}: {fid} tiene el resumen RECORTADO; no se cita por contenido "
                               "hasta pegar el resumen completo en el inventario")
            if len(ab.strip()) < 120:
                errores.append(f"{etq}: {fid} no tiene abstract; solo puede citarse a nivel de título")
            elif frase and norm(frase).strip(" .") not in norm(ab):
                errores.append(f"{etq}: la frase no aparece literal en el abstract de {fid}")
        if dirc == "no informa":
            avisos.append(f"{etq}: la fuente no informa la dirección; la hipótesis queda con respaldo escaso")
    return errores, avisos


def leer_xlsx(ruta):
    from openpyxl import load_workbook
    ws = load_workbook(ruta).active
    filas = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if any(c not in (None, "") for c in fila[:7]):
            filas.append([str(c if c is not None else "").strip() for c in fila[:7]])
    return filas


def escribir_xlsx(ruta, filas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Trazabilidad"
    ws.append(COLUMNAS)
    for c in ws[1]:
        c.font = Font(bold=True, name="Calibri", size=11)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        ws.append(f)
    for i, w in enumerate(ANCHOS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{max(1, len(filas) + 1)}"
    wb.save(ruta)


def informar(filas, errores, avisos):
    tipos = [norm(f[5]) for f in filas]
    print(f"{len(filas)} filas: {tipos.count('contenido')} por contenido, "
          f"{tipos.count('titulo')} por título, {tipos.count('cifra')} cifras")
    for e in errores:
        print("✗ " + e)
    for a in avisos:
        print("! " + a)
    if not errores:
        print("✓ sin errores")


def main():
    ap = argparse.ArgumentParser(description="Matriz de trazabilidad (7 columnas)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("crear")
    c.add_argument("--out", required=True)
    c.add_argument("--filas", required=True)
    c.add_argument("--inventario")
    a = sub.add_parser("agregar")
    a.add_argument("--base", required=True)
    a.add_argument("--filas", required=True)
    a.add_argument("--out", help="por defecto, la misma base")
    a.add_argument("--inventario")
    v = sub.add_parser("verificar")
    v.add_argument("xlsx")
    v.add_argument("--inventario")
    args = ap.parse_args()
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.exit("Falta openpyxl: pip install openpyxl")

    inv = cargar_inventario(getattr(args, "inventario", None))
    if args.cmd == "verificar":
        filas = leer_xlsx(args.xlsx)
    elif args.cmd == "crear":
        if Path(args.out).exists():
            sys.exit(f"{args.out} ya existe: usa 'agregar' para no perder filas")
        filas = cargar_filas(args.filas)
    else:
        previas = leer_xlsx(args.base)
        vistas = {(norm(f[1]), norm(f[2]), norm(f[3])) for f in previas}
        nuevas = [f for f in cargar_filas(args.filas) if (norm(f[1]), norm(f[2]), norm(f[3])) not in vistas]
        print(f"{len(previas)} filas previas, {len(nuevas)} nuevas")
        filas = previas + nuevas

    errores, avisos = revisar(filas, inv)
    if args.cmd != "verificar":
        destino = args.out if args.cmd == "crear" else (args.out or args.base)
        escribir_xlsx(destino, filas)
        print(f"-> {destino}")
    informar(filas, errores, avisos)
    sys.exit(1 if errores else 0)


if __name__ == "__main__":
    main()
