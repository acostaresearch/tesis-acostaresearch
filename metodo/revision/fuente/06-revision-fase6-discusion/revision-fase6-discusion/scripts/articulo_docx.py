#!/usr/bin/env python3
"""Motor único del Word de la Ruta del Artículo — empírico y de revisión (Acosta | IA & Research).

Una sola herramienta para todas las fases, para que el documento salga idéntico
sin importar quién lo arme ni en qué sesión. Solo necesita python-docx
(pip install python-docx). No usa Node, ni LibreOffice, ni rutas fijas.

Subcomandos
-----------
  nuevo      Crea el documento desde un .md (Fase 0 ficha, Fase 2 artículo).
  agregar    Añade una sección .md a un .docx existente y devuelve la versión
             siguiente. No reescribe lo anterior: lo deja tal cual.
  esquema    Resumen compacto de un .docx (títulos, palabras, tablas, figuras,
             referencias, pendientes). Para que la fase siguiente no tenga que
             leer el artículo entero.
  texto      Vuelca el texto de una sección concreta (o de todo) del .docx.
  verificar  Control de calidad: formato, numeración de tablas y figuras, que
             cada tabla se cite antes de aparecer, referencias ordenadas y sin
             duplicados, citas del texto sin referencia, pendientes, palabras.
  contar     Cuenta palabras de un .md o un .docx (opcionalmente por sección).
  normalizar Convierte un Word que no salió de la ruta (títulos en negrita sin
             estilo) al formato de la ruta o de una revista (--perfil), sin tocar el texto.
  reemplazar Cambia un texto en el .docx y en su _RESPALDO a la vez (p. ej. un
             [POR CONFIRMAR] que la fase siguiente resolvió). Sin control de cambios.
  perfil     `perfil plantilla` imprime el REVISTA_[clave].json a rellenar desde la
             guía de autores; `perfil validar` lo comprueba.

Perfil de revista: con --perfil REVISTA_[clave].json, `nuevo` y `agregar` arman
el Word con el formato de la revista y, en paralelo, el mismo texto con el
formato neutro de la ruta en *_RESPALDO.docx (para reenviar a otra revista si
hay rechazo).

Sintaxis del .md de entrada (deliberadamente pequeña)
-----------------------------------------------------
  # Título 1 / ## Título 2 / ### Título 3
  Párrafo normal (una línea o varias seguidas; una línea vacía lo cierra).
  *cursiva*  **negrita**  [POR CONFIRMAR: texto]  -> rojo y negrita
  - viñeta
  [TITULO] Título del artículo           -> centrado, negrita
  [CENTRADO] texto                       -> centrado sin sangría
  [SIN SANGRIA] texto                    -> párrafo sin sangría de primera línea
  [TABLA 2] Título de la tabla           -> seguido de una tabla con barras |
  | Col A | Col B |
  |---|---|
  | x | 1.23 |
  [NOTA] Texto de la nota                 -> "Nota." en cursiva, 11 pt
  [FIGURA 1] Título de la figura | ruta/imagen.png
  [SALTO]                                 -> salto de página
  # Referencias  (o # References)        -> cada línea siguiente es una referencia
                                            con sangría francesa; al agregar se
                                            fusionan con la lista existente,
                                            ordenadas y sin duplicados.
                                            (Modo antiguo: sigue funcionando.)

Citas por clave (recomendado; con --fuentes inventario.json)
------------------------------------------------------------
  [F012]                 cita parentética          (Hair et al., 2019) / [3] / (3)
  [F012:n]               narrativa                 Hair et al. (2019) / Hair et al. [3]
  [F012:p. 45]           con página                (Hair et al., 2019, p. 45)
  [F012:n, p. 45]        narrativa con página
  [F012][AR51541814]     grupo: una sola cita       (Chen et al., 2025; Hair et al., 2019)
El motor las escribe en el estilo_cita del perfil (APA 7 sin perfil; Chicago
autor-fecha 18.ª, Vancouver, IEEE) y GENERA la lista de Referencias con todas las
fuentes citadas: no se escribe a mano. Vancouver/IEEE numeran por orden de primera
aparición en todo el documento y se renumeran solas en cada `agregar`. Una clave que
no está en el inventario sale como [POR CONFIRMAR: fuente CLAVE] en rojo y
`verificar` la marca ✗. Cada cita queda entre marcadores ocultos de Word y el mapa
marcador→claves (más una copia mínima de las fuentes citadas) se guarda en las
propiedades personalizadas del .docx: por eso `normalizar --perfil` puede reconvertir
todas las citas y la lista a la norma de otra revista (reenvío), aunque el autor haya
editado el Word, mientras no borre o pegue de nuevo las citas.
"""
import argparse
import copy
import os
import re
import sys
import unicodedata

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("Falta python-docx. Instálalo con:  pip install python-docx")

import json

# Formato neutro de la ruta (el del respaldo). Un perfil de revista lo cambia.
NEUTRO = dict(fuente="Times New Roman", tamano=12, interlineado=1.5, margen_cm=2.54, justificado=True,
              sangria_cm=1.27, tabla_tamano=11, numerar_lineas=False, numerar_paginas=False,
              titulo1_centrado=False, titulo1_mayusculas=False, titulo2_cursiva=True,
              papel="A4", numero_pagina="pie")
FUENTE = NEUTRO["fuente"]
TAMANO = NEUTRO["tamano"]
TAM_TABLA = NEUTRO["tabla_tamano"]     # 10 en modo compacto
INTERLINEADO = NEUTRO["interlineado"]  # 1.15 en modo compacto (Ficha y documentos de trabajo)
FMT = dict(NEUTRO)


def aplicar_formato(fmt=None):
    """Fija el formato activo: el neutro, o el del perfil de la revista."""
    global FUENTE, TAMANO, TAM_TABLA, INTERLINEADO, FMT
    FMT = dict(NEUTRO)
    for k, v in (fmt or {}).items():
        if k in FMT and v not in (None, ""):
            FMT[k] = v
    FUENTE, TAMANO = FMT["fuente"], float(FMT["tamano"])
    TAM_TABLA, INTERLINEADO = float(FMT["tabla_tamano"]), float(FMT["interlineado"])


def leer_perfil(ruta):
    if not ruta:
        return None
    with open(ruta, encoding="utf-8") as f:
        return json.load(f)


def _romano(s):
    v = dict(I=1, V=5, X=10, L=50, C=100)
    t = 0
    for i, ch in enumerate(s):
        t += -v[ch] if i + 1 < len(s) and v[ch] < v[s[i + 1]] else v[ch]
    return t


def _ruta_respaldo(ruta):
    b, e = os.path.splitext(ruta)
    return b + "_RESPALDO" + e
NEGRO = RGBColor(0, 0, 0)
ROJO = RGBColor(0xC0, 0, 0)
REF_TITULOS = ("referencias", "references", "referencias bibliograficas", "bibliografia", "bibliography",
               "lista de referencias", "reference list", "literatura citada", "works cited")
RE_TABLA_CAP = re.compile(r"^(Tabla|Table)\s+(\d+)\s*$")
RE_FIG_CAP = re.compile(r"^(Figura|Figure|Gráfico|Grafico)\s+(\d+|[IVXLC]+)\s*$")
RE_PEND = re.compile(r"\[POR CONFIRMAR[^\]]*\]|\[PENDIENTE[^\]]*\]")
RE_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[POR CONFIRMAR[^\]]*\]|\[PENDIENTE[^\]]*\]|"
                       r"(?:\[[A-Z]{1,3}\d{2,}(?::[^\[\]\n]*)?\])+)")


# ----------------------------------------------------------------- utilidades
def _norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def _palabras(texto):
    return len(re.findall(r"\b[\wÁÉÍÓÚáéíóúÑñÜü'’-]+\b", texto))


def _es_titulo_ref(texto):
    return _norm(texto) in REF_TITULOS


def _siguiente_version(ruta):
    base, ext = os.path.splitext(ruta)
    m = re.search(r"_v(\d+)$", base)
    if m:
        return f"{base[:m.start()]}_v{int(m.group(1)) + 1}{ext}"
    return f"{base}_v2{ext}"


# ------------------------------------------------------------------- estilos
def _sin_metadatos(doc):
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""


def _numeracion(sec):
    sp = sec._sectPr
    for e in sp.findall(qn("w:lnNumType")):
        sp.remove(e)
    if FMT["numerar_lineas"]:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        # orden del esquema: ... pgSz, pgMar, paperSrc, pgBorders, lnNumType, pgNumType, cols ...
        ancla = None
        for tag in ("w:pgBorders", "w:paperSrc", "w:pgMar", "w:pgSz"):
            ancla = sp.find(qn(tag))
            if ancla is not None:
                break
        if ancla is not None:
            ancla.addnext(ln)
        else:
            sp.insert(0, ln)
    if FMT["numerar_paginas"]:
        zona = sec.header if str(FMT["numero_pagina"]).lower().startswith("enc") else sec.footer
        pie = zona.paragraphs[0] if zona.paragraphs else zona.add_paragraph()
        if "PAGE" not in pie._p.xml:
            pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = pie.add_run()
            for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
                if tipo:
                    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); r._r.append(fc)
                else:
                    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = texto; r._r.append(it)


def preparar_estilos(doc, margen_cm=None):
    """A4, márgenes, fuente, interlineado y títulos en negro, según el formato activo."""
    _sin_metadatos(doc)
    margen_cm = margen_cm or float(FMT["margen_cm"])
    for sec in doc.sections:
        if str(FMT["papel"]).lower() in ("carta", "letter"):
            sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        else:
            sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(margen_cm)
        sec.left_margin = sec.right_margin = Cm(margen_cm)
        _numeracion(sec)
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(TAMANO)
    normal.font.color.rgb = NEGRO
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    pf = normal.paragraph_format
    pf.line_spacing = INTERLINEADO
    pf.space_after = Pt(6 if INTERLINEADO > 1.2 else 3)
    for nivel, cursiva in ((1, False), (2, bool(FMT["titulo2_cursiva"])), (3, True)):
        h = st[f"Heading {nivel}"]
        h.font.name = FUENTE
        h.font.size = Pt(TAMANO)
        h.font.bold = True
        h.font.italic = cursiva
        if nivel == 1:
            h.font.all_caps = bool(FMT["titulo1_mayusculas"])
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if FMT["titulo1_centrado"] else WD_ALIGN_PARAGRAPH.LEFT
        h.font.color.rgb = NEGRO
        rpr = h.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(a), FUENTE)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing = INTERLINEADO
        if INTERLINEADO < 1.2:
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True


def _run(p, texto, bold=False, italic=False, size=None, color=None):
    r = p.add_run(texto)
    r.bold, r.italic = bold or None, italic or None
    r.font.name = FUENTE
    r.font.color.rgb = color or NEGRO
    if size:
        r.font.size = Pt(size)
    return r


def _inline(p, texto, size=None, base_italic=False):
    for trozo in RE_INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            _run(p, trozo[2:-2], bold=True, italic=base_italic, size=size)
        elif trozo.startswith("*") and trozo.endswith("*"):
            _run(p, trozo[1:-1], italic=not base_italic, size=size)
        elif RE_PEND.fullmatch(trozo):
            _run(p, trozo, bold=True, size=size, color=ROJO)
        elif RE_GRUPO.fullmatch(trozo):
            _cita_marcada(p, trozo, size=size)
        else:
            _run(p, trozo, italic=base_italic, size=size)


def _parrafo(doc, texto, sangria=True, centro=False, bold=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else (
        WD_ALIGN_PARAGRAPH.JUSTIFY if FMT["justificado"] else WD_ALIGN_PARAGRAPH.LEFT)
    pf.first_line_indent = Cm(float(FMT["sangria_cm"])) if (sangria and not centro) else Cm(0)
    if bold:
        _run(p, texto, bold=True)
    else:
        _inline(p, texto)
    return p


def _titulo(doc, texto, nivel):
    h = doc.add_heading(level=nivel)
    _inline(h, texto)
    for r in h.runs:
        r.font.color.rgb = NEGRO
        r.font.name = FUENTE
    return h


def _referencia(doc, texto):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.left_indent = Cm(1.27)
    pf.first_line_indent = Cm(-1.27)
    _inline(p, texto)
    return p


def _bordes_apa(tabla, n_filas):
    """APA 7: regla arriba y abajo del encabezado y al cierre; nada vertical."""
    tbl = tabla._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)
    for i, fila in enumerate(tabla.rows):
        for celda in fila.cells:
            tcPr = celda._tc.get_or_add_tcPr()
            cb = OxmlElement("w:tcBorders")
            for lado in ("top", "bottom", "left", "right"):
                e = OxmlElement(f"w:{lado}")
                regla = (lado == "top" and i == 0) or (lado == "bottom" and i in (0, n_filas - 1))
                if regla:
                    e.set(qn("w:val"), "single")
                    e.set(qn("w:sz"), "6")
                    e.set(qn("w:color"), "000000")
                else:
                    e.set(qn("w:val"), "nil")
                cb.append(e)
            tcPr.append(cb)


def _es_numero(s):
    s = s.strip()
    if s in ("—", "-", "–", ""):
        return True
    # cifras con signo, %, asteriscos, rangos (18–45), M (DE) e intervalos [0.41, 0.58]
    return bool(re.fullmatch(r"[\[(]?[<>≤≥=]?\s*[-−–+]?\.?\d[\d.,]*%?\**[\])]?"
                             r"(\s*(?:[-–—]|,|\()\s*[-−–+]?\.?\d[\d.,]*%?\**[\])]?)*", s))


def _anchos(filas, m, util, pt):
    """Ancho por columna proporcional al contenido, sin partir palabras."""
    cm_car = 0.0185 * pt          # ancho medio de un carácter en Times
    largo, minimo = [], []
    for j in range(m):
        celdas = [re.sub(r"\*", "", f[j]) if j < len(f) else "" for f in filas]
        palabra = max((len(w) for c in celdas for w in c.split()), default=4)
        cuerpo = [len(c) for c in celdas[1:]] or [len(celdas[0])]
        largo.append(max(8, sum(cuerpo) / len(cuerpo) + 0.3 * max(cuerpo)))
        minimo.append((palabra + 1.5) * cm_car)
    total = sum(largo)
    w = [util * l / total for l in largo]
    for _ in range(5):            # subir las que parten palabras, quitando a las holgadas
        falta = sum(max(0, mn - x) for x, mn in zip(w, minimo))
        if falta <= 0.01:
            break
        holgura = [max(0, x - mn) for x, mn in zip(w, minimo)]
        h = sum(holgura) or 1
        w = [max(x, mn) if x < mn else x - falta * hg / h for x, mn, hg in zip(w, minimo, holgura)]
    return w


def _tabla(doc, filas):
    n, m = len(filas), max(len(f) for f in filas)
    t = doc.add_table(rows=n, cols=m)
    t.autofit = False
    util = (21.59 if str(FMT["papel"]).lower() in ("carta", "letter") else 21) - 2 * float(FMT["margen_cm"])
    anchos = [Cm(w) for w in _anchos(filas, m, util, TAM_TABLA)]
    for j, col in enumerate(t.columns):
        col.width = anchos[j]
    tblW = t._tbl.tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        t._tbl.tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(util / 2.54 * 1440)))
    tblW.set(qn("w:type"), "dxa")
    col_num = []
    for j in range(m):
        cuerpo = [f[j] for f in filas[1:] if j < len(f) and f[j].strip()]
        col_num.append(j > 0 and bool(cuerpo) and sum(_es_numero(x) for x in cuerpo) >= 0.6 * len(cuerpo))
    for i, fila in enumerate(filas):
        for j in range(m):
            txt = fila[j] if j < len(fila) else ""
            c = t.cell(i, j)
            c.width = anchos[j]
            p = c.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            centro = col_num[j]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
            if i == 0:
                _run(p, txt.replace("**", ""), bold=True, size=TAM_TABLA)
            else:
                _inline(p, txt, size=TAM_TABLA)
    _bordes_apa(t, n)
    return t


def _leyenda(doc, tipo, num, titulo):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12 if INTERLINEADO > 1.2 else 6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    _run(p, f"{tipo} {num}", bold=True)
    q = doc.add_paragraph()
    q.paragraph_format.first_line_indent = Cm(0)
    q.paragraph_format.keep_with_next = True
    _inline(q, titulo, base_italic=True)


def _nota(doc, texto, idioma):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)
    _run(p, "Note. " if idioma == "en" else "Nota. ", italic=True, size=max(TAM_TABLA, TAMANO - 1))
    _inline(p, texto, size=max(TAM_TABLA, TAMANO - 1))


# ------------------------------------------------------------------- parser
def parsear_md(texto):
    """Devuelve (bloques, referencias)."""
    bloques, refs = [], []
    lineas = texto.replace("\r\n", "\n").split("\n")
    i, en_refs, buf = 0, False, []

    def cerrar():
        if buf:
            bloques.append(("p", " ".join(s.strip() for s in buf)))
            buf.clear()

    while i < len(lineas):
        ln = lineas[i].rstrip()
        s = ln.strip()
        if en_refs:
            if s.startswith("# "):
                en_refs = False
                continue
            if s:
                refs.append(s.lstrip("-• ").strip())
            i += 1
            continue
        if not s:
            cerrar(); i += 1; continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            cerrar()
            if len(m.group(1)) == 1 and _es_titulo_ref(m.group(2)):
                en_refs = True
            else:
                bloques.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1; continue
        m = re.match(r"^\[(TABLA|TABLE)\s+(\d+)\]\s*(.*)$", s, re.I)
        if m:
            cerrar()
            tipo = "Table" if m.group(1).upper() == "TABLE" else "Tabla"
            num, tit = m.group(2), m.group(3)
            i += 1
            filas = []
            while i < len(lineas) and (lineas[i].strip().startswith("|") or not lineas[i].strip()):
                f = lineas[i].strip()
                i += 1
                if not f:
                    if filas:
                        break
                    continue
                celdas = [c.strip() for c in f.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in celdas if c):
                    continue
                filas.append(celdas)
            bloques.append(("tabla", tipo, num, tit, filas))
            continue
        m = re.match(r"^\[(FIGURA|FIGURE)\s+(\d+)\]\s*(.*)$", s, re.I)
        if m:
            cerrar()
            tipo = "Figure" if m.group(1).upper() == "FIGURE" else "Figura"
            partes = m.group(3).split("|")
            bloques.append(("figura", tipo, m.group(2), partes[0].strip(),
                            partes[1].strip() if len(partes) > 1 else ""))
            i += 1; continue
        m = re.match(r"^\[(NOTA|NOTE)\]\s*(.*)$", s, re.I)
        if m:
            cerrar(); bloques.append(("nota", m.group(2))); i += 1; continue
        for marca, tipo in (("[TITULO]", "titulo"), ("[CENTRADO]", "centro"),
                            ("[SIN SANGRIA]", "sin")):
            if s.upper().startswith(marca):
                cerrar(); bloques.append((tipo, s[len(marca):].strip())); break
        else:
            if s.upper() == "[SALTO]":
                cerrar(); bloques.append(("salto",))
            elif re.match(r"^[-•]\s+", s):
                cerrar(); bloques.append(("vineta", re.sub(r"^[-•]\s+", "", s)))
            else:
                buf.append(s)
        i += 1
    cerrar()
    return bloques, refs


def escribir_bloques(doc, bloques, idioma, base_dir):
    for b in bloques:
        k = b[0]
        if k == "p":
            _parrafo(doc, b[1])
        elif k == "h":
            _titulo(doc, b[2], b[1])
        elif k == "titulo":
            _parrafo(doc, b[1], centro=True, bold=True)
        elif k == "centro":
            _parrafo(doc, b[1], centro=True)
        elif k == "sin":
            _parrafo(doc, b[1], sangria=False)
        elif k == "vineta":
            p = doc.add_paragraph(style="List Bullet")
            _inline(p, b[1])
        elif k == "salto":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif k == "nota":
            _nota(doc, b[1], idioma)
        elif k == "tabla":
            _, tipo, num, tit, filas = b
            _leyenda(doc, tipo, num, tit)
            if filas:
                _tabla(doc, filas)
        elif k == "figura":
            _, tipo, num, tit, ruta = b
            _leyenda(doc, tipo, num, tit)
            rp = ruta if os.path.isabs(ruta) else os.path.join(base_dir, ruta)
            if ruta and os.path.exists(rp):
                p = doc.add_paragraph()
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.add_run().add_picture(rp, width=Cm(15))
            else:
                _parrafo(doc, f"[POR CONFIRMAR: falta la imagen de la {tipo.lower()} {num} ({ruta or 'sin ruta'})]", sangria=False)


# --------------------------------------------------------- referencias
def _bloque_refs(doc):
    """(índice del título de Referencias en body, lista de párrafos de refs)."""
    body = doc.element.body
    hijos = list(body.iterchildren())
    for idx, el in enumerate(hijos):
        if el.tag == qn("w:p"):
            txt = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
            if _es_titulo_ref(txt):
                refs = []
                for h in hijos[idx + 1:]:
                    if h.tag != qn("w:p"):
                        if h.tag == qn("w:tbl"):
                            break
                        continue
                    ppr = h.find(qn("w:pPr"))
                    est = ppr.find(qn("w:pStyle")).get(qn("w:val")) if ppr is not None and ppr.find(qn("w:pStyle")) is not None else ""
                    t = "".join(x.text or "" for x in h.iter(qn("w:t"))).strip()
                    if est.lower() in ("heading1", "ttulo1") or (t and _titulo_implicito(h, t) == 1) \
                            or RE_TABLA_CAP.match(t) or RE_FIG_CAP.match(t):
                        break      # termina la lista: empieza otra sección (anexos, tablas al final)
                    if t:
                        refs.append(h)
                return el, refs
    return None, []


def _clave_ref(texto):
    doi = re.search(r"10\.\d{4,9}/\S+", texto)
    if doi:
        return doi.group(0).rstrip(".").lower()
    return _norm(texto)[:70]


def fusionar_referencias(doc, nuevas, idioma):
    titulo_el, existentes = _bloque_refs(doc)
    textos = {}
    for el in existentes:
        txt = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
        textos[_clave_ref(txt)] = ("xml", el, txt)
    añadidas = 0
    for r in nuevas:
        k = _clave_ref(r)
        if k not in textos:
            textos[k] = ("md", None, r)
            añadidas += 1
    if not textos:
        return 0, 0
    # retirar las existentes y el título; reconstruir al final
    for el in existentes:
        el.getparent().remove(el)
    if titulo_el is not None:
        titulo_el.getparent().remove(titulo_el)
    _titulo(doc, "References" if idioma == "en" else "Referencias", 1)
    vals = list(textos.values())
    if vals and all(re.match(r"^\[?\d+[\].]", v[2].strip()) for v in vals):
        # IEEE / Vancouver: orden de citación, por número
        orden = sorted(vals, key=lambda v: int(re.match(r"^\[?(\d+)", v[2].strip()).group(1)))
    else:
        orden = sorted(vals, key=lambda v: _norm(re.sub(r"^\W+", "", v[2])))
    for tipo, el, txt in orden:
        if tipo == "xml":
            doc.element.body.insert(len(doc.element.body) - 1, el)  # antes de sectPr
        else:
            _referencia(doc, txt)
    return añadidas, len(orden)


# ------------------------------------------------------- citas por clave
# En el .md se cita por clave del inventario: [F012] parentética, [F012:n] narrativa,
# [F012:p. 45] con página, [F012][AR51541814] cita agrupada. El motor escribe cada
# cita entre dos marcadores ocultos de Word (bookmark "_AcCitaN") y guarda en las
# propiedades personalizadas del .docx (AcostaCitas_001...) el mapa marcador→claves,
# el estilo y una copia mínima de las fuentes citadas. Con eso, en cada `nuevo`,
# `agregar` y `normalizar` se recorren TODAS las citas en el orden del documento,
# se renumeran (Vancouver, IEEE), se reescriben en el estilo vigente y se regenera
# la lista de referencias. Word conserva marcadores y propiedades al guardar.
RE_CLAVE = r"\[[A-Z]{1,3}\d{2,}(?::[^\[\]\n]*)?\]"
RE_GRUPO = re.compile(r"(?:" + RE_CLAVE + r")+")
RE_ITEM = re.compile(r"\[([A-Z]{1,3}\d{2,})(?::([^\[\]\n]*))?\]")
PREFIJO_CITA, PREFIJO_REF, PROP = "_AcCita", "_AcRef_", "AcostaCitas_"
FUENTES = {}          # clave (id o alias) -> registro del inventario
ESTILO = None         # "apa" | "chicago" | "vancouver" | "ieee" (None: el guardado o APA)
IDIOMA_CITA = None
_CTX = {"n": 0, "id": 0, "nuevas": {}}
ESTILOS = {"apa": "APA 7", "chicago": "Chicago autor-fecha 18.ª", "vancouver": "Vancouver", "ieee": "IEEE",
           "zotero": "marcas {Autor, año} para el RTF Scan de Zotero (sin lista: la pone Zotero)"}
CAMPOS_SNAP = ("id", "alias", "autores", "autores_completos", "anio", "titulo", "fuente", "doi", "volumen",
               "numero", "paginas", "art_no", "tipo", "editorial", "editores", "url", "consultado")


def estilo_de(texto):
    t = _norm(texto or "")
    for k in ("zotero", "vancouver", "ieee", "chicago", "apa"):
        if k in t:
            return k
    if any(x in t for x in ("icmje", "nlm")):
        return "vancouver"
    if t:
        print(f"  ! Estilo de cita «{texto}» sin formateador: se usa APA 7 (revisar con estilos_cita.md).")
    return "apa"


def cargar_fuentes(ruta):
    """Lee inventario.json (lista o {"fuentes": [...]}) y lo indexa por id y alias."""
    if not ruta:
        return {}
    with open(ruta, encoding="utf-8") as f:
        datos = json.load(f)
    datos = datos.get("fuentes", datos) if isinstance(datos, dict) else datos
    if not isinstance(datos, list) or (datos and not isinstance(datos[0], dict)):
        sys.exit(f"{ruta} no es un inventario de fuentes (se espera el inventario.json de inventario_fuentes.py).")
    idx = {}
    for r in datos:
        for k in [r.get("id")] + list(r.get("alias") or []):
            if k:
                idx[str(k)] = r
    return idx


# ---- propiedades personalizadas del .docx (mapa invisible)
_NS_CP = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
_NS_VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"


def _parte_custom(doc, crear=False):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT, CONTENT_TYPE as CT
    pkg = doc.part.package
    try:
        return pkg.part_related_by(RT.CUSTOM_PROPERTIES)
    except KeyError:
        if not crear:
            return None
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    blob = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<Properties xmlns="{_NS_CP}" xmlns:vt="{_NS_VT}"></Properties>').encode()
    parte = Part(PackURI("/docProps/custom.xml"), CT.OFC_CUSTOM_PROPERTIES, blob, pkg)
    pkg.relate_to(parte, RT.CUSTOM_PROPERTIES)
    return parte


def leer_mapa(doc):
    import xml.etree.ElementTree as ET
    parte = _parte_custom(doc)
    if parte is None:
        return {}
    raiz = ET.fromstring(parte.blob)
    trozos = sorted((p.get("name"), "".join(p.itertext())) for p in raiz
                    if (p.get("name") or "").startswith(PROP))
    if not trozos:
        return {}
    try:
        return json.loads("".join(t for _, t in trozos))
    except ValueError:
        print("  ! El mapa de citas del documento está dañado: se ignora.")
        return {}


def guardar_mapa(doc, mapa):
    import xml.etree.ElementTree as ET
    parte = _parte_custom(doc, crear=True)
    ET.register_namespace("", _NS_CP)
    ET.register_namespace("vt", _NS_VT)
    raiz = ET.fromstring(parte.blob)
    for p in list(raiz):
        if (p.get("name") or "").startswith(PROP):
            raiz.remove(p)
    pid = max([int(p.get("pid") or 1) for p in raiz] or [1]) + 1
    texto = json.dumps(mapa, ensure_ascii=False, separators=(",", ":"))
    for i in range(0, len(texto), 240):            # Word recorta las propiedades a 255 caracteres
        p = ET.SubElement(raiz, f"{{{_NS_CP}}}property", fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}",
                          pid=str(pid), name=f"{PROP}{i // 240 + 1:04d}")
        ET.SubElement(p, f"{{{_NS_VT}}}lpwstr").text = texto[i:i + 240]
        pid += 1
    parte._blob = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + ET.tostring(raiz)


# ---- marcadores al escribir el .md
def _citas_inicio(doc):
    """Prepara la numeración de marcadores para escribir citas nuevas en `doc`."""
    ids = [int(b.get(qn("w:id"))) for b in doc.element.body.iter(qn("w:bookmarkStart"))
           if (b.get(qn("w:id")) or "").isdigit()]
    nums = [int(b.get(qn("w:name"))[len(PREFIJO_CITA):]) for b in doc.element.body.iter(qn("w:bookmarkStart"))
            if (b.get(qn("w:name")) or "").startswith(PREFIJO_CITA) and b.get(qn("w:name"))[len(PREFIJO_CITA):].isdigit()]
    mapa = leer_mapa(doc)
    nums += [int(k[len(PREFIJO_CITA):]) for k in mapa.get("citas", {}) if k[len(PREFIJO_CITA):].isdigit()]
    _CTX.update(n=max(nums or [0]), id=max(ids or [0]) + 1000, nuevas={})


def _marcador(nombre, inicio=True, bid=None):
    e = OxmlElement("w:bookmarkStart" if inicio else "w:bookmarkEnd")
    e.set(qn("w:id"), str(bid))
    if inicio:
        e.set(qn("w:name"), nombre)
    return e


def _cita_marcada(p, spec, size=None):
    _CTX["n"] += 1
    _CTX["id"] += 1
    nombre, bid = f"{PREFIJO_CITA}{_CTX['n']}", _CTX["id"]
    _CTX["nuevas"][nombre] = spec
    p._p.append(_marcador(nombre, True, bid))
    _run(p, spec, size=size)                 # provisional: _procesar_citas lo reescribe
    p._p.append(_marcador(nombre, False, bid))


# ---- autores
_ORG = re.compile(r"\b(organi[sz]ation|organizaci[oó]n|ministerio|ministry|institut\w*|universi\w*|association|"
                  r"asociaci[oó]n|council|consejo|group|grupo|committee|comit[eé]|comisi[oó]n|commission|agency|"
                  r"agencia|fondo|office|oficina|naciones|nations|bank|banco|foundation|fundaci[oó]n|department|"
                  r"departamento|society|sociedad|centre|center|centro|unesco|oecd|ocde|cepal|unicef|"
                  r"world health|oms|inei|minedu|sunedu|concytec)\b", re.I)


def _iniciales(pila):
    pila = (pila or "").strip()
    if re.fullmatch(r"[A-Z]{1,4}", pila):                       # WoS: 'ME'
        return [c + "." for c in pila]
    out = []
    for tok in re.split(r"[\s.]+", pila):
        if not tok or tok[0].islower():                          # partículas: de, van, da
            continue
        partes = [x for x in tok.split("-") if x]
        out.append("-".join(x[0].upper() + "." for x in partes))
    return out


def _autor(s):
    s = re.sub(r"\s*\(\d+\)\s*$", "", (s or "").strip()).strip(" ;")
    if not s:
        return None
    if "," in s:
        ap, pila = [x.strip() for x in s.split(",", 1)]
    else:
        m = re.match(r"^(.*\S)\s+((?:[A-Z]\.?-?\s?){1,4})$", s)
        if m and not _ORG.search(s):
            ap, pila = m.group(1), m.group(2)
        elif _ORG.search(s) or len(s.split()) > 4 or (s.isupper() and " " not in s):
            return {"org": s}
        else:
            partes = s.split()
            ap, pila = (partes[-1], " ".join(partes[:-1])) if len(partes) > 1 else (s, "")
    completo = bool(re.search(r"[^\W\d_]{2,}(?![.\w])", pila)) and not re.fullmatch(r"[A-Z]{1,4}", pila)
    return {"ap": ap, "pila": pila if completo else "", "ini": _iniciales(pila)}


def _autores(r):
    base = [a for a in (r.get("autores") or []) if a]
    comp = [a for a in (r.get("autores_completos") or []) if a]
    lista = comp if comp and (len(comp) == len(base) or not base) else base
    return [x for x in (_autor(a) for a in lista) if x]


def _ap(a):
    return a.get("org") or a["ap"]


def _fmt_apa(a):
    return a["org"] if "org" in a else f"{a['ap']}, {' '.join(a['ini'])}".rstrip(", ")


def _fmt_ieee(a):
    return a["org"] if "org" in a else (" ".join(a["ini"]) + " " + a["ap"]).strip()


def _fmt_van(a):
    return a["org"] if "org" in a else (a["ap"] + " " + "".join(i.replace(".", "").replace("-", "")
                                                                 for i in a["ini"])).strip()


def _fmt_chi(a, primero):
    if "org" in a:
        return a["org"]
    pila = a["pila"] or " ".join(a["ini"])
    return f"{a['ap']}, {pila}".rstrip(", ") if primero else f"{pila} {a['ap']}".strip()


def _y(idioma):
    return "and" if idioma == "en" else "y"


def _unir(nombres, conj, coma_final=True, coma_dos=None):
    if len(nombres) <= 1:
        return "".join(nombres)
    if len(nombres) == 2:
        coma = (conj == "&") if coma_dos is None else coma_dos
        return f"{nombres[0]}{',' if coma_final and coma else ''} {conj} {nombres[1]}"
    return ", ".join(nombres[:-1]) + ("," if coma_final else "") + f" {conj} {nombres[-1]}"


def _anio(r, idioma, sufijo=""):
    a = r.get("anio") or ("n.d." if idioma == "en" else "s. f.")
    return a + (("-" if not r.get("anio") else "") + sufijo if sufijo else "")


def _etiqueta(r, estilo, idioma, conj):
    """Nombre en el texto: 1-2 autores todos; 3 o más, et al. (APA, Chicago 18.ª, IEEE, Vancouver)."""
    aus = _autores(r)
    if not aus:
        t = (r.get("titulo") or "?").split()
        return "“" + " ".join(t[:4]) + ("…" if len(t) > 4 else "") + "”"
    if len(aus) == 1:
        return _ap(aus[0])
    if len(aus) == 2:
        return f"{_ap(aus[0])} {conj} {_ap(aus[1])}"
    return f"{_ap(aus[0])} et al."


def _tc(s):
    """Revista en MAYÚSCULAS (WoS) → tipo título."""
    if not s or not s.isupper() or len(s) < 5:
        return s or ""
    chicas = {"of", "and", "the", "in", "for", "on", "at", "by", "to", "a", "an", "de", "la", "y", "en",
              "del", "el", "los", "las", "e", "und", "fur", "et"}
    pal = s.lower().split()
    return " ".join(w if (i and w in chicas) else w[:1].upper() + w[1:] for i, w in enumerate(pal))


def _pags(p, guion="–"):
    return re.sub(r"\s*[-–—]+\s*", guion, p or "")


def _loc(loc, estilo):
    loc = (loc or "").strip()
    if not loc:
        return ""
    if estilo == "chicago":
        return _pags(re.sub(r"^pp?\.\s*", "", loc))
    if re.fullmatch(r"[\d\s\-–]+", loc):
        return ("pp. " if re.search(r"[-–]", loc) else "p. ") + _pags(loc)
    return re.sub(r"^(pp?)\.?\s*", lambda m: m.group(1) + ". ", _pags(loc))


def _parse_spec(spec):
    items = []
    for clave, mod in RE_ITEM.findall(spec):
        mod = (mod or "").strip()
        narr = False
        m = re.match(r"^n(?:\s*[,:;]\s*|$)(.*)$", mod)
        if m:
            narr, mod = True, m.group(1).strip()
        items.append({"clave": clave, "narr": narr, "loc": mod})
    return items


# ---- referencias
def _doi_url(doi):
    return "https://doi.org/" + doi if doi else ""


def _ref_segmentos(r, estilo, idioma, num=None, sufijo=""):
    """Lista de (texto, 'i'|''|'rojo') con la referencia completa en el estilo."""
    en = idioma == "en"
    aus = _autores(r)
    tipo = r.get("tipo") or ("articulo" if r.get("fuente") else "otro")
    tit = (r.get("titulo") or "").strip().rstrip(".")
    fte = _tc(r.get("fuente") or "")
    vol, nro, pag, artno = r.get("volumen", ""), r.get("numero", ""), r.get("paginas", ""), r.get("art_no", "")
    edit, doi, url = r.get("editorial", ""), r.get("doi", ""), r.get("url", "")
    eds = [x for x in (_autor(e) for e in (r.get("editores") or [])) if x]
    anio = _anio(r, idioma, sufijo)
    S = []
    add = lambda t, k="": S.append((t, k)) if t else None
    punto = lambda t: t if re.search(r"[.?!]$", t) else t + "."
    falta = (not fte and not doi and not url and tipo not in ("libro",)) or not tit
    if estilo == "apa":
        if aus:
            nombres = [_fmt_apa(a) for a in aus]
            if len(nombres) >= 21:
                add(", ".join(nombres[:19]) + ", . . . " + nombres[-1])
            else:
                add(_unir(nombres, "&"))
            add(f" ({anio}). ")
        if tipo == "articulo" or tipo not in ("libro", "capitulo", "congreso", "web", "otro"):
            add(punto(tit) if aus else punto(tit) + f" ({anio}).")
            if fte:
                add(" ")
                add(fte + (", " + vol if vol else ""), "i")
                add(f"({nro})" if vol and nro else "")
                if pag:
                    add(", " + _pags(pag))
                elif artno:
                    add(f", {'Article' if en else 'Artículo'} {artno}")
                add(".")
        elif tipo in ("capitulo", "congreso"):
            add(punto(tit) + (f" ({anio})." if not aus else ""))
            if fte:
                add(" In " if en else " En ")
                if eds:
                    ed = _unir([(" ".join(e["ini"]) + " " + e["ap"]).strip() if "org" not in e else e["org"]
                                for e in eds], "&")
                    add(f"{ed} ({'Eds.' if len(eds) > 1 else 'Ed.'}), ")
                add(fte, "i")
                add(f" (pp. {_pags(pag)})." if pag else ".")
            if edit:
                add(" " + punto(edit))
        else:   # libro, web, otro
            add(punto(tit), "i")
            if not aus:
                add(f" ({anio}).")
            if tipo == "web" and fte:
                add(" " + punto(fte))
            elif edit:
                add(" " + punto(edit))
            elif fte:
                add(" " + punto(fte))
        if doi:
            add(" " + _doi_url(doi))
        elif url:
            if r.get("consultado"):
                add(f" Retrieved {r['consultado']}, from {url}" if en else f" Recuperado el {r['consultado']} de {url}")
            else:
                add(" " + url)
    elif estilo == "chicago":
        if aus:
            nombres = [_fmt_chi(a, i == 0) for i, a in enumerate(aus)]
            if len(nombres) >= 7:
                nombres = nombres[:3]
                add(", ".join(nombres) + ", et al.")
            else:
                add(punto(_unir(nombres, _y(idioma), True, True)))
            add(f" {anio}. ")
        if tipo in ("libro", "otro") and not (tipo == "otro" and fte):
            add(punto(tit), "i")
            if not aus:
                add(f" {anio}.")
            if edit:
                add(" " + punto(edit))
        else:
            add(f"“{punto(tit)}”")
            if not aus:
                add(f" {anio}.")
            if tipo in ("capitulo", "congreso") and fte:
                add(" In " if en else " En ")
                add(fte, "i")
                if eds:
                    add((", edited by " if en else ", editado por ") +
                        _unir([_fmt_chi(e, False) for e in eds], _y(idioma), False))
                add(f", {_pags(pag)}." if pag else ".")
                if edit:
                    add(" " + punto(edit))
            elif tipo == "web":
                if fte:
                    add(" " + punto(fte))
                if r.get("consultado"):
                    add(f" Accessed {r['consultado']}." if en else f" Consultado el {r['consultado']}.")
            elif fte:
                add(" ")
                add(fte, "i")
                if vol:
                    add(" " + vol + (f" ({nro})" if nro else ""))
                if pag or artno:
                    add(": " + (_pags(pag) if pag else artno))
                add(".")
        if doi:
            add(" " + _doi_url(doi) + ".")
        elif url:
            add(" " + url + ".")
    elif estilo == "ieee":
        add(f"[{num}] ")
        if aus:
            nombres = [_fmt_ieee(a) for a in aus]
            add((nombres[0] + " et al.") if len(nombres) > 6 else _unir(nombres, "and", len(nombres) > 2))
            add(", ")
        if tipo in ("libro", "otro") and not (tipo == "otro" and fte):
            add(tit, "i")
            add(". " + (f"{edit}, " if edit else "") + f"{anio}.")
        else:
            add(f"“{tit},” ")
            if tipo in ("capitulo", "congreso") and fte:
                add("in ")
                add(fte, "i")
                if eds:
                    add(", " + _unir([_fmt_ieee(e) for e in eds], "and", len(eds) > 2) +
                        (", Eds." if len(eds) > 1 else ", Ed."))
                add(". " + (f"{edit}, " if edit else "") + anio)
                add(f", pp. {_pags(pag)}" if pag else "")
            elif tipo == "web":
                add((fte + ". " if fte else "") + (f"Accessed: {r['consultado']}. " if r.get("consultado") else ""))
                add("[Online]. Available: " + (url or _doi_url(doi)))
                url = doi = ""
            else:
                if fte:
                    add(fte, "i")
                    add(", ")
                partes = ([f"vol. {vol}"] if vol else []) + ([f"no. {nro}"] if nro else []) + \
                         ([f"pp. {_pags(pag)}"] if pag and re.search(r"[-–]", pag) else [f"p. {pag}"] if pag else
                          [f"Art. no. {artno}"] if artno else []) + [anio]
                add(", ".join(partes))
            if doi:
                add(f", doi: {doi}.")
            elif url:
                add(f". [Online]. Available: {url}")
            else:
                add(".")
    else:   # vancouver
        add(f"{num}. ")
        if aus:
            nombres = [_fmt_van(a) for a in aus]
            add(", ".join(nombres[:6]) + (", et al." if len(nombres) > 6 else "."))
            add(" ")
        if tipo in ("libro", "otro") and not (tipo == "otro" and fte):
            add(punto(tit) + " " + (f"{edit}; " if edit else "") + f"{anio}.")
        elif tipo in ("capitulo", "congreso") and fte:
            add(punto(tit) + (" In: " if en else " En: "))
            if eds:
                add(", ".join(_fmt_van(e) for e in eds) + (", editors. " if en else ", editores. "))
            add(punto(fte) + " " + (f"{edit}; " if edit else "") + f"{anio}." + (f" p. {_pags(pag, '-')}." if pag else ""))
        elif tipo == "web":
            add(f"{tit} [Internet]. " + (f"{fte}; " if fte else "") + f"{anio}" +
                (f" [{'cited' if en else 'citado'} {r['consultado']}]" if r.get("consultado") else "") + ". ")
            add(("Available from: " if en else "Disponible en: ") + (url or _doi_url(doi)))
            url = doi = ""
        else:
            add(punto(tit) + " ")
            if fte:
                add(punto(fte) + " ")
            add(anio + (f";{vol}" if vol else "") + (f"({nro})" if nro else "") +
                (f":{_pags(pag, '-')}" if pag else f":{artno}" if artno else "") + ".")
        if doi:
            add(f" doi:{doi}")
        elif url:
            add(f" {'Available from' if en else 'Disponible en'}: {url}")
    if falta:
        add(" ")
        add("[POR CONFIRMAR: " + ("título y " if not tit else "") + "datos de publicación]", "rojo")
    return [(re.sub(r"\s{2,}", " ", t), k) for t, k in S]


# ---- procesamiento del documento
def _rango(nums, estilo):
    nums = sorted(set(nums))
    trozos, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            trozos.append((nums[i], nums[j]))
        else:
            trozos += [(n, None) for n in nums[i:j + 1]]
        i = j + 1
    if estilo == "ieee":
        return ", ".join(f"[{a}]–[{b}]" if b else f"[{a}]" for a, b in trozos)
    return ",".join(f"{a}–{b}" if b else f"{a}" for a, b in trozos)


def _texto_cita(items, estilo, idioma, numeros, sufijos):
    """Devuelve [(texto, rojo?)] de una cita (posible grupo)."""
    out, faltan = [], []
    ok = []
    for it in items:
        r = FUENTES.get(it["clave"])
        if r is None:
            faltan.append(it["clave"])
        else:
            ok.append((it, r))
    if ok:
        narr = any(it["narr"] for it, _ in ok)
        if estilo == "zotero":      # formato que reconoce Zotero → Herramientas → RTF Scan
            if narr:
                trozos = [f"{_etiqueta(r, estilo, idioma, '&')} {{{r.get('anio') or 'n.d.'}"
                          + (", " + _loc(it["loc"], "chicago") if it["loc"] else "") + "}" for it, r in ok]
                out.append((_unir(trozos, _y(idioma), False), False))
            else:
                out.append(("{" + "; ".join(f"{_etiqueta(r, estilo, idioma, '&')}, {r.get('anio') or 'n.d.'}"
                                            + (", " + _loc(it["loc"], "chicago") if it["loc"] else "")
                                            for it, r in ok) + "}", False))
        elif estilo in ("apa", "chicago"):
            conj_p = "&" if estilo == "apa" else _y(idioma)
            sep = ", " if estilo == "apa" else " "
            if not narr:
                orden = sorted(ok, key=lambda x: (_norm(_etiqueta(x[1], estilo, idioma, conj_p)), x[1].get("anio") or ""))
                partes, prev = [], None
                for it, r in orden:
                    et = _etiqueta(r, estilo, idioma, conj_p)
                    an = _anio(r, idioma, sufijos.get(r["id"], "")) + (", " + _loc(it["loc"], estilo) if it["loc"] else "")
                    if et == prev:
                        partes[-1] += ", " + an
                    else:
                        partes.append(f"{et}{sep}{an}")
                    prev = et
                out.append(("(" + "; ".join(partes) + ")", False))
            else:
                trozos = [f"{_etiqueta(r, estilo, idioma, _y(idioma))} ({_anio(r, idioma, sufijos.get(r['id'], ''))}"
                          + (", " + _loc(it["loc"], estilo) if it["loc"] else "") + ")" for it, r in ok]
                out.append((_unir(trozos, _y(idioma), False), False))
        else:
            con_loc = [(it, r) for it, r in ok if it["loc"]]
            if narr:
                trozos = []
                for it, r in ok:
                    n, l = numeros[r["id"]], _loc(it["loc"], estilo)
                    ref = f"[{n}{', ' + l if l else ''}]" if estilo == "ieee" else f"({n}{', ' + l if l else ''})"
                    trozos.append(f"{_etiqueta(r, estilo, idioma, _y(idioma))} {ref}")
                out.append((_unir(trozos, _y(idioma), False), False))
            elif estilo == "ieee":
                partes = [f"[{numeros[r['id']]}, {_loc(it['loc'], estilo)}]" for it, r in con_loc]
                resto = [numeros[r["id"]] for it, r in ok if not it["loc"]]
                out.append((", ".join(([_rango(resto, estilo)] if resto else []) + partes), False))
            else:
                resto = [numeros[r["id"]] for it, r in ok if not it["loc"]]
                partes = [f"{numeros[r['id']]}, {_loc(it['loc'], estilo)}" for it, r in con_loc]
                out.append(("(" + "; ".join(([_rango(resto, estilo)] if resto else []) + partes) + ")", False))
    for k in faltan:
        out.append((("" if not out else " ") + f"[POR CONFIRMAR: fuente {k}]", True))
    return out, faltan


def _reescribir(start, trozos):
    """Sustituye los runs entre el marcador de inicio y su fin por los trozos nuevos."""
    bid = start.get(qn("w:id"))
    medio, e = [], start.getnext()
    while e is not None and not (e.tag == qn("w:bookmarkEnd") and e.get(qn("w:id")) == bid):
        medio.append(e)
        e = e.getnext()
    if e is None:
        return False
    plantilla = next((m.find(qn("w:rPr")) for m in medio if m.tag == qn("w:r") and m.find(qn("w:rPr")) is not None), None)
    for m in medio:
        if m.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:ins"), qn("w:smartTag")):
            start.getparent().remove(m)
    ancla = start
    for texto, rojo in trozos:
        r = OxmlElement("w:r")
        rpr = copy.deepcopy(plantilla) if plantilla is not None else OxmlElement("w:rPr")
        for tag in ("w:b", "w:color", "w:i"):
            for x in rpr.findall(qn(tag)):
                rpr.remove(x)
        if rojo:
            b = OxmlElement("w:b"); rpr.append(b)
            c = OxmlElement("w:color"); c.set(qn("w:val"), "C00000"); rpr.append(c)
        _ordenar_rpr(rpr)
        r.append(rpr)
        t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = texto
        r.append(t)
        ancla.addnext(r)
        ancla = r
    return True


_ORDEN_RPR = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
              "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
              "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
              "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def _ordenar_rpr(rpr):
    hijos = list(rpr)
    for h in hijos:
        rpr.remove(h)
    clave = lambda h: _ORDEN_RPR.index(h.tag.split("}")[1]) if h.tag.split("}")[1] in _ORDEN_RPR else 99
    for h in sorted(hijos, key=clave):
        rpr.append(h)


def _citas_del_doc(doc, mapa):
    """[(start, items)] en el orden del documento, solo las citas que siguen vivas."""
    vivas = []
    for b in doc.element.body.iter(qn("w:bookmarkStart")):
        nombre = b.get(qn("w:name")) or ""
        if not nombre.startswith(PREFIJO_CITA) or nombre not in mapa:
            continue
        bid, e, texto, fin = b.get(qn("w:id")), b.getnext(), "", False
        while e is not None:
            if e.tag == qn("w:bookmarkEnd") and e.get(qn("w:id")) == bid:
                fin = True
                break
            texto += "".join(t.text or "" for t in e.iter(qn("w:t")))
            e = e.getnext()
        if fin and texto.strip():          # vacío: el autor borró la cita en Word
            vivas.append((b, _parse_spec(mapa[nombre])))
        elif not fin:
            print(f"  ! La cita {nombre} quedó partida entre párrafos: se deja como está.")
    return vivas


def _es_generada(el):
    return any((b.get(qn("w:name")) or "").startswith(PREFIJO_REF) for b in el.iter(qn("w:bookmarkStart")))


def _poner_lista(doc, entradas, idioma, numerica):
    """entradas: [(clave, segmentos)]. Regenera la lista conservando las referencias escritas a mano."""
    titulo_el, existentes = _bloque_refs(doc)
    viejas = [el for el in existentes if _es_generada(el)]
    manuales = [el for el in existentes if not _es_generada(el)]
    if not entradas and not viejas:
        return 0
    for el in viejas:
        el.getparent().remove(el)
    if not entradas:                   # marcas para Zotero: la bibliografía la inserta Zotero
        if titulo_el is not None and not manuales:
            titulo_el.getparent().remove(titulo_el)
        return 0
    if titulo_el is None:
        titulo_el = _titulo(doc, "References" if idioma == "en" else "Referencias", 1)._p
    elif not manuales:
        txt = "References" if idioma == "en" else "Referencias"
        ts = list(titulo_el.iter(qn("w:t")))
        if ts and _norm("".join(t.text or "" for t in ts)) in ("referencias", "references"):
            ts[0].text = txt
            for t in ts[1:]:
                t.text = ""
    nuevos = []
    for k, (clave, segs) in enumerate(entradas):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent = Cm(1.27 if not numerica else 0.75)
        pf.first_line_indent = Cm(-1.27 if not numerica else -0.75)
        p._p.append(_marcador(f"{PREFIJO_REF}{clave}"[:40], True, 900000 + k))
        p._p.append(_marcador("", False, 900000 + k))
        for texto, est in segs:
            _run(p, texto, italic=est == "i", bold=est == "rojo", color=ROJO if est == "rojo" else None)
        nuevos.append((p._p, "".join(t for t, _ in segs)))
    texto_de = lambda el: "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
    todos = nuevos + [(el, texto_de(el)) for el in manuales]
    if not numerica:
        todos.sort(key=lambda x: _norm(re.sub(r"^\W+", "", x[1])))
    elif manuales:
        print(f"  ! Hay {len(manuales)} referencias escritas a mano en un estilo numérico: quedan al final, sin número."
              " Pásalas al inventario y cítalas por clave.")
    ancla = titulo_el
    for el, _ in todos:
        el.getparent().remove(el)
        ancla.addnext(el)
        ancla = el
    return len(entradas)


def _procesar_citas(doc, idioma, avisar=True):
    """Reescribe todas las citas por clave del documento y regenera la lista. Devuelve claves que faltan."""
    mapa = leer_mapa(doc)
    citas = dict(mapa.get("citas", {}))
    citas.update(_CTX.get("nuevas") or {})
    _CTX["nuevas"] = {}
    if not citas:
        return []
    estilo = ESTILO or mapa.get("estilo") or "apa"
    idioma = idioma or mapa.get("idioma") or "es"
    # fuentes: inventario (--fuentes) por encima de la copia guardada en el documento
    for k, r in (mapa.get("fuentes") or {}).items():
        FUENTES.setdefault(k, r)
        for al in r.get("alias") or []:
            FUENTES.setdefault(al, r)
    vivas = _citas_del_doc(doc, citas)
    numeros, orden = {}, []
    for _, items in vivas:
        for it in items:
            r = FUENTES.get(it["clave"])
            if r is not None and r["id"] not in numeros:
                numeros[r["id"]] = len(numeros) + 1
                orden.append(r)
    sufijos = {}
    if estilo in ("apa", "chicago"):
        grupos = {}
        for r in orden:
            grupos.setdefault((_norm(_etiqueta(r, estilo, idioma, "&")), r.get("anio") or ""), []).append(r)
        for g in grupos.values():
            if len(g) > 1:
                for i, r in enumerate(sorted(g, key=lambda x: _norm(x.get("titulo") or ""))):
                    sufijos[r["id"]] = "abcdefghijklmnopqrstuvwxyz"[i]
    faltan = []
    for start, items in vivas:
        trozos, f = _texto_cita(items, estilo, idioma, numeros, sufijos)
        faltan += f
        _reescribir(start, trozos)
    entradas = [(r["id"], _ref_segmentos(r, estilo, idioma, numeros[r["id"]], sufijos.get(r["id"], "")))
                for r in orden] if estilo != "zotero" else []
    _poner_lista(doc, entradas, idioma, estilo in ("ieee", "vancouver"))
    snap = {r["id"]: {k: r[k] for k in CAMPOS_SNAP if r.get(k)} for r in orden}
    guardar_mapa(doc, {"v": 1, "estilo": estilo, "idioma": idioma, "citas": citas, "fuentes": snap})
    faltan = list(dict.fromkeys(faltan))
    if avisar:
        print(f"citas por clave: {len(vivas)} en el documento, {len(orden)} fuentes citadas · estilo {ESTILOS[estilo]}")
        if faltan:
            print(f"  ✗ Claves sin fuente en el inventario: {', '.join(faltan)} → quedan [POR CONFIRMAR: fuente …] en rojo."
                  " Agrégalas al inventario (o corrige la clave) y vuelve a correr con --fuentes.")
        sin_pila = sum(1 for r in orden if estilo == "chicago" and any("org" not in a and not a["pila"] for a in _autores(r)))
        if sin_pila:
            print(f"  ! {sin_pila} referencias sin nombre de pila completo (Chicago lo pide): se usan iniciales."
                  " Exporta con nombres completos (AF de WoS, 'Author full names' de Scopus) o confírmalos.")
    return faltan


# --------------------------------------------------------------- comandos
def _modo(a):
    global TAM_TABLA, INTERLINEADO
    if getattr(a, "compacto", False):
        TAM_TABLA, INTERLINEADO = 10, 1.15
        FMT["tabla_tamano"], FMT["interlineado"] = 10, 1.15


PLANTILLA_REVISTA = None   # .docx/.dotx oficial de la revista, si el perfil la trae


def _documento_base():
    """Documento vacío: la plantilla oficial de la revista (con sus estilos) o uno nuevo."""
    if PLANTILLA_REVISTA and os.path.exists(PLANTILLA_REVISTA):
        doc = Document(PLANTILLA_REVISTA)
        body = doc.element.body
        for el in list(body.iterchildren()):
            if el.tag != qn("w:sectPr"):
                body.remove(el)
        return doc, True
    return Document(), False


def _construir_nuevo(a, out, margen):
    doc, de_plantilla = _documento_base()
    if de_plantilla:
        _sin_metadatos(doc)    # se respetan los estilos y márgenes de la plantilla oficial
    else:
        preparar_estilos(doc, margen)
    with open(a.md, encoding="utf-8") as f:
        bloques, refs = parsear_md(f.read())
    _citas_inicio(doc)
    escribir_bloques(doc, bloques, a.idioma, os.path.dirname(os.path.abspath(a.md)))
    if refs:
        fusionar_referencias(doc, refs, a.idioma)
    _procesar_citas(doc, IDIOMA_CITA or a.idioma)
    doc.save(out)
    print(f"ok: {out}")
    _resumen_corto(out)


def _plantilla_de(perfil, ruta_perfil):
    pl = (perfil or {}).get("plantilla") or ""
    if pl and not os.path.isabs(pl) and ruta_perfil:
        pl = os.path.join(os.path.dirname(os.path.abspath(ruta_perfil)), pl)
    if pl and not os.path.exists(pl):
        print(f"  ! El perfil nombra la plantilla «{pl}», que no está: se usa el formato del perfil.")
        return None
    return pl or None


def _preparar_citas(a, perfil, mapa=None):
    """Fuentes (--fuentes), estilo de cita (perfil > el guardado en el documento > APA 7) e idioma."""
    global ESTILO, IDIOMA_CITA
    FUENTES.clear()
    FUENTES.update(cargar_fuentes(getattr(a, "fuentes", None)))
    ESTILO = estilo_de(perfil.get("estilo_cita")) if perfil and perfil.get("estilo_cita") else None
    idioma_perfil = str((perfil or {}).get("idioma") or "").lower()[:2]
    if getattr(a, "idioma", None) is None:
        a.idioma = idioma_perfil if idioma_perfil in ("es", "en") else (mapa or {}).get("idioma") or "es"
    IDIOMA_CITA = a.idioma


def cmd_nuevo(a):
    global PLANTILLA_REVISTA
    perfil = leer_perfil(getattr(a, "perfil", None))
    _preparar_citas(a, perfil)
    aplicar_formato(perfil.get("formato") if perfil else None)
    PLANTILLA_REVISTA = _plantilla_de(perfil, getattr(a, "perfil", None))
    _modo(a)
    _construir_nuevo(a, a.out, a.margen)
    if perfil and not a.sin_respaldo:
        aplicar_formato(None)          # mismo texto, formato neutro de la ruta
        PLANTILLA_REVISTA = None
        _construir_nuevo(a, _ruta_respaldo(a.out), None)
        print("  (versión con formato de la revista + respaldo neutro por si hay rechazo)")


def cmd_agregar(a):
    perfil = leer_perfil(getattr(a, "perfil", None))
    _preparar_citas(a, perfil, leer_mapa(Document(a.base)))
    aplicar_formato(perfil.get("formato") if perfil else None)
    out = a.out or _siguiente_version(a.base)
    _agregar_en(a, a.base, out)
    if perfil and not a.sin_respaldo:
        base_r = _ruta_respaldo(a.base)
        if os.path.exists(base_r):
            aplicar_formato(None)
            _agregar_en(a, base_r, _ruta_respaldo(out))
            print("  (respaldo neutro actualizado)")
        else:
            print(f"  ! No existe {os.path.basename(base_r)}: el respaldo neutro no se actualizó. "
                  f"Créalo con `normalizar {os.path.basename(a.base)} --out {os.path.basename(base_r)}`.")


def _agregar_en(a, base, out):
    doc = Document(base)
    with open(a.md, encoding="utf-8") as f:
        bloques, refs = parsear_md(f.read())
    _citas_inicio(doc)
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    previos = [e for e in body.iterchildren() if e is not sectPr]
    n_previos = len(previos)
    escribir_bloques(doc, bloques, a.idioma, os.path.dirname(os.path.abspath(a.md)))
    nuevos = [e for e in body.iterchildren() if e.tag != qn("w:sectPr")][n_previos:]
    # ¿dónde colocarlos?
    destino = None
    if a.posicion == "inicio":
        # después del título del artículo (primer párrafo con texto)
        for e in body.iterchildren():
            if e.tag == qn("w:p") and "".join(t.text or "" for t in e.iter(qn("w:t"))).strip():
                destino = ("despues", e)
                break
    elif a.posicion == "antes-de":
        for e in body.iterchildren():
            if e.tag == qn("w:p") and _norm("".join(t.text or "" for t in e.iter(qn("w:t")))) == _norm(a.seccion):
                destino = ("antes", e)
                break
        if destino is None:
            sys.exit(f"No encontré la sección «{a.seccion}» en {base}.")
    else:  # antes de Referencias
        titulo_ref, _ = _bloque_refs(doc)
        if titulo_ref is not None:
            destino = ("antes", titulo_ref)
    if destino:
        modo, ancla = destino
        if modo == "antes":
            for e in nuevos:
                ancla.addprevious(e)
        else:
            for e in reversed(nuevos):
                ancla.addnext(e)
    if a.titulo:
        for p in doc.paragraphs:
            if p.text.strip():
                for r in p.runs[1:]:
                    r._element.getparent().remove(r._element)
                if p.runs:
                    p.runs[0].text = a.titulo
                break
    añadidas, total = fusionar_referencias(doc, refs, a.idioma) if refs else (0, None)
    _procesar_citas(doc, IDIOMA_CITA or a.idioma)
    _sin_metadatos(doc)
    if os.path.abspath(out) == os.path.abspath(base):
        sys.exit("No sobrescribo la versión anterior: usa otro nombre de salida.")
    doc.save(out)
    print(f"ok: {out}")
    if refs:
        print(f"referencias: {añadidas} nuevas, {total} en la lista")
    _resumen_corto(out)


SECCIONES_1 = {_n for _n in (
    "resumen", "abstract", "introduccion", "introduction", "revision de la literatura", "literature review",
    "marco teorico", "metodo", "metodos", "metodologia", "method", "methods", "materials and methods",
    "materiales y metodos", "resultados", "results", "discusion", "discussion", "conclusiones", "conclusion",
    "conclusions", "referencias", "references", "declaraciones", "declarations", "agradecimientos",
    "acknowledgments", "limitaciones", "anexos", "apendice", "resultados y discusion", "results and discussion")}


def _titulo_implicito(el, txt):
    """Párrafo corto, todo en negrita y sin punto final: título en un Word externo."""
    if not txt or len(txt.split()) > 14 or txt.endswith((".", ":", ";")) or RE_TABLA_CAP.match(txt) or RE_FIG_CAP.match(txt):
        return 0
    if txt.startswith("[") or RE_PEND.search(txt):   # pendientes en rojo y negrita no son títulos
        return 0
    ppr = el.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:   # viñetas y listas
        return 0
    runs = [r for r in el.iter(qn("w:r")) if "".join(t.text or "" for t in r.iter(qn("w:t"))).strip()]
    if not runs:
        return 0
    for r in runs:
        rpr = r.find(qn("w:rPr"))
        b = rpr.find(qn("w:b")) if rpr is not None else None
        if b is None or b.get(qn("w:val")) in ("0", "false"):
            return 0
    return 1 if _norm(txt) in SECCIONES_1 else 2


def _recorrer(doc):
    """Genera (tipo, texto, sección actual) en orden de documento."""
    seccion = ""
    body = doc.element.body
    primero = True
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            estilo = el.find(qn("w:pPr"))
            nombre = ""
            if estilo is not None and estilo.find(qn("w:pStyle")) is not None:
                nombre = estilo.find(qn("w:pStyle")).get(qn("w:val")) or ""
            txt = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
            implicito = 0 if (nombre.lower().startswith(("heading", "ttulo")) or (primero and txt)) else _titulo_implicito(el, txt)
            if txt:
                primero = False
            if nombre.lower().startswith("heading") or nombre.lower().startswith("ttulo") or implicito:
                nivel = str(implicito) if implicito else (re.sub(r"\D", "", nombre) or "1")
                if nivel == "1":
                    seccion = txt
                yield ("h" + nivel, txt, seccion, el)
            elif txt:
                yield ("p", txt, seccion, el)
        elif el.tag == qn("w:tbl"):
            yield ("tabla", "", seccion, el)


def _resumen_corto(ruta):
    doc = Document(ruta)
    secs, total = {}, 0
    for tipo, txt, sec, _ in _recorrer(doc):
        if tipo == "p" and not _es_titulo_ref(sec):
            n = _palabras(txt)
            secs[sec] = secs.get(sec, 0) + n
            total += n
    print("secciones: " + "; ".join(f"{k or '(inicio)'} {v}" for k, v in secs.items()))
    print(f"palabras sin referencias: {total}")


def cmd_esquema(a):
    doc = Document(a.docx)
    secs, tablas, figs, pend = [], [], [], []
    palabras = {}
    for tipo, txt, sec, _ in _recorrer(doc):
        if tipo.startswith("h"):
            secs.append(("  " * (int(tipo[1:]) - 1)) + txt)
        elif tipo == "p":
            if RE_TABLA_CAP.match(txt):
                tablas.append(txt)
            elif RE_FIG_CAP.match(txt):
                figs.append(txt)
            for m in RE_PEND.findall(txt):
                pend.append(f"{sec}: {m}")
            if not _es_titulo_ref(sec):
                palabras[sec] = palabras.get(sec, 0) + _palabras(txt)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for m in RE_PEND.findall(c.text):
                    pend.append(f"(tabla): {m}")
    _, refs = _bloque_refs(doc)
    for r in refs:
        for m in RE_PEND.findall("".join(x.text or "" for x in r.iter(qn("w:t")))):
            pend.append(f"(referencias): {m}")
    pend = list(dict.fromkeys(pend))
    primero = next((p.text.strip() for p in doc.paragraphs if p.text.strip()), "")
    print(f"ARCHIVO: {os.path.basename(a.docx)}")
    print(f"TÍTULO: {primero}")
    print("ESTRUCTURA:")
    for s in secs:
        sec = s.strip()
        extra = f"  ({palabras.get(sec, 0)} palabras)" if sec in palabras else ""
        print(f"  {s}{extra}")
    print(f"TABLAS: {', '.join(tablas) or 'ninguna'}")
    print(f"FIGURAS: {', '.join(figs) or 'ninguna'}")
    print(f"REFERENCIAS: {len(refs)}")
    print(f"PALABRAS (sin referencias): {sum(palabras.values())}")
    print(f"PENDIENTES ({len(pend)}):")
    for p in pend:
        print(f"  - {p}")
    if a.acronimos:
        acr = sorted(set(re.findall(r"\(([A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ0-9-]{1,9})\)", " ".join(p.text for p in doc.paragraphs))))
        print(f"ACRÓNIMOS DEFINIDOS: {', '.join(acr) or 'ninguno'}")


def cmd_texto(a):
    doc = Document(a.docx)
    dentro, nivel_sec = (not a.seccion), 0
    for tipo, txt, sec, el in _recorrer(doc):
        if a.seccion:
            if tipo.startswith("h"):
                niv = int(tipo[1:])
                if _norm(txt) == _norm(a.seccion):
                    dentro, nivel_sec = True, niv
                elif dentro and niv <= nivel_sec:
                    dentro = False
            if not dentro:
                continue
        if tipo.startswith("h"):
            print("#" * int(tipo[1:]) + " " + txt)
        elif tipo == "p":
            print(txt + "\n")
        elif tipo == "tabla":
            for fila in el.iter(qn("w:tr")):
                celdas = ["".join(t.text or "" for t in c.iter(qn("w:t"))) for c in fila.iter(qn("w:tc"))]
                print("| " + " | ".join(celdas) + " |")
            print()


RE_CITA = re.compile(
    r"([A-ZÁÉÍÓÚÑ][\wáéíóúñü'’-]+)(?:\s+(?:et al\.|y|and|&)\s*(?:[A-ZÁÉÍÓÚÑ][\wáéíóúñü'’-]+)?)?,?\s*\(?((?:19|20)\d{2}[a-z]?)\)?")


def cmd_verificar(a):
    doc = Document(a.docx)
    problemas, avisos = [], []
    # 1. formato base
    sec = doc.sections[0]
    _pf = (leer_perfil(getattr(a, "perfil", None)) or {}).get("formato", {})
    carta = str(_pf.get("papel", "")).lower() in ("carta", "letter")
    ancho, alto = (21.59, 27.94) if carta else (21, 29.7)
    if abs(sec.page_width.cm - ancho) > 0.2 or abs(sec.page_height.cm - alto) > 0.2:
        (problemas if _pf else avisos).append(f"El tamaño de página no es {'carta' if carta else 'A4'}.")
    st = doc.styles["Normal"].font
    if not getattr(a, "perfil", None) and st.name != FUENTE:
        avisos.append(f"La fuente base es {st.name}, no {FUENTE}.")
    for n in (1, 2):
        try:
            c = doc.styles[f"Heading {n}"].font.color.rgb
            if c is not None and c != NEGRO:
                problemas.append(f"El Título {n} no está en negro ({c}).")
        except KeyError:
            pass
    _fuente_esp = ((leer_perfil(getattr(a, "perfil", None)) or {}).get("formato", {}) or {}).get("fuente")
    if _fuente_esp:
        otras = {r.font.name for p in doc.paragraphs for r in p.runs if r.font.name and r.font.name != _fuente_esp}
        otras |= {r.font.name for t in doc.tables for f in t.rows for c in f.cells for q in c.paragraphs
                  for r in q.runs if r.font.name and r.font.name != _fuente_esp}
        if otras:
            problemas.append(f"Hay texto en otra fuente ({', '.join(sorted(otras))}); la revista pide {_fuente_esp}.")
    # 2. tablas y figuras
    vistos = {"Tabla": [], "Figura": []}
    citadas = {"Tabla": set(), "Figura": set()}
    cuerpo_texto = []
    for tipo, txt, seccion, el in _recorrer(doc):
        if tipo != "p":
            continue
        mt, mf = RE_TABLA_CAP.match(txt), RE_FIG_CAP.match(txt)
        if mt or mf:
            clase = "Tabla" if mt else "Figura"
            _g = (mt or mf).group(2)
            num = int(_g) if _g.isdigit() else _romano(_g)
            if num not in citadas[clase]:
                problemas.append(f"{clase} {num} aparece antes de ser citada en el texto.")
            vistos[clase].append(num)
            continue
        if _es_titulo_ref(seccion):
            continue
        cuerpo_texto.append(txt)
        for m in re.finditer(r"\b(Tablas?|Tables?|Figuras?|Figures?)\s+(\d+)(?:\s*(?:y|and|a|-|–)\s*(\d+))?", txt):
            clase = "Tabla" if m.group(1).lower().startswith("tab") else "Figura"
            ini = int(m.group(2))
            fin = int(m.group(3)) if m.group(3) else ini
            for k in range(ini, fin + 1):
                citadas[clase].add(k)
    for clase, nums in vistos.items():
        if nums and nums != list(range(1, len(nums) + 1)):
            problemas.append(f"Numeración de {clase.lower()}s no continua: {nums}.")
        for k in citadas[clase] - set(nums):
            avisos.append(f"El texto cita la {clase} {k}, que no existe en el documento.")
    # 3. bordes verticales en tablas
    for i, tbl in enumerate(doc.tables, 1):
        xml = tbl._tbl.xml
        if re.search(r'<w:(left|right|insideV)\s+w:val="(single|double|thick)"', xml):
            problemas.append(f"La tabla {i} tiene bordes verticales (APA 7 no los lleva).")
    # 4. referencias
    _, refs = _bloque_refs(doc)
    textos = ["".join(t.text or "" for t in r.iter(qn("w:t"))).strip() for r in refs]
    claves = [_clave_ref(x) for x in textos]
    dup = {k for k in claves if claves.count(k) > 1}
    if dup:
        problemas.append(f"Referencias duplicadas: {len(dup)}.")
    numerica = textos and all(re.match(r"^\[?\d+[\].]", x) for x in textos)
    if numerica:
        nums = [int(re.match(r"^\[?(\d+)", x).group(1)) for x in textos]
        if nums != list(range(1, len(nums) + 1)):
            avisos.append(f"Referencias numeradas fuera de orden o con saltos: {nums[:15]}")
    elif [_norm(x) for x in textos] != sorted(_norm(x) for x in textos):
        avisos.append("La lista de referencias no está en orden alfabético.")
    if textos:
        apellidos_ref = set()
        for x in textos:   # todas las palabras de la parte de autores (apellidos compuestos)
            autores = re.split(r"\(\s*(?:19|20)\d{2}|\b(?:19|20)\d{2}\b", x)[0]
            apellidos_ref |= {_norm(w) for w in re.findall(r"[A-Za-zÁÉÍÓÚáéíóúÑñÜü'’-]{3,}", autores)}
        sin_ref = set()
        for m in RE_CITA.finditer(" ".join(cuerpo_texto)):
            ap = _norm(m.group(1))
            if ap in {"tabla", "table", "figura", "figure", "nota", "en", "el", "la", "los", "las", "desde", "hasta",
                      "durante", "entre", "segun", "hacia", "año", "ano", "ods"}:
                continue
            if ap not in apellidos_ref and len(ap) > 2:
                sin_ref.add(f"{m.group(1)} {m.group(2)}")
        if sin_ref:
            avisos.append("Posibles citas sin referencia (revisar a mano): " + ", ".join(sorted(sin_ref)[:25]))
    # 4b. citas por clave
    mapa, linea_citas = leer_mapa(doc), ""
    crudas = sorted({m.group(0) for t in cuerpo_texto for m in RE_GRUPO.finditer(t)})
    if crudas:
        problemas.append("Claves de cita sin convertir en el texto (¿pegadas a mano?): " + ", ".join(crudas[:10])
                         + ". Escríbelas en el .md y pásalas por el motor.")
    if mapa.get("citas"):
        FUENTES.clear()
        FUENTES.update(cargar_fuentes(getattr(a, "fuentes", None)))
        inventario = dict(FUENTES)
        for k, r in (mapa.get("fuentes") or {}).items():
            FUENTES.setdefault(k, r)
        vivas = _citas_del_doc(doc, mapa["citas"])
        claves_c = list(dict.fromkeys(it["clave"] for _, items in vivas for it in items))
        sin_fuente = [k for k in claves_c if k not in FUENTES]
        if sin_fuente:
            problemas.append("Citas con clave que no está en el inventario: " + ", ".join(sin_fuente))
        ids_c = {FUENTES[k]["id"] for k in claves_c if k in FUENTES}
        ids_l = {(b.get(qn("w:name")) or "")[len(PREFIJO_REF):] for el in refs
                 for b in el.iter(qn("w:bookmarkStart")) if (b.get(qn("w:name")) or "").startswith(PREFIJO_REF)}
        if mapa.get("estilo") == "zotero":
            avisos.append("Versión con marcas {Autor, año} para el RTF Scan de Zotero: no es la versión para enviar.")
            ids_l = ids_c
        if ids_c - ids_l:
            problemas.append("Fuentes citadas que no están en la lista de referencias: " + ", ".join(sorted(ids_c - ids_l))
                             + ". Regenerar con `agregar` o `normalizar` y --fuentes.")
        if ids_l - ids_c:
            problemas.append("Referencias generadas que ya no se citan: " + ", ".join(sorted(ids_l - ids_c))
                             + ". Regenerar con `normalizar --fuentes`.")
        if getattr(a, "fuentes", None):
            nuevas = [k for k in claves_c if k in inventario and k not in (mapa.get("fuentes") or {})
                      and inventario[k]["id"] not in (mapa.get("fuentes") or {})]
            if nuevas:
                avisos.append("Claves que ya están en el inventario pero el Word no las resolvió: " + ", ".join(nuevas)
                              + ". Regenerar con `normalizar --fuentes`.")
        linea_citas = (f"  citas por clave: {len(vivas)} · fuentes citadas: {len(ids_c)} · estilo "
                       f"{ESTILOS.get(mapa.get('estilo'), mapa.get('estilo'))}")
    # 5. huecos de plantilla sin rellenar
    todo = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    huecos = sorted({h for x in todo for h in re.findall(r"\{[^{}]{1,60}\}|<[^<>]{2,60}>", x)
                     if not (mapa.get("estilo") == "zotero" and h.startswith("{"))})
    if huecos:
        problemas.append("Quedaron huecos de plantilla sin rellenar: " + ", ".join(huecos[:10]))
    # 6. pendientes y palabras
    pend = [m for t in todo for m in RE_PEND.findall(t)]   # cuerpo, viñetas, tablas y referencias
    total = sum(_palabras(t) for t in cuerpo_texto)
    if a.con_tablas:
        total += sum(_palabras(c.text) for t in doc.tables for r in t.rows for c in r.cells)
    perfil = leer_perfil(getattr(a, "perfil", None))
    if perfil:
        fm = perfil.get("formato", {})
        if fm.get("fuente") and st.name != fm["fuente"]:
            problemas.append(f"Fuente {st.name}; la revista pide {fm['fuente']}.")
        if fm.get("tamano") and st.size and abs(st.size.pt - float(fm["tamano"])) > 0.1:
            problemas.append(f"Tamaño {st.size.pt} pt; la revista pide {fm['tamano']}.")
        li = doc.styles["Normal"].paragraph_format.line_spacing
        if fm.get("interlineado") and li and abs(float(li) - float(fm["interlineado"])) > 0.05:
            problemas.append(f"Interlineado {li}; la revista pide {fm['interlineado']}.")
        if fm.get("numerar_lineas") and "lnNumType" not in doc.sections[0]._sectPr.xml:
            problemas.append("La revista pide numeración de líneas y el documento no la tiene.")
        a.limite = a.limite or perfil.get("limite_palabras")
        cuenta = _norm(perfil.get("limite_cuenta", ""))
        total_refs = sum(_palabras(x) for x in textos) if ("con referencias" in cuenta or "todo" in cuenta) else 0
        if ("tabla" in cuenta or "todo" in cuenta) and not a.con_tablas:
            total += sum(_palabras(c.text) for t in doc.tables for r in t.rows for c in r.cells)
            a.con_tablas = True
        if perfil.get("limite_paginas"):
            por_pag = 250 if float(fm.get("interlineado") or 1.5) >= 2 else 330
            est = (total + total_refs) / por_pag
            (avisos if est <= float(perfil["limite_paginas"]) else problemas).append(
                f"Páginas estimadas ≈ {est:.1f} (a {por_pag} palabras por página) frente a un límite de "
                f"{perfil['limite_paginas']}: confirmar en Word.")
        esperadas = [_norm(x) for x in perfil.get("estructura", [])]
        presentes = [_norm(t) for tp, t, _, _ in _recorrer(doc) if tp == "h1"]
        faltan = [x for x in perfil.get("estructura", []) if _norm(x) not in presentes]
        no_secciones = {"resumen", "abstract", "palabras clave", "keywords", "declaraciones", "declarations",
                        "agradecimientos", "acknowledgments", "anexos", "apendice", "appendix", "financiamiento",
                        "conflicto de intereses", "disponibilidad de datos", "contribucion de los autores"}
        extra = [t for t in presentes if esperadas and t not in esperadas and t not in REF_TITULOS
                 and t not in no_secciones]
        if faltan:
            avisos.append("Secciones de la revista que aún no están: " + ", ".join(faltan))
        if extra:
            avisos.append("Secciones que la revista no nombra (revisar el título): " + ", ".join(extra))
        orden = [x for x in presentes if x in esperadas]
        if orden != [x for x in esperadas if x in orden]:
            problemas.append("Las secciones no siguen el orden que pide la revista.")
    else:
        total_refs = 0
    print(f"VERIFICACIÓN de {os.path.basename(a.docx)}" + (f" · perfil {perfil.get('revista', '')}" if perfil else ""))
    print(f"  palabras ({'con' if a.con_tablas else 'sin'} tablas, sin referencias): {total}"
          + (f"  · con referencias: {total + total_refs}" if total_refs else "")
          + (f"  · límite {a.limite} (se compara {total + total_refs})" if a.limite else ""))
    if a.limite and total + total_refs > a.limite:
        problemas.append(f"Se pasa del límite: {total + total_refs} > {a.limite} palabras"
                         + (" (contando referencias)" if total_refs else "") + ".")
    if not textos and re.search(r"\((?:[^()]*?\b(?:19|20)\d{2}[a-z]?)\)|\[\d+\]", " ".join(cuerpo_texto)):
        problemas.append("Hay citas en el texto pero no encontré la lista de referencias (¿título de la sección distinto?).")
    print(f"  tablas: {vistos['Tabla'] or '-'}  figuras: {vistos['Figura'] or '-'}  referencias: {len(textos)}")
    print(f"  pendientes [POR CONFIRMAR]: {len(pend)}")
    if linea_citas:
        print(linea_citas)
    for p in problemas:
        print(f"  ✗ {p}")
    for w in avisos:
        print(f"  ! {w}")
    if not problemas:
        print("  ✓ Sin errores de formato ni de numeración.")
    sys.exit(1 if problemas else 0)


def cmd_normalizar(a):
    """Pasa un Word externo al formato de la ruta sin tocar su texto."""
    global ESTILO
    perfil = leer_perfil(getattr(a, "perfil", None))
    aplicar_formato(perfil.get("formato") if perfil else None)
    doc = Document(a.docx)
    preparar_estilos(doc)
    n1 = n2 = 0
    primero = next((p for p in doc.paragraphs if p.text.strip()), None)
    for p in doc.paragraphs:
        txt = p.text.strip()
        if p.style.name.lower().startswith("heading"):
            n1 += p.style.name.endswith("1")
            n2 += not p.style.name.endswith("1")
            for r in p.runs:
                r.font.name = FUENTE
                r.font.size = None          # hereda el tamaño del estilo de título
            continue
        nivel = 0 if p is primero else _titulo_implicito(p._p, txt)   # el título del artículo no es sección
        if nivel:
            p.style = doc.styles[f"Heading {nivel}"]
            n1 += nivel == 1
            n2 += nivel == 2
        if not nivel and p.text.strip() and not RE_TABLA_CAP.match(txt) and not RE_FIG_CAP.match(txt):
            p.paragraph_format.line_spacing = INTERLINEADO
        for r in p.runs:
            r.font.name = FUENTE
            if r.font.color is not None and r.font.color.rgb not in (None, NEGRO, ROJO):
                r.font.color.rgb = NEGRO
    for t in doc.tables:
        for fila in t.rows:
            for c in fila.cells:
                for q in c.paragraphs:
                    for r in q.runs:
                        r.font.name = FUENTE
                        r.font.size = Pt(TAM_TABLA)
    mapa = leer_mapa(doc)
    if mapa.get("citas"):
        # reenvío: las citas conservan sus claves → se reescriben en el estilo de la revista nueva
        a.idioma = None
        _preparar_citas(a, perfil, mapa)
        if getattr(a, "estilo", None):
            ESTILO = estilo_de(a.estilo)
        _citas_inicio(doc)
        antes = mapa.get("estilo")
        _procesar_citas(doc, a.idioma)
        if ESTILO and antes and ESTILO != antes:
            print(f"  citas y referencias reconvertidas: {ESTILOS.get(antes, antes)} → {ESTILOS[ESTILO]}")
    elif perfil and perfil.get("estilo_cita"):
        print("  ! Este Word no conserva las claves de las citas: si la revista nueva usa otro estilo, las citas "
              "y la lista se convierten a mano (estilos_cita.md).")
    out = a.out or _siguiente_version(a.docx)
    doc.save(out)
    print(f"ok: {out} · {n1} títulos de sección y {n2} subtítulos reconocidos")
    print("Revisar con `esquema`: si un título no se reconoció, ponerlo en negrita y sin punto final.")


PLANTILLA_PERFIL = {
    "revista": "Nombre exacto de la revista",
    "guia_url": "https://... (guía de autores)",
    "consultado": "AAAA-MM-DD",
    "idioma": "es",
    "estilo_cita": "APA 7 | Chicago autor-fecha | Vancouver | IEEE | Harvard | otro (copiar el nombre de la guía)",
    "estructura": ["Introducción", "Método", "Resultados", "Discusión", "Conclusiones"],
    "estructura_en": [],
    "limite_palabras": None,
    "limite_cuenta": "sin referencias | con referencias | con tablas y referencias | todo el manuscrito",
    "limite_paginas": None,
    "resumen": {"tipo": "no estructurado | estructurado", "max_palabras": 250, "idiomas": ["es", "en"]},
    "palabras_clave": "3 a 5",
    "titulo_max_palabras": 20,
    "tablas": "en el texto | al final del manuscrito",
    "figuras": "en el texto | al final | archivos aparte",
    "anonimizar": True,
    "plantilla": "ruta al .docx/.dotx oficial de la revista si es obligatorio (o vacío)",
    "formato": {"fuente": "Times New Roman", "tamano": 12, "interlineado": 1.5, "margen_cm": 2.54,
                "justificado": True, "sangria_cm": 1.27, "tabla_tamano": 10, "numerar_lineas": False,
                "numerar_paginas": True, "titulo1_centrado": False, "titulo1_mayusculas": False,
                "titulo2_cursiva": True, "papel": "A4 | carta", "numero_pagina": "pie | encabezado"},
    "por_confirmar": ["campos que la guía no dice"],
}


def cmd_perfil(a):
    if a.accion == "plantilla":
        print(json.dumps(PLANTILLA_PERFIL, ensure_ascii=False, indent=2))
        return
    p = leer_perfil(a.archivo)
    errores = [k for k in ("revista", "estilo_cita", "estructura", "formato") if not p.get(k)]
    sin_rellenar = []
    def _mirar(v, ruta=""):
        if isinstance(v, dict):
            for k, x in v.items():
                _mirar(x, f"{ruta}.{k}" if ruta else k)
        elif isinstance(v, list):
            for x in v:
                _mirar(x, ruta)
        elif isinstance(v, str) and ruta not in ("por_confirmar", "plantilla", "otros") and (
                " | " in v or v.startswith(("Nombre exacto", "https://...", "AAAA", "ruta al ")) or "campos que la guía" in v):
            sin_rellenar.append(ruta)
    _mirar(p)
    for k in sorted(set(sin_rellenar)):
        errores.append(f"{k} (sin rellenar: quedó el texto de la plantilla)")
    if not p.get("limite_palabras"):
        print("  ! límite de palabras vacío: ponerlo, o anotarlo en por_confirmar")
    raros = [k for k in p.get("formato", {}) if k not in NEUTRO]
    if raros:
        print("  ! campos de formato que el motor no usa:", ", ".join(raros))
    for k in errores:
        print(f"  ✗ {k}" if "sin rellenar" in k else f"  ✗ falta «{k}»")
    pc = p.get("por_confirmar") or []
    print(f"Perfil de {p.get('revista', '?')}: cita {p.get('estilo_cita', '?')}, "
          f"{len(p.get('estructura', []))} secciones, límite {p.get('limite_palabras', '?')} palabras, "
          f"{len(pc)} campos por confirmar")
    sys.exit(1 if errores else 0)


def _reemplazar_en(doc, buscar, por):
    """Reemplaza texto en párrafos y celdas. Si cabe en un run, conserva su formato;
    si cruza runs, une los runs del párrafo. Un pendiente rojo pasa a texto normal."""
    n = 0
    parrafos = list(doc.paragraphs) + [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    for p in parrafos:
        if buscar not in p.text:
            continue
        hecho = False
        for r in p.runs:
            if buscar in r.text:
                era_pendiente = bool(RE_PEND.fullmatch(buscar.strip())) or r.text.strip() == buscar.strip() and RE_PEND.search(buscar)
                r.text = r.text.replace(buscar, por)
                if era_pendiente:
                    r.font.color.rgb = NEGRO
                    r.bold = None
                n += 1
                hecho = True
        if not hecho and p.runs:
            # cruza runs: se reescribe el párrafo; los [POR CONFIRMAR] que queden siguen en rojo
            texto = p.text.replace(buscar, por)
            tam = p.runs[0].font.size.pt if p.runs[0].font.size else None
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            _inline(p, texto, size=tam)
            n += 1
    return n


def cmd_reemplazar(a):
    import json as _j
    pares = [(a.buscar, a.por)] if a.buscar else [(d["buscar"], d["por"]) for d in _j.load(open(a.cambios, encoding="utf-8"))]
    destinos = [a.docx] + ([_ruta_respaldo(a.docx)] if os.path.exists(_ruta_respaldo(a.docx)) and not a.solo_este else [])
    for ruta in destinos:
        doc = Document(ruta)
        for b, r in pares:
            k = _reemplazar_en(doc, b, r)
            print(f"{os.path.basename(ruta)}: «{b[:50]}» → {k} reemplazo(s)" + ("  ✗ NO ENCONTRADO" if not k else ""))
        _sin_metadatos(doc)
        doc.save(ruta)


def cmd_contar(a):
    if a.archivo.lower().endswith(".docx"):
        doc = Document(a.archivo)
        secs = {}
        for tipo, txt, sec, _ in _recorrer(doc):
            if tipo == "p" and not _es_titulo_ref(sec):
                secs[sec] = secs.get(sec, 0) + _palabras(txt)
    else:
        with open(a.archivo, encoding="utf-8") as f:
            bloques, _ = parsear_md(f.read())
        secs, actual = {}, ""
        for b in bloques:
            if b[0] == "h" and b[1] == 1:
                actual = b[2]
            elif b[0] in ("p", "sin", "centro", "titulo", "vineta", "nota"):
                secs[actual] = secs.get(actual, 0) + _palabras(b[1])
            elif b[0] == "tabla":
                secs[actual] = secs.get(actual, 0) + _palabras(b[3]) + sum(_palabras(" ".join(f)) for f in b[4])
    for k, v in secs.items():
        print(f"{k or '(inicio)'}: {v}")
    print(f"TOTAL: {sum(secs.values())}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("nuevo"); p.add_argument("--md", required=True); p.add_argument("--out", required=True)
    p.add_argument("--idioma", default=None, choices=["es", "en"])
    p.add_argument("--fuentes", help="inventario.json: resuelve las citas por clave [F012] y genera las referencias"); p.add_argument("--margen", type=float, default=None)
    p.add_argument("--compacto", action="store_true", help="interlineado 1.15 y tablas a 10 pt (Ficha, carta, documentos de trabajo)")
    p.add_argument("--perfil", help="REVISTA_[clave].json: formato de la revista + respaldo neutro")
    p.add_argument("--sin-respaldo", action="store_true", help="con --perfil, no generar el _RESPALDO")
    p.set_defaults(f=cmd_nuevo)
    p = sub.add_parser("agregar"); p.add_argument("--base", required=True); p.add_argument("--md", required=True)
    p.add_argument("--out"); p.add_argument("--idioma", default=None, choices=["es", "en"])
    p.add_argument("--fuentes", help="inventario.json: resuelve las citas por clave y regenera las referencias")
    p.add_argument("--posicion", default="antes-referencias", choices=["antes-referencias", "inicio", "antes-de"])
    p.add_argument("--seccion", help="con --posicion antes-de: título exacto de la sección ancla")
    p.add_argument("--titulo", help="reemplaza el título del artículo (Fase 7)")
    p.add_argument("--perfil", help="REVISTA_[clave].json: formato de la revista; actualiza también el _RESPALDO")
    p.add_argument("--sin-respaldo", action="store_true")
    p.set_defaults(f=cmd_agregar)
    p = sub.add_parser("esquema"); p.add_argument("docx"); p.add_argument("--acronimos", action="store_true")
    p.set_defaults(f=cmd_esquema)
    p = sub.add_parser("texto"); p.add_argument("docx"); p.add_argument("--seccion"); p.set_defaults(f=cmd_texto)
    p = sub.add_parser("verificar"); p.add_argument("docx"); p.add_argument("--limite", type=int)
    p.add_argument("--con-tablas", action="store_true", help="cuenta también el texto de las tablas (Ficha)")
    p.add_argument("--perfil", help="compara formato, estructura y límite con REVISTA_[clave].json")
    p.add_argument("--fuentes", help="inventario.json: comprueba que cada clave citada exista")
    p.set_defaults(f=cmd_verificar)
    p = sub.add_parser("contar"); p.add_argument("archivo"); p.set_defaults(f=cmd_contar)
    p = sub.add_parser("normalizar", help="Word externo → formato de la ruta (estilos de título, fuente)")
    p.add_argument("docx"); p.add_argument("--out"); p.add_argument("--perfil", help="aplica el formato de esta revista")
    p.add_argument("--fuentes", help="inventario.json (si no, usa la copia de las fuentes guardada en el Word)")
    p.add_argument("--estilo", choices=["apa", "chicago", "vancouver", "ieee", "zotero"],
                   help="fuerza el estilo de las citas por clave; 'zotero' deja marcas {Autor, año} para RTF Scan")
    p.set_defaults(f=cmd_normalizar)
    p = sub.add_parser("reemplazar", help="cambia un texto (p. ej. un [POR CONFIRMAR] resuelto) en el .docx y en su _RESPALDO")
    p.add_argument("docx"); p.add_argument("--buscar"); p.add_argument("--por", default="")
    p.add_argument("--cambios", help='JSON [{"buscar": "...", "por": "..."}]')
    p.add_argument("--solo-este", action="store_true", help="no tocar el _RESPALDO")
    p.set_defaults(f=cmd_reemplazar)
    p = sub.add_parser("perfil", help="plantilla del perfil de la revista, o validarlo")
    p.add_argument("accion", choices=["plantilla", "validar"]); p.add_argument("archivo", nargs="?")
    p.set_defaults(f=cmd_perfil)
    a = ap.parse_args()
    a.f(a)


if __name__ == "__main__":
    main()
