#!/usr/bin/env python3
"""Matriz de Estrategia (Fase 1) en Excel, siempre con la misma forma.

Uso:
  python3 matriz_excel.py --json matriz.json [--out MATRIZ_clave.xlsx]
  python3 matriz_excel.py --ejemplo cuantitativo|secundarios|cualitativo|escala|mixto|revision > matriz.json

Solo necesita openpyxl (pip install openpyxl). Sin rutas fijas: lee y escribe
donde se le indique. Si el archivo de salida ya existe, guarda la versión
siguiente (_v2, _v3...) y no pisa la anterior.

Hojas fijas, en este orden:
  1 Diseño y muestra   (en revisiones: Diseño de la revisión)
  2 Objetivos          (es también el chequeo de coherencia)
  3 Operacionalización (columnas según el tipo; en revisiones: Variables de extracción)
  4 Papers de referencia
  5 Pendientes         (se completa sola con todo [POR CONFIRMAR] de las hojas 1 a 4)

Estructura del JSON: la muestra `--ejemplo <tipo>`, completa y con el formato
exacto de cada clave (también `pendientes`: lista de objetos
{"pendiente", "quien", "como"}; se aceptan cadenas sueltas). Se reemplaza cada
"<...>". En revisiones las claves de `diseno` son propias (bases, periodo,
ecuacion, registro, criterios_inclusion, criterios_exclusion, calidad,
analisis...), la hoja 3 va en `variables_extraccion` y la ECUACIÓN DE BÚSQUEDA
es obligatoria para cada base de `bases` (la fija la Fase 1).
Datos secundarios (tipo `secundarios`): `diseno` lleva claves explícitas
(fuente_datos, acceso, periodo, ponderador, variables_base, diccionario,
licencia...) y la hoja 3 cambia el instrumento por la variable de la base.
Objetivos específicos: 2 a 4; 5 se acepta con aviso; más de 5, error.
Sale con código 1 (y sin escribir) si falta algo obligatorio o hay celdas vacías.
"""
import argparse
import json
import os
import re
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")

TIPOS = ("cuantitativo", "secundarios", "cualitativo", "escala", "mixto", "revision")
PC = "[POR CONFIRMAR]"
RE_PC = re.compile(r"\[POR CONFIRMAR[^\]]*\]")

CAMPOS_DISENO = {
    "cuantitativo": ["Enfoque", "Diseño específico", "Alcance", "Temporalidad", "Población",
                     "Muestra prevista", "Tipo de muestreo", "Criterios de inclusión",
                     "Criterios de exclusión", "Justificación del tamaño", "Término unificador",
                     "Aprobación ética"],
    "escala": ["Enfoque", "Diseño específico", "Escala de origen (autor, año)", "Población",
               "Muestra prevista", "Tipo de muestreo", "Criterios de inclusión",
               "Criterios de exclusión", "Justificación del tamaño", "Traducción y juicio de expertos",
               "Permiso de uso", "Aprobación ética"],
    "cualitativo": ["Enfoque", "Diseño específico", "Participantes", "Muestreo",
                    "Criterio de cierre (saturación)", "Criterios de inclusión",
                    "Criterios de exclusión", "Técnica de recolección", "Análisis previsto",
                    "Consentimiento y grabación", "Aprobación ética"],
    "mixto": ["Enfoque", "Diseño específico", "Orden e integración de las ramas", "Temporalidad",
              "Población", "Muestra cuantitativa", "Participantes cualitativos",
              "Tipo de muestreo", "Criterios de inclusión", "Criterios de exclusión",
              "Justificación del tamaño", "Término unificador", "Aprobación ética"],
}

# Revisiones: claves explícitas en el JSON → rótulo en la hoja. `ecuacion` va
# aparte: un objeto {base: ecuación literal}, obligatorio para cada base.
CAMPOS_REVISION = [
    ("tipo_revision", "Tipo de revisión"), ("guia", "Guía de reporte"),
    ("pregunta", "Pregunta de revisión (PICO, PEO o PCC)"), ("bases", "Bases de datos"),
    ("periodo", "Periodo"), ("idiomas", "Idiomas"), ("tipo_documento", "Tipo de documento"),
    ("fecha_busqueda_bases", "Fecha de la búsqueda en bases (hecha o prevista)"),
    ("registro", "Registro de protocolo"), ("criterios_inclusion", "Criterios de inclusión"),
    ("criterios_exclusion", "Criterios de exclusión"), ("revisores", "Revisores del cribado"),
    ("calidad", "Evaluación de calidad o riesgo de sesgo"), ("analisis", "Análisis previsto"),
]
CAMPOS_DISENO["revision"] = [r for _, r in CAMPOS_REVISION]

# Datos secundarios (encuestas nacionales, bases públicas o institucionales,
# panel, series): claves explícitas → rótulo. Sin instrumento propio.
CAMPOS_SECUNDARIOS = [
    ("enfoque", "Enfoque"), ("diseno_especifico", "Diseño específico"),
    ("fuente_datos", "Fuente de datos (base y propietario)"),
    ("acceso", "Acceso (enlace, fecha de descarga, versión)"),
    ("periodo", "Año, ola o periodo"), ("unidad_analisis", "Unidad de análisis"),
    ("poblacion", "Población y muestra de la base"),
    ("diseno_muestral", "Diseño muestral (estratos, conglomerados)"),
    ("ponderador", "Factor de expansión o ponderador"),
    ("variables_base", "Módulos y variables de la base"),
    ("diccionario", "Diccionario de variables"),
    ("filtros", "Filtros y tamaño final"), ("perdidos", "Tratamiento de perdidos"),
    ("licencia", "Licencia de uso"), ("termino_unificador", "Término unificador"),
    ("etica", "Ética (exención o aprobación)"),
]
CAMPOS_DISENO["secundarios"] = [r for _, r in CAMPOS_SECUNDARIOS]
CAMPOS_CLAVE = {"revision": CAMPOS_REVISION, "secundarios": CAMPOS_SECUNDARIOS}

# Hoja 2: nombres de las columnas "unidad" y "medio" según el tipo
COLS_OBJ = {
    "cuantitativo": ("Variables", "Instrumento", "Hipótesis"),
    "secundarios": ("Variables", "Fuente y variable de la base", "Hipótesis"),
    "escala": ("Escala y dimensiones", "Instrumento (versión original)", "Modelo de medida esperado"),
    "cualitativo": ("Categorías", "Técnica e instrumento", "Hipótesis"),
    "mixto": ("Variables o categorías", "Instrumento o técnica", "Hipótesis"),
    "revision": ("Criterio de elegibilidad", "Variable de extracción", "Hipótesis"),
}

COLS_OPER = {
    "cuantitativo": ["Variable (acrónimo)", "Papel", "Definición conceptual (fuente)",
                     "Definición operacional", "Dimensiones", "Indicadores", "Escala de medición",
                     "Instrumento propuesto (autor, año, ítems)", "Estado"],
    "secundarios": ["Variable (acrónimo)", "Papel", "Definición conceptual (fuente)",
                    "Variable(s) de la base (código)", "Módulo, año u ola",
                    "Construcción o recodificación", "Escala de medición", "Fuente de datos", "Estado"],
    "escala": ["Escala y versión original (autor, año)", "Dimensión", "N.º de ítems",
               "Formato de respuesta", "Idioma original", "Propiedad a examinar", "Estado"],
    "cualitativo": ["Categoría", "Subcategorías", "Indicadores observables", "Técnica",
                    "Instrumento", "Origen (previa o emergente)", "Estado"],
    "mixto": ["Rama (cuanti o cuali)", "Variable o categoría", "Papel", "Definición conceptual (fuente)",
              "Definición operacional o subcategorías", "Indicadores", "Escala o técnica",
              "Instrumento propuesto (autor, año)", "Estado"],
    "revision": ["Variable de extracción", "Qué se registra", "Pregunta que responde", "Estado"],
}

COLS_PAPERS = {
    "default": [("autor_anio", "Autor y año"), ("titulo", "Título"), ("revista", "Revista"),
                ("diseno", "Diseño"), ("instrumento", "Instrumento o base de datos"), ("muestra", "Muestra"),
                ("analisis", "Análisis"), ("enlace", "Enlace")],
    "revision": [("autor_anio", "Autor y año"), ("titulo", "Título"), ("revista", "Revista"),
                 ("tipo_revision", "Tipo de revisión"), ("bases_periodo", "Bases y periodo"),
                 ("n_estudios", "N.º de estudios"), ("analisis", "Análisis"), ("enlace", "Enlace")],
}

NEGRITA = Font(bold=True)
GRIS = PatternFill("solid", fgColor="D9D9D9")
AJUSTE = Alignment(wrap_text=True, vertical="top")


def _ancho(texto):
    return max(14, min(60, int(len(texto) * 1.1) + 4))


def _hoja(wb, titulo, encabezados, filas, nota=None, primera=False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo[:31]
    ws.append(encabezados)
    for c in ws[1]:
        c.font, c.fill, c.alignment = NEGRITA, GRIS, AJUSTE
    for f in filas:
        ws.append(f)
    anchos = [_ancho(h) for h in encabezados]
    for f in filas:
        for j, v in enumerate(f):
            anchos[j] = max(anchos[j], min(60, len(str(v)) // 3 + 12))
    for j, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = AJUSTE
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{max(1, len(filas) + 1)}"
    if nota:
        ws.cell(row=len(filas) + 3, column=1, value=f"Nota. {nota}").font = Font(italic=True)
    return ws


def _vacias(nombre, filas, errores):
    for i, f in enumerate(filas, 2):
        for j, v in enumerate(f, 1):
            if re.fullmatch(r"<.*>", str(v).strip()):
                errores.append(f"{nombre}: fila {i}, columna {j} quedó sin rellenar ({v}).")
            elif str(v).strip() == "":
                errores.append(f"{nombre}: celda vacía en fila {i}, columna {j} "
                               "(poner dato, [POR CONFIRMAR: ...] o 'no aplica a este diseño').")


def _sin_rellenar(x, ruta, errores):
    """Cualquier valor "<...>" que quedó del esqueleto de --ejemplo."""
    if isinstance(x, dict):
        for k, v in x.items():
            _sin_rellenar(v, f"{ruta}.{k}" if ruta else k, errores)
    elif isinstance(x, list):
        for i, v in enumerate(x):
            _sin_rellenar(v, f"{ruta}[{i}]", errores)
    elif isinstance(x, str) and re.fullmatch(r"<.*>", x.strip()):
        errores.append(f"{ruta}: quedó sin rellenar ({x}).")


def _diseno_claves(tipo, dis, errores, avisos):
    """Hoja 1 con claves explícitas (revisión: más una ecuación por base;
    secundarios). Acepta también el rótulo como clave."""
    campos = CAMPOS_CLAVE[tipo]
    por_rotulo = {r: k for k, r in campos}
    dis = {por_rotulo.get(k, k): v for k, v in dis.items()}
    validas = [k for k, _ in campos] + (["ecuacion"] if tipo == "revision" else [])
    for k in dis:
        if k not in validas:
            errores.append(f"Diseño: clave desconocida '{k}'. Claves válidas: {validas}")
    filas = []
    for k, rotulo in campos:
        v = str(dis.get(k, "")).strip()
        if not v:
            avisos.append(f"Diseño: '{k}' sin dato; se anota {PC}.")
            v = f"[POR CONFIRMAR: {rotulo.lower()}]"
        filas.append([rotulo, v])
        if tipo == "revision" and k == "bases":
            filas.extend(_ecuaciones(dis.get("ecuacion"), v, errores, avisos))
    return filas


def _ecuaciones(ec, bases, errores, avisos):
    nombres = [b for b in re.split(r"\s*[;,]\s*|\s+y\s+", bases) if b and not RE_PC.fullmatch(b)]
    if isinstance(ec, str) and len(nombres) == 1:
        ec = {nombres[0]: ec}
    if not isinstance(ec, dict) or not ec:
        errores.append("Diseño: falta 'ecuacion', un objeto {\"Base\": \"ecuación literal\"} con una "
                       "entrada por cada base. Es obligatoria: la Fase 4 la reporta tal cual.")
        return []
    norm = {k.strip().lower(): k for k in ec}
    for b in nombres:
        if b.lower() not in norm:
            errores.append(f"Diseño: la base '{b}' no tiene ecuación en 'ecuacion'.")
    filas = []
    for base, txt in ec.items():
        txt = str(txt).strip()
        if not txt:
            errores.append(f"Diseño: la ecuación de '{base}' está vacía.")
        elif RE_PC.search(txt):
            avisos.append(f"Diseño: ecuación de '{base}' pendiente; pruébala en la base antes de la Fase 2.")
        filas.append([f"Ecuación de búsqueda · {base}", txt])
    return filas


def construir(d):
    errores, avisos = [], []
    tipo = d.get("tipo")
    if tipo not in TIPOS:
        sys.exit(f"✗ 'tipo' debe ser uno de {TIPOS}.")
    if not d.get("clave"):
        errores.append("Falta 'clave' (la palabra del tema fijada en la Fase 0).")

    # Hoja 1 · Diseño
    dis = d.get("diseno", {})
    f_dis = _diseno_claves(tipo, dis, errores, avisos) if tipo in CAMPOS_CLAVE else []
    if tipo not in CAMPOS_CLAVE:
        for k in dis:
            if k not in CAMPOS_DISENO[tipo]:
                errores.append(f"Diseño: campo desconocido '{k}'. Campos válidos: {CAMPOS_DISENO[tipo]}")
        for campo in CAMPOS_DISENO[tipo]:
            v = str(dis.get(campo, "")).strip()
            if not v:
                avisos.append(f"Diseño: '{campo}' sin dato; se anota {PC}.")
                v = f"[POR CONFIRMAR: {campo.lower()}]"
            f_dis.append([campo, v])

    # Hoja 2 · Objetivos
    unidad, medio, hip = COLS_OBJ[tipo]
    enc_obj = ["Nivel", "Objetivo", "Pregunta", hip, unidad, medio, "Análisis previsto"]
    objs = d.get("objetivos", [])
    if not objs or str(objs[0].get("nivel", "")).lower() != "general":
        errores.append("Objetivos: la primera fila debe ser el objetivo 'General'.")
    esp = [o for o in objs if str(o.get("nivel", "")).lower() != "general"]
    if len(esp) > 5:
        errores.append(f"Objetivos: hay {len(esp)} específicos; el tope es 5. Agrupa (un específico "
                       "puede reunir varias hipótesis o, en una bibliométrica, una familia de indicadores).")
    elif len(esp) == 5:
        avisos.append("Objetivos: 5 específicos; se acepta, pero lo habitual son 2 a 4. "
                      "Comprueba que caben en el límite de palabras de la revista.")
    elif len(esp) < 2:
        avisos.append(f"Objetivos: hay {len(esp)} específico(s); lo habitual son 2 a 4.")
    f_obj = [[str(o.get(k, "")).strip() for k in
              ("nivel", "objetivo", "pregunta", "hipotesis", "unidad", "medio", "analisis")] for o in objs]
    _vacias("Objetivos", f_obj, errores)

    # Hoja 3 · Operacionalización
    cols = COLS_OPER[tipo]
    filas_op = d.get("variables_extraccion") if tipo == "revision" else None
    f_op = [[str(x).strip() for x in fila] for fila in (filas_op or d.get("operacionalizacion", []))]
    if not f_op:
        errores.append("Operacionalización: no hay filas.")
    for i, fila in enumerate(f_op, 2):
        if len(fila) != len(cols):
            errores.append(f"Operacionalización fila {i}: {len(fila)} celdas, se esperan {len(cols)}: {cols}")
    _vacias("Operacionalización", f_op, errores)

    # Hoja 4 · Papers
    cp = COLS_PAPERS["revision" if tipo == "revision" else "default"]
    viejas = {"tipo_revision": "diseno", "bases_periodo": "instrumento", "n_estudios": "muestra"}
    f_pap = [[(str(p.get(k, p.get(viejas.get(k, ""), ""))).strip() or "no consta en el resultado")
              for k, _ in cp] for p in d.get("papers", [])]
    _vacias("Papers", f_pap, errores)
    nota_pap = (f"Búsqueda web del {d.get('fecha_busqueda') or PC}. Sondeo de orientación "
                "metodológica; no sustituye la búsqueda en bases de la Fase 2.")

    _sin_rellenar(d, "", errores)
    if errores:
        for e in errores:
            print(f"✗ {e}")
        sys.exit(1)

    # Hoja 5 · Pendientes (los declarados + los [POR CONFIRMAR] de las hojas 1 a 4)
    f_pen = [[p, "Investigador", "—"] if isinstance(p, str) else
             [p.get("pendiente", ""), p.get("quien", "Investigador"), p.get("como", "")]
             for p in d.get("pendientes", [])]
    ya = " ".join(x[0] for x in f_pen).lower()
    hojas = [("Diseño", f_dis), ("Objetivos", f_obj), ("Operacionalización", f_op), ("Papers", f_pap)]
    for nombre, filas in hojas:
        for i, fila in enumerate(filas, 2):
            for v in fila:
                for m in RE_PC.findall(str(v)):
                    txt = m[1:-1].replace("POR CONFIRMAR", "").strip(": ").strip() or "dato"
                    if txt.lower() not in ya:
                        f_pen.append([txt, "Investigador", f"Hoja {nombre}, fila {i}"])
                        ya += " " + txt.lower()
    if not f_pen:
        f_pen = [["No quedaron pendientes", "—", "—"]]

    wb = Workbook()
    _hoja(wb, "Diseño de la revisión" if tipo == "revision" else "Diseño y muestra",
          ["Campo", "Contenido"], f_dis, primera=True)
    _hoja(wb, "Objetivos", enc_obj, f_obj, nota=d.get("nota_objetivos") or None)
    _hoja(wb, "Variables de extracción" if tipo == "revision" else "Operacionalización", cols, f_op)
    _hoja(wb, "Papers de referencia", [h for _, h in cp], f_pap, nota=nota_pap)
    _hoja(wb, "Pendientes", ["Pendiente", "Quién lo resuelve", "Cómo"], f_pen)
    return wb, avisos, len(f_pen) if f_pen[0][0] != "No quedaron pendientes" else 0


def _salida(ruta):
    if not os.path.exists(ruta):
        return ruta
    base, ext = os.path.splitext(ruta)
    base = re.sub(r"_v\d+$", "", base)
    n = 2
    while os.path.exists(f"{base}_v{n}{ext}"):
        n += 1
    return f"{base}_v{n}{ext}"


def ejemplo(tipo):
    """Esqueleto completo y válido como JSON: se reemplaza cada "<...>"."""
    sin_hip = tipo in ("cualitativo", "revision")
    obj = lambda n: {"nivel": n, "objetivo": "<objetivo>", "pregunta": "<pregunta en espejo>",
                     "hipotesis": "no aplica" if sin_hip else "<H1 y H2, o 'no aplica'>",
                     "unidad": f"<{COLS_OBJ[tipo][0]}>", "medio": f"<{COLS_OBJ[tipo][1]}>",
                     "analisis": "<análisis según la pregunta (disenos.md): correlación, comparación, "
                                 "regresión lineal/logística/ordinal/conteo, multinivel, panel, SEM...>"}
    d = {"clave": "<clave de la Ficha>", "tipo": tipo}
    if tipo == "revision":
        d["diseno"] = {k: f"<{r}>" for k, r in CAMPOS_REVISION}
        d["diseno"]["bases"] = "Scopus; Web of Science"
        d["diseno"]["ecuacion"] = {"Scopus": "<TITLE-ABS-KEY(...) literal>",
                                   "Web of Science": "<TS=(...) literal>"}
    elif tipo == "secundarios":
        d["diseno"] = {k: f"<{r}>" for k, r in CAMPOS_SECUNDARIOS}
    else:
        d["diseno"] = {c: f"<{c}>" for c in CAMPOS_DISENO[tipo]}
    d["objetivos"] = [obj("General"), obj("OE1"), obj("OE2")]
    d["nota_objetivos"] = ("<por qué no hay hipótesis, en una línea>" if sin_hip else "")
    filas = [[f"<{c}>" for c in COLS_OPER[tipo]]]
    d["variables_extraccion" if tipo == "revision" else "operacionalizacion"] = filas
    cp = COLS_PAPERS["revision" if tipo == "revision" else "default"]
    d["papers"] = [{k: f"<{h}>" for k, h in cp}]
    d["fecha_busqueda"] = "<fecha de la búsqueda web del Paso 3>"
    d["pendientes"] = [{"pendiente": "<qué falta>", "quien": "<Investigador / coautor>",
                        "como": "<cómo se resuelve>"}]
    return d


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--json")
    ap.add_argument("--out")
    ap.add_argument("--ejemplo", choices=TIPOS, help="imprime un JSON vacío con los campos de ese tipo")
    a = ap.parse_args()
    if a.ejemplo:
        print(json.dumps(ejemplo(a.ejemplo), ensure_ascii=False, indent=2))
        return
    if not a.json:
        ap.error("falta --json")
    with open(a.json, encoding="utf-8") as f:
        d = json.load(f)
    wb, avisos, npen = construir(d)
    out = _salida(a.out or f"MATRIZ_{d['clave']}.xlsx")
    wb.save(out)
    print(f"ok: {out}")
    print(f"  hojas: {', '.join(wb.sheetnames)}")
    print(f"  pendientes: {npen}")
    for w in avisos:
        print(f"  ! {w}")


if __name__ == "__main__":
    main()
