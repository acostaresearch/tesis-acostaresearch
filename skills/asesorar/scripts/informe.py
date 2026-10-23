r"""Pasa el informe de revisión de Markdown sencillo a Word.

    python informe.py <informe.md> <04-revision\Informe_de_revision.docx>

Entiende «# título», «## subtítulo», «- viñeta», «1. numerada», tablas con | y párrafos. **negrita**.
"""
import re
import sys

from docx import Document
from docx.shared import Pt

import comun  # noqa: F401

ent, sal = sys.argv[1:3]
d = Document()
d.styles["Normal"].font.name = "Arial"
d.styles["Normal"].font.size = Pt(11)


def poner(p, texto):
    for tr in re.split(r"(\*\*[^*]+\*\*)", texto):
        if tr.startswith("**") and tr.endswith("**"):
            p.add_run(tr[2:-2]).bold = True
        elif tr:
            p.add_run(tr)


lineas = open(ent, encoding="utf-8").read().splitlines()
k = 0
while k < len(lineas):
    l = lineas[k].rstrip()
    if not l.strip():
        k += 1
    elif l.startswith("|"):
        filas = []
        while k < len(lineas) and lineas[k].startswith("|"):
            c = [x.strip() for x in lineas[k].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", x) for x in c):
                filas.append(c)
            k += 1
        t = d.add_table(rows=len(filas), cols=max(len(f) for f in filas))
        t.style = "Table Grid"
        for i, f in enumerate(filas):
            for j, x in enumerate(f):
                poner(t.cell(i, j).paragraphs[0], f"**{x}**" if i == 0 and x else x)
        d.add_paragraph()
    else:
        m = re.match(r"(#{1,3})\s+(.*)", l)
        if m:
            d.add_heading(m.group(2), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
        elif re.match(r"\s*[-•]\s+", l):
            poner(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-•]\s+", "", l))
        elif re.match(r"\s*\d+[.)]\s+", l):
            poner(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", l))
        else:
            poner(d.add_paragraph(), l)
        k += 1
d.save(sal)
print(f"Guardado: {sal}")
