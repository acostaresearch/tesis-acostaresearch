import sys, re, copy, docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
src, dst = sys.argv[1], sys.argv[2]
d = docx.Document(src)
body = d.element.body
W = lambda t: qn('w:'+t)

def has_sect(el):
    return el is not None and el.tag == W('p') and el.find(W('pPr')+'/'+W('sectPr')) is not None

def blank(line=480):
    p = OxmlElement('w:p'); ppr = OxmlElement('w:pPr'); sp = OxmlElement('w:spacing')
    sp.set(W('after'), '0'); sp.set(W('line'), str(line)); sp.set(W('lineRule'), 'auto')
    ppr.append(sp); p.append(ppr); return p

def set_pbb(p_el, val):
    ppr = p_el.get_or_add_pPr()
    for e in ppr.findall(W('pageBreakBefore')): ppr.remove(e)
    e = OxmlElement('w:pageBreakBefore'); e.set(W('val'), '1' if val else '0')
    # pageBreakBefore va después de pStyle/keepNext/keepLines
    anchor = None
    for tag in ('pStyle','keepNext','keepLines'):
        x = ppr.find(W(tag))
        if x is not None: anchor = x
    if anchor is not None: anchor.addnext(e)
    else: ppr.insert(0, e)

def bajar_a_linea(p, linea):
    """deja el título en la línea `linea` de una hoja nueva, con líneas en blanco delante"""
    el = p._p
    prev = el.getprevious()
    set_pbb(el, False)
    blanks = [blank() for _ in range(linea - 1)]
    for b in blanks: el.addprevious(b)
    if not has_sect(prev):
        set_pbb(blanks[0], True)

P = d.paragraphs
idx = {p.text.strip(): i for i, p in enumerate(P) if p.style.name == 'Heading 1'}

# 1. el salto de sección romano/arábigo pasa del inicio de la introducción al inicio del Capítulo 1
p_sec = P[idx['INTRODUCCIÓN'] - 1]._p
assert has_sect(p_sec) and p_sec.find(W('r')) is None
cap1 = [p for p in P if p.text.startswith('CAPÍTULO I\n')][0]._p
nuevo = blank(); nuevo.find(W('pPr')).append(p_sec.find(W('pPr')).find(W('sectPr')))
cap1.addprevious(nuevo)
p_sec.getparent().remove(p_sec)

# 2. márgenes: izq 3, der 2,5, sup 2,5, inf 3
for sp in body.iter(W('sectPr')):
    m = sp.find(W('pgMar'))
    m.set(W('left'), '1701'); m.set(W('right'), '1417'); m.set(W('top'), '1417'); m.set(W('bottom'), '1701')

# 3. número de página abajo a la derecha
for s in d.sections:
    if s.header.is_linked_to_previous: continue
    hp = s.header._element.findall(W('p'))
    tiene_page = any('PAGE' in (x.text or '') for p in hp for x in p.iter(W('instrText')))
    if not tiene_page: continue
    campo = hp[0]
    fp = s.footer._element.findall(W('p'))[0]
    for r in list(campo.findall(W('r'))):
        campo.remove(r); fp.append(r)
    ppr = fp.get_or_add_pPr()
    jc = OxmlElement('w:jc'); jc.set(W('val'), 'right'); ppr.append(jc)
    jh = campo.find(W('pPr')).find(W('jc'))
    if jh is not None: campo.find(W('pPr')).remove(jh)

# 4. estilos de título: subrayado y negrita; subtítulos al margen
for sid in ('Ttulo1','Ttulo2','Ttulo3','Ttulo4'):
    st = d.styles.element.find(f"{W('style')}[@{W('styleId')}='{sid}']")
    rpr = st.find(W('rPr'))
    if rpr.find(W('u')) is None:
        u = OxmlElement('w:u'); u.set(W('val'), 'single')
        col = rpr.find(W('color'))
        (col.addprevious(u) if col is not None else rpr.append(u))
    ppr = st.find(W('pPr'))
    for e in ppr.findall(W('ind')): ppr.remove(e)

# 5. capítulos: «CAPÍTULO 1» en la línea 6 y el título en la línea 8
ROM = {'I':'1','II':'2','III':'3','IV':'4','V':'5'}
for p in d.paragraphs:
    if p.style.name == 'Heading 1' and p.text.startswith('CAPÍTULO '):
        for t in p._p.iter(W('t')):
            m = re.fullmatch(r'CAPÍTULO ([IVX]+)', t.text or '')
            if m: t.text = 'CAPÍTULO ' + ROM[m.group(1)]
        br = next(p._p.iter(W('br')))
        br.addprevious(OxmlElement('w:br'))
        bajar_a_linea(p, 6)

# 6. RESUMEN, INTRODUCCIÓN, CONCLUSIONES, BIBLIOGRAFÍA (y sus pares) en la línea 8
LINEA8 = ['RESUMEN','ABSTRACT','INTRODUCCIÓN','CONCLUSIONES','RECOMENDACIONES','GLOSARIO','BIBLIOGRAFÍA']
for p in d.paragraphs:
    if p.style.name == 'Heading 1' and p.text.strip() in LINEA8:
        bajar_a_linea(p, 8)

# 7. referencias a los capítulos en el texto, en arábigos
for p in d.paragraphs:
    for t in p._p.iter(W('t')):
        if t.text and re.search(r'Capítulo [IVX]+\b', t.text):
            t.text = re.sub(r'Capítulo (V|IV|III|II|I)\b', lambda m: 'Capítulo ' + ROM[m.group(1)], t.text)

# 8. bibliografía: cuatro espacios simples tras el título, entradas a espacio simple al margen
bib = [p for p in d.paragraphs if p.style.name == 'Heading 1' and p.text.strip() == 'BIBLIOGRAFÍA'][0]
el = bib._p
for _ in range(4):
    b = blank(240); el.addnext(b)
n = 0
x = el.getnext()
while x is not None:
    if x.tag == W('p') and x.find(W('pPr')) is not None and x.find(W('pPr')).find(W('ind')) is not None:
        ppr = x.find(W('pPr'))
        for e in ppr.findall(W('spacing')): ppr.remove(e)
        sp = OxmlElement('w:spacing'); sp.set(W('after'), '240'); sp.set(W('line'), '240'); sp.set(W('lineRule'), 'auto')
        ppr.find(W('ind')).addprevious(sp); n += 1
    x = x.getnext()
print('referencias a espacio simple:', n)

d.save(dst)
print('ok', dst)
