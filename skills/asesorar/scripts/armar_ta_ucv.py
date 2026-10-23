"""Arma un trabajo académico de segunda especialidad UCV (guía PP-G-02.07) desde cero.

    python armar_ta_ucv.py texto.txt tablas.json logo.png salida.docx --titulo "…" --especialidad "…"

Carátula del Anexo 1 (autor y asesor en amarillo), índices en hoja ii-iii, cuerpo desde 1, anexos
apaisados sin número. Después: indice_word.py --despues "Índice de contenidos" --sin-titulo,
indice_tablas_word.py --despues "Índice de tablas" e indice.ps1. tablas.json: {id: {n, titulo,
encabezado, filas, nota, letra, anchos}}. Ejemplo de uso: Wilmer, 21-oct-2026.

texto.txt: «# Título» = nivel 1 · «[TABLA id]» = tabla de tablas.json · «[REFERENCIAS]» = las líneas
siguientes hasta «[FIN]» son referencias (sangría francesa) · «[ANEXOS]» abre la sección de anexos
(horizontal, sin número) · «[SALTO]» = salto de página · el resto, párrafos. *texto* va en cursiva.
"""
import json
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

a = sys.argv[1:]
texto, tablas_json, logo, salida = a[:4]
op = lambda n, d: a[a.index(n) + 1] if n in a else d
TAB = json.load(open(tablas_json, encoding="utf-8-sig"))
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

TITULO = op("--titulo", "Título del trabajo académico")
ESPECIALIDAD = op("--especialidad", "ESPECIALIDAD EN …")

d = Document()
# ---------- estilos ----------
st = d.styles["Normal"]
st.font.name = "Arial"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.element.rPr.rFonts.set(qn("w:cs"), "Arial")
pf = st.paragraph_format
pf.line_spacing = 1.5
pf.space_after = Pt(0)
pf.space_before = Pt(0)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for n in ("Heading 1", "Heading 2"):
    h = d.styles[n]
    h.font.name = "Arial"
    h.font.size = Pt(12)
    h.font.bold = True
    h.font.italic = False
    h.font.color.rgb = RGBColor(0, 0, 0)
    rpr = h.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = parse_xml(f"<w:rFonts {W}/>")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Arial")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    h.paragraph_format.space_before = Pt(0)
    h.paragraph_format.space_after = Pt(12)
    h.paragraph_format.line_spacing = 1.5
    h.paragraph_format.keep_with_next = True
cap = d.styles["Caption"]
cap.font.name = "Arial"
cap.font.size = Pt(12)
cap.font.bold = True
cap.font.italic = False
cap.font.color.rgb = RGBColor(0, 0, 0)


def margenes(s, horizontal=False):
    if horizontal:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(s, m, Cm(2.54))


def runs(p, txt, bold=False, size=None, hl=False):
    """Escribe txt con *cursiva* y [[amarillo]]."""
    for trozo in re.split(r"(\*[^*]+\*|\[\[[^\]]+\]\])", txt):
        if not trozo:
            continue
        it = trozo.startswith("*") and trozo.endswith("*")
        am = trozo.startswith("[[")
        t = trozo[1:-1] if it else (trozo[2:-2] if am else trozo)
        r = p.add_run(t)
        r.bold = bold or None
        r.italic = it or None
        if size:
            r.font.size = Pt(size)
        if am or hl:
            r.font.highlight_color = 7  # amarillo
    return p


def par(txt="", align=None, bold=False, size=None, sangria=True, hl=False):
    p = d.add_paragraph()
    if align is not None:
        p.alignment = align
    if sangria and align in (None, WD_ALIGN_PARAGRAPH.JUSTIFY):
        p.paragraph_format.first_line_indent = Cm(1.27)
    runs(p, txt, bold, size, hl)
    return p


def titulo1(txt):
    p = d.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(txt)
    return p


def numeracion(s, fmt, inicio):
    s.footer.is_linked_to_previous = False
    s.header.is_linked_to_previous = False
    sp = s._sectPr
    for v in sp.findall(qn("w:pgNumType")):
        sp.remove(v)
    if fmt:
        sp.append(parse_xml(f'<w:pgNumType {W} w:fmt="{fmt}" w:start="{inicio}"/>'))
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for x in ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
                  '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>',
                  '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', '<w:r><w:t>1</w:t></w:r>',
                  '<w:r><w:fldChar w:fldCharType="end"/></w:r>'):
            fp._p.append(parse_xml(x.replace("<w:r>", f"<w:r {W}>", 1)))
    else:
        s.footer.paragraphs[0].text = ""


def borde(celda, lados):
    tcPr = celda._tc.get_or_add_tcPr()
    b = parse_xml(f"<w:tcBorders {W}>" + "".join(
        f'<w:{l} w:val="{v}" w:sz="8" w:space="0" w:color="000000"/>' for l, v in lados.items()) + "</w:tcBorders>")
    tcPr.append(b)


def tabla(cfg):
    n = cfg["n"]
    p = d.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Tabla {n}")
    r.bold = True
    p2 = d.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.keep_with_next = True
    r = p2.add_run(cfg["titulo"])
    r.italic = True
    filas = [cfg["encabezado"]] + cfg["filas"]
    t = d.add_table(rows=len(filas), cols=len(cfg["encabezado"]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    tblPr.append(parse_xml(f'<w:tblW {W} w:w="5000" w:type="pct"/>'))
    tblPr.append(parse_xml(f'<w:tblLook {W} w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="1" w:noVBand="1"/>'))
    tam = cfg.get("letra", 10)
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = ""
            q = c.paragraphs[0]
            q.paragraph_format.line_spacing = 1.0
            q.paragraph_format.first_line_indent = Cm(0)
            q.alignment = WD_ALIGN_PARAGRAPH.LEFT
            runs(q, str(v), bold=(i == 0), size=tam)
            lados = {}
            if i == 0:
                lados = {"top": "single", "bottom": "single"}
            if i == len(filas) - 1:
                lados["bottom"] = "single"
            if lados:
                borde(c, lados)
        if i == 0:
            trPr = t.rows[0]._tr.get_or_add_trPr()
            trPr.append(parse_xml(f"<w:tblHeader {W}/>"))
    if cfg.get("anchos"):
        # ancho fijo: Word respeta las columnas solo con tblLayout fixed, tblW en dxa y gridCol
        tw = parse_xml(f'<w:tblW {W} w:w="{int(sum(cfg["anchos"]) * 567)}" w:type="dxa"/>')
        tblPr.replace(tblPr.find(qn("w:tblW")), tw)
        tblPr.append(parse_xml(f'<w:tblLayout {W} w:type="fixed"/>'))
        for gc, a in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), cfg["anchos"]):
            gc.set(qn("w:w"), str(int(a * 567)))
        for j, a in enumerate(cfg["anchos"]):
            for fila in t.rows:
                fila.cells[j].width = Cm(a)
    if cfg.get("nota"):
        q = d.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        q.paragraph_format.line_spacing = 1.0
        q.paragraph_format.space_before = Pt(4)
        r = q.add_run("Nota. ")
        r.italic = True
        r.font.size = Pt(10)
        for trozo in re.split(r"(\*[^*]+\*)", cfg["nota"]):
            if trozo:
                rr = q.add_run(trozo.strip("*") if trozo.startswith("*") else trozo)
                rr.font.size = Pt(10)
                rr.italic = trozo.startswith("*") or None
    d.add_paragraph()


# ---------- carátula ----------
s0 = d.sections[0]
margenes(s0)
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(logo, width=Cm(11.7))
C = WD_ALIGN_PARAGRAPH.CENTER
par("", C)
par("PROGRAMA DE SEGUNDA ESPECIALIDAD", C, size=18)
par(ESPECIALIDAD, C, size=16)
par("", C)
par(TITULO, C, bold=True, size=14)
par("", C)
par("Autor(es):", C, bold=True)
par("[[Apellidos, Nombres (orcid.org/0000-0000-0000-0000)]]", C)
par("", C)
par("Asesor(a)(es):", C, bold=True)
par("[[Grado. Apellidos, Nombres (orcid.org/0000-0000-0000-0000)]]", C)
par("", C)
par("Línea de investigación:", C, bold=True)
par("Gestión y Calidad Educativa", C)
par("", C)
par("Línea de responsabilidad social universitaria:", C, bold=True)
par("Apoyo a la reducción de brechas y carencias en la educación en todos sus niveles", C)
par("", C)
par("TRUJILLO — PERÚ", C)
par("2026", C)

# ---------- índices ----------
s1 = d.add_section(WD_SECTION.NEW_PAGE)
margenes(s1)
titulo1("Índice de contenidos")
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
titulo1("Índice de tablas")

# ---------- cuerpo ----------
lineas = [l.rstrip("\n") for l in open(texto, encoding="utf-8-sig")]
i = 0
primero = True
en_refs = False
while i < len(lineas):
    l = lineas[i].strip()
    i += 1
    if not l:
        continue
    if l == "[FIN]":
        en_refs = False
        continue
    if en_refs:
        q = d.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        q.paragraph_format.left_indent = Cm(1.27)
        q.paragraph_format.first_line_indent = Cm(-1.27)
        runs(q, l)
        continue
    if l == "[REFERENCIAS]":
        en_refs = True
        continue
    if l == "[SALTO]":
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        continue
    if l == "[ANEXOS]":
        s3 = d.add_section(WD_SECTION.NEW_PAGE)
        margenes(s3, horizontal=True)
        continue
    if l.startswith("[TABLA "):
        tabla(TAB[l[7:-1]])
        continue
    if l.startswith("# "):
        if primero:
            s2 = d.add_section(WD_SECTION.NEW_PAGE)
            margenes(s2)
            primero = False
        titulo1(l[2:])
        continue
    if l.startswith("## "):
        q = d.add_paragraph(style="Heading 2")
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        q.add_run(l[3:])
        continue
    par(l)

for k, s in enumerate(d.sections):
    numeracion(s, *({1: ("lowerRoman", 2), 2: ("decimal", 1)}.get(k, (None, None))))
d.save(salida)
print("guardado", salida)
