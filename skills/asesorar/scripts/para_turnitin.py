r"""Saca la versión que se pasa por Turnitin: sin carátula, sin índices y sin anexos.

    python para_turnitin.py <corregida.docx> <03-turnitin\salida.docx> --desde "I. GENERALIDADES" --hasta "VI. ANEXOS" [--pdf]

--desde es cómo empieza el primer párrafo que se conserva; --hasta, cómo empieza el primero que ya no va
(se omite si no hay anexos). Las referencias se dejan: se descuentan con el filtro de bibliografía.
--pdf guarda además el PDF con Word.
"""
import subprocess
import sys

from docx import Document

import comun  # noqa: F401

a = sys.argv[1:]
ent, sal = a[0], a[1]
op = lambda n: a[a.index(n) + 1] if n in a else None
desde, hasta = op("--desde"), op("--hasta")

d = Document(ent)
cuerpo = d.element.body
hijos = [e for e in cuerpo if not e.tag.endswith("}sectPr")]
texto = lambda e: "".join(t.text or "" for t in e.iter() if t.tag.endswith("}t")).strip()


def donde(marca):
    hay = [k for k, e in enumerate(hijos) if e.tag.endswith("}p") and texto(e).lower().startswith(marca.lower())]
    if not hay:
        sys.exit(f"No hay ningún párrafo que empiece por «{marca}».")
    return hay


ini = donde(desde)[-1] if desde else 0  # el último: el primero suele ser la línea del índice
fin = [k for k in donde(hasta) if k > ini][-1] if hasta else len(hijos)
for k, e in enumerate(hijos):
    if k < ini or k >= fin:
        cuerpo.remove(e)
d.save(sal)
pal = sum(len(p.text.split()) for p in Document(sal).paragraphs)
print(f"Guardado: {sal} · {pal} palabras · quitados {ini} bloques del inicio y {len(hijos) - fin} del final")

if "--pdf" in a:
    pdf = sal[:-5] + ".pdf"
    ps = ("$w=New-Object -ComObject Word.Application; $w.Visible=$false; "
          f"$d=$w.Documents.Open('{sal}'); $d.SaveAs([ref]'{pdf}',[ref]17); $d.Close(); $w.Quit()")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    print(f"PDF: {pdf}" if r.returncode == 0 else f"No se pudo crear el PDF con Word: {r.stderr[:200]}")
