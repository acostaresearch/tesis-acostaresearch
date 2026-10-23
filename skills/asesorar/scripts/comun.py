"""Piezas compartidas por los scripts de la skill asesorar."""
import re
import sys
from pathlib import Path

from rutas import ASESORAR as RAIZ

REPO = RAIZ.parent
ETAPAS = ("01-original", "02-corregida", "03-turnitin", "04-revision", "05-pasadoporturnitin")

for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8")


def norm(s):
    return re.sub(r"\W+", " ", s.lower()).strip()


def carpeta(arg):
    """Acepta «CURASMA», «Josué Segundo Silva Ora» o una ruta completa."""
    p = Path(arg)
    if p.is_dir():
        return p
    todas = [d for d in RAIZ.iterdir() if d.is_dir()]
    hay = [d for d in todas if arg.lower() in d.name.lower()]
    if len(hay) == 1:
        return hay[0]
    sys.exit(f"No encuentro la carpeta «{arg}» en {RAIZ}. Hay: {', '.join(d.name for d in todas)}")


def con_estilos(d):
    """True si el documento usa estilos de título: entonces no hace falta adivinar por el aspecto."""
    return sum(bool(re.search(r"(?:heading|t[íi]?tulo)\s*\d", (p.style.name.lower() if p.style is not None else ""))) for p in d.paragraphs) >= 3


def es_titulo(p, solo_estilo=False):
    """Nivel del título (1-4) o 0. Mira el estilo; si no hay, una línea corta en mayúsculas o numerada."""
    e = p.style.name.lower() if p.style is not None else ""
    m = re.search(r"(?:heading|t[íi]?tulo)\s*(\d)", e)
    if m:
        return int(m.group(1))
    t = p.text.strip()
    if solo_estilo or not t or len(t.split()) > 14:
        return 0
    if re.match(r"[IVX]+\.\s+\S", t):
        return 1
    m = re.match(r"(\d+(?:\.\d+)*)\.?\s+[A-ZÁÉÍÓÚÑ]", t)
    if m and not t.endswith((".", ";", ",")):
        return min(4, m.group(1).count(".") + 1)
    if t.isupper() and len(t) > 3 and not t.endswith((".", ":", ";", ",")):
        return 1
    return 0


def bloques(docx):
    """Recorre el cuerpo en orden: ('p', índice en doc.paragraphs, párrafo) y ('t', nº de tabla, tabla)."""
    from docx import Document
    d = Document(docx)
    ip = {p._p: (i, p) for i, p in enumerate(d.paragraphs)}
    it = {t._tbl: (i, t) for i, t in enumerate(d.tables)}
    out = []
    for el in d.element.body:
        if el in ip:
            out.append(("p",) + ip[el])
        elif el in it:
            out.append(("t",) + it[el])
    return d, out


def texto_pdf(pdf):
    import pymupdf
    return "\n".join(pg.get_text() for pg in pymupdf.open(pdf))


def clave_env(nombre):
    """Lee una clave de la variable de entorno o del .env local del backend. No se imprime nunca."""
    import os
    if os.environ.get(nombre):
        return os.environ[nombre]
    f = REPO / "backend" / ".env"
    if f.exists():
        for linea in f.read_text(encoding="utf-8", errors="ignore").splitlines():
            if linea.startswith(nombre + "="):
                return linea.split("=", 1)[1].strip().strip("\"'") or None
    return None


def http_json(url, cabeceras=None, espera=25):
    import json
    import urllib.request
    req = urllib.request.Request(url, headers={"User-Agent": "asesorar/1.0", **(cabeceras or {})})
    try:
        with urllib.request.urlopen(req, timeout=espera) as r:
            return json.loads(r.read().decode("utf-8"))
    except Exception as e:  # sin red, 404, cupo agotado: quien llama decide
        return {"_error": str(e)}
