"""Deja en «Subir a Turnitin» la copia de un Word lista para subir: sin páginas
preliminares (carátula, jurado, declaratoria, dedicatoria, agradecimientos,
índices) ni anexos. Conserva del resumen a las referencias, sin tocar el texto.

    python para_turnitin.py <word de entrada> "<Subir a Turnitin/salida.docx>" [--ver]
"""
import copy
import io
import re
import sys
import unicodedata

from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
INICIO = re.compile(r'^(resumen|abstract|introduccion|capitulo i\b.*|i\.? introduccion|(i\.? )?generalidades)$')
FIN = re.compile(r'^anexos?$|^apendices?$')


def texto(e):
    return ''.join(x.text or '' for x in e.iter(W + 't')).strip()


def llano(s):
    s = unicodedata.normalize('NFD', s.lower())
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn').strip(' .:')


def es_titulo(e):
    st = e.find(W + 'pPr/' + W + 'pStyle')
    return st is not None and re.match(r'(Ttulo|Heading|Title)', st.get(W + 'val') or '')


def palabras(els):
    return sum(len(' '.join(x.text or '' for x in e.iter(W + 't')).split()) for e in els)


def main():
    args = [a for a in sys.argv[1:] if a != '--ver']
    ver = '--ver' in sys.argv
    src, dst = args[0], args[1]
    d = Document(src)
    body = d.element.body
    els = list(body)
    ps = [i for i, e in enumerate(els) if e.tag == W + 'p']
    ini = next((i for i in ps if es_titulo(els[i]) and INICIO.match(llano(texto(els[i])))), None)
    if ini is None:
        ini = next((i for i in ps if INICIO.match(llano(texto(els[i])))), None)
    fin = next((i for i in reversed(ps) if es_titulo(els[i]) and FIN.match(llano(texto(els[i])))), None)
    if fin is None:
        # título ANEXOS con estilo propio (no «Título»): vale el último párrafo que diga solo eso
        fin = next((i for i in reversed(ps) if FIN.match(llano(texto(els[i])))), None)
    if ini is None:
        sys.exit('No encuentro dónde empieza el cuerpo (RESUMEN, INTRODUCCIÓN, CAPÍTULO I o GENERALIDADES).')
    final = els[-1] if els[-1].tag == W + 'sectPr' else None
    tope = fin if fin is not None else (len(els) - 1 if final is not None else len(els))
    print(f'Empieza en: {texto(els[ini])}')
    print(f'Preliminares que se quitan: {palabras(els[:ini])} palabras')
    if fin is None:
        print('No hay título ANEXOS: no se quita nada al final.')
    else:
        print(f'Anexos que se quitan: {palabras(els[fin:len(els) - 1])} palabras')
    print(f'Queda: {palabras(els[ini:tope])} palabras')
    if ver:
        return
    quitar = els[:ini]
    if fin is not None:
        # los párrafos vacíos justo antes de ANEXOS (salto de página y de sección) también se van,
        # y la sección que cerraban pasa a ser la última del documento
        k = fin
        sect = None
        while k - 1 > ini and els[k - 1].tag == W + 'p' and not texto(els[k - 1]) \
                and els[k - 1].find('.//' + W + 'drawing') is None:
            k -= 1
            s = els[k].find(W + 'pPr/' + W + 'sectPr')
            if s is not None and sect is None:
                sect = copy.deepcopy(s)
        quitar += els[k:len(els) - 1] if final is not None else els[k:]
        if sect is not None and final is not None:
            body.replace(final, sect)
    for e in quitar:
        body.remove(e)
    a = {b.get(W + 'id'): b for b in body.iter(W + 'bookmarkStart')}
    z = {b.get(W + 'id'): b for b in body.iter(W + 'bookmarkEnd')}
    for k in set(a) ^ set(z):
        b = a.get(k) if k in a else z.get(k)
        b.getparent().remove(b)
    d.save(dst)
    r = Document(dst)
    t = [p.text.strip() for p in r.paragraphs if p.text.strip()]
    print(f'Primero: {t[0][:70]} | Último: {t[-1][:70]}')
    print(f'Tablas: {len(r.tables)} · Figuras: {len(r.inline_shapes)} · Guardado: {dst}')


if __name__ == '__main__':
    main()
