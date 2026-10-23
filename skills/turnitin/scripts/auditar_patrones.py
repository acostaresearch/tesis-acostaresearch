#!/usr/bin/env python3
"""Audita el texto reescrito contra el catalogo de patrones de IA y mide su ritmo.

Uso:
  python auditar_patrones.py texto.txt
  python auditar_patrones.py doc.docx --parrafos 38,39,45,47

Existe porque bajar similitud y bajar patrones de IA tiran en direcciones
opuestas: al reescribir para no parecerse a otras tesis es facil producir prosa
con la firma tipica de un modelo. Este script se corre DESPUES de reescribir y
ANTES de mostrarle nada al usuario.

Las metas de dispersion y de oraciones cortas son un piso, no una garantia.
Hay contenido, como una enumeracion de estudios con autor y año, que tiene un
techo de variacion que ninguna reescritura le quita. Cuando dos pasadas no
mueven el numero, decirlo en vez de seguir retorciendo el texto.
"""
import argparse
import re
import statistics
import sys

sys.stdout.reconfigure(encoding='utf-8')

ABREV = ["et al.", "Dr.", "Dra.", "Sr.", "Sra.", "S.A.", "S.A.C.", "E.I.R.L.", "p. ej.", "etc."]

PROMOCIONAL = r"\b(crucial|vital|esencial|fundamental|robusto|valios[oa]s?|hol[íi]stic[oa]s?|innovador\w*|imprescindible|primordial)\b"
CONECTORES = r"(^|\. )(asimismo|adem[áa]s|por otro lado|en este sentido|en tal sentido|en consecuencia|cabe destacar|cabe se[ñn]alar|en resumen|en general|finalmente|por consiguiente|de igual manera|en otras palabras)\b"
RELLENO = r"\b(en [úu]ltima instancia|en el fondo|no es casualidad|vale la pena (?:destacar|mencionar)|es importante (?:destacar|mencionar|se[ñn]alar))\b"


def oraciones(texto):
    t = texto
    for a in ABREV:
        t = t.replace(a, a.replace(".", "\x00"))
    return [s.replace("\x00", ".").strip()
            for s in re.split(r"(?<=[.?!])\s+", t.strip()) if s.strip()]


def contar(patron, texto, flags=re.I):
    return len(re.findall(patron, texto, flags))


def auditar(texto, etiqueta=""):
    oo = oraciones(texto)
    largos = [len(re.findall(r"\b\w+\b", o)) for o in oo]
    palabras = re.findall(r"\b[\wáéíóúñü]+\b", texto, re.I)
    if not largos or not palabras:
        sys.exit("Texto vacio.")

    media = statistics.mean(largos)
    desv = statistics.pstdev(largos)
    disp = desv / media if media else 0
    cortas = sum(1 for n in largos if n <= 8)
    func = {"el", "la", "los", "las", "un", "una", "unos", "unas", "de", "del", "a", "al",
            "en", "con", "por", "para", "sin", "sobre", "entre", "hasta", "desde", "y", "e",
            "o", "u", "pero", "sino", "que", "si", "no", "se", "su", "sus", "lo", "le", "les",
            "es", "son", "ser", "fue", "hay", "como", "cuando", "donde", "porque", "este",
            "esta", "estos", "estas", "ese", "esa", "esto", "eso", "ya", "mas", "más", "muy",
            "tambien", "también", "asi", "así", "solo", "sólo", "cada", "todo", "toda"}
    pct_func = 100 * sum(1 for p in palabras if p.lower() in func) / len(palabras)
    largo_pal = sum(len(p) for p in palabras) / len(palabras)

    if etiqueta:
        print(f"--- {etiqueta} ---")
    print(f"oraciones: {len(oo)} | palabras: {len(palabras)}")
    print(f"largo medio de oracion : {media:.1f}   (referencia humana ~23)")
    print(f"dispersion relativa    : {disp:.2f}   (piso >=0.45)")
    print(f"min-max                : {min(largos)}-{max(largos)}")
    print(f"oraciones cortas (<=8) : {cortas} de {len(oo)}   (piso >=15%)")
    print(f"palabras funcionales   : {pct_func:.1f}%   (piso >=40%)")
    print(f"largo medio de palabra : {largo_pal:.2f}   (techo <=5.3)")

    print("\nCatalogo de patrones:")
    checks = [
        ("rayas largas o cortas", texto.count("\u2014") + texto.count("\u2013"), 0),
        ("gerundios colgados (coma + gerundio)", contar(r", [a-záéíóúñ]+(ando|iendo)\b", texto), 0),
        ("dos puntos", texto.count(":"), 1),
        ("negacion paralela (no solo... sino)", contar(r"no s[óo]lo.{0,80}sino", texto), 0),
        ("lenguaje promocional", contar(PROMOCIONAL, texto), 0),
        ("conectores de apertura de oracion", contar(CONECTORES, texto, re.I | re.M), 0),
        ("frases de relleno", contar(RELLENO, texto), 0),
    ]
    for nombre, valor, techo in checks:
        marca = "ok" if valor <= techo else "REVISAR"
        print(f"  {nombre:38} {valor:>3}   {marca}")

    racha = mx = 0
    for n in largos:
        racha = racha + 1 if n >= 30 else 0
        mx = max(mx, racha)
    print(f"  {'racha de oraciones >=30 palabras':38} {mx:>3}   {'ok' if mx < 3 else 'REVISAR: paralelismo de plantilla'}")

    verbos = ["sostiene", "señala", "subraya", "apunta", "advierte", "explica", "destaca",
              "indica", "afirma", "coincide", "concluye", "muestra", "plantea", "establece"]
    rep = {v: contar(r"\b" + v, texto) for v in verbos if contar(r"\b" + v, texto) >= 3}
    if rep:
        print(f"  verbos de atribucion repetidos 3+ veces: {rep}   REVISAR")

    largas = [(n, o[:90]) for n, o in zip(largos, oo) if n >= 42]
    if largas:
        print("\nOraciones de 42 palabras o mas (candidatas a partir):")
        for n, o in largas:
            print(f"  {n:>3} | {o}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("archivo")
    ap.add_argument("--parrafos", help="indices separados por coma, solo para .docx")
    args = ap.parse_args()

    if args.archivo.endswith(".docx"):
        from docx import Document
        d = Document(args.archivo)
        if args.parrafos:
            idx = [int(x) for x in args.parrafos.split(",")]
            texto = "\n\n".join(d.paragraphs[i].text.strip() for i in idx)
        else:
            texto = "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
    else:
        texto = open(args.archivo, encoding="utf-8", errors="replace").read()

    auditar(texto, args.archivo)


if __name__ == "__main__":
    main()
