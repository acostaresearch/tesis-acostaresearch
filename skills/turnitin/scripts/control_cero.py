#!/usr/bin/env python3
"""Control antes de entregar: ¿la prosa se parece a la de los Word que dieron 0 % de IA?

Uso:
  python control_cero.py <Listo\\salida.docx>
  python control_cero.py <salida.docx> --original <actual\\word.docx> --marcado <marcado.json>
  python control_cero.py <salida.docx> --desde 27 --hasta 180      solo ese tramo de párrafos
  python control_cero.py <a.docx> <b.docx> ...                     varios (calibrar)

Existe por la Ruta 0 % (18-oct-2026). Con «*%» Turnitin no resalta nada y cada ronda
se hacía a ciegas: Roy, Pamela, Eduardo y 14-Parte I pasaron de 6 a 10 rondas en «*%».
Este control no reemplaza a Turnitin, pero frena la entrega cuando la prosa todavía
tiene lo que sí se repitió en los Word que se quedaron en «*%»:

  - párrafos que el primer reporte dejó LIMPIOS y que se reescribieron (frena);
  - el mismo arranque en muchos párrafos y la plantilla repetida por objetivo (aviso);
  - ritmo y texto conservado del autor (solo dato).

Calibración (18-oct-2026) con 0 %: 07, 08, 09, 10-Roy, 11, 12 (v2), contra «*%»: 12 (v1),
13-Pamela, 14-Parte I, 06-Actividad 4. NI el ritmo (media, dispersión, cortas) NI el % de
texto del autor separaron un grupo del otro: 11 y 12 dieron 0 % con peor ritmo que Pamela,
y los 0 % conservan solo 14-23 % del texto original. Por eso esos números son INFORMATIVOS.
Lo único que frena la entrega es lo que sí explicó los «*%» largos (Roy, Eduardo, Pamela):
párrafos que el primer reporte con marcas dejó limpios y que se reescribieron. Pamela v1
tenía 78; los 0 % con reporte de partida, 0 a 2. Si aparece otra señal que separe, se añade
aquí con su calibración y se anota en references/ruta-ia-cero.md.

No reescribe nada. Sale con código 0 si ENTREGAR y 1 si NO ENTREGAR.
"""
import argparse
import difflib
import json
import re
import statistics
import sys
from collections import Counter

import docx

sys.stdout.reconfigure(encoding="utf-8")

# Umbrales (ver calibración en references/ruta-ia-cero.md)
U = {
    "cortas_min": 6.0,      # % de oraciones de 8 palabras o menos
    "dispersion_min": 0.45,  # desviación / media del largo de oración
    "media_max": 32.0,      # largo medio de oración
    "arranque_max": 4,      # un mismo arranque (2 palabras) en más párrafos que esto
    "plantilla_max": 3,     # párrafos seguidos con el mismo arranque
    "original_min": 25.0,   # % de palabras del cuerpo que siguen siendo del autor (con --original)
}

NOMINAL = re.compile(r"\b\w+(ción|sión|miento|idad|encia|ancia)\b", re.I)
ART = {"el", "la", "los", "las", "un", "una", "unos", "unas"}


def cuerpo(doc, desde, hasta):
    for i, p in enumerate(doc.paragraphs):
        if i < desde or (hasta is not None and i > hasta):
            continue
        t = " ".join(p.text.split())
        if len(t.split()) < 25:
            continue
        estilo = (p.style.name or "").lower()
        if any(e in estilo for e in ("heading", "título", "titulo", "caption", "toc", "bibliog")):
            continue
        yield i, t


def oraciones(t):
    t = t.replace("et al.", "et al\x00").replace("p. ej.", "p\x00 ej\x00")
    ss = re.split(r"(?<=[.?!])\s+(?=[A-ZÁÉÍÓÚÑ¿(])", t)
    return [s.replace("\x00", ".").strip() for s in ss if s.strip()]


def palabras(t):
    return re.findall(r"[\wáéíóúüñÁÉÍÓÚÜÑ]+", t)


def arranque(t):
    ws = [w.lower() for w in palabras(t)[:2]]
    return " ".join(ws)


def medir(ruta, desde=0, hasta=None):
    doc = docx.Document(ruta)
    pars = list(cuerpo(doc, desde, hasta))
    largos, n_pal, nomin, art = [], 0, 0, 0
    for _, t in pars:
        for o in oraciones(t):
            n = len(palabras(o))
            if n:
                largos.append(n)
        ws = palabras(t)
        n_pal += len(ws)
        nomin += len(NOMINAL.findall(t))
        if ws and ws[0].lower() in ART:
            art += 1
    if not largos:
        return None
    media = statistics.mean(largos)
    arr = Counter(arranque(t) for _, t in pars)
    # racha: párrafos seguidos que abren igual (la plantilla por objetivo)
    racha, mayor, donde, prev = 1, 1, None, None
    for i, t in pars:
        a = arranque(t)
        racha = racha + 1 if a == prev else 1
        if racha > mayor:
            mayor, donde = racha, i
        prev = a
    return {
        "pars": pars,
        "parrafos": len(pars),
        "palabras": n_pal,
        "media": media,
        "dispersion": statistics.pstdev(largos) / media,
        "cortas": 100 * sum(1 for n in largos if n <= 8) / len(largos),
        "largas": 100 * sum(1 for n in largos if n >= 40) / len(largos),
        "nominal": 100 * nomin / max(n_pal, 1),
        "art": 100 * art / max(len(pars), 1),
        "arranques": [(a, c) for a, c in arr.most_common() if c > U["arranque_max"]],
        "plantilla": (mayor, donde),
    }


def original_conservado(pars, ruta_original):
    """% de palabras del cuerpo en párrafos casi idénticos (≥0,9) a uno del original."""
    orig = [" ".join(p.text.split()) for p in docx.Document(ruta_original).paragraphs]
    orig_set = set(orig)
    iguales, total, cambiados = 0, 0, []
    for i, t in pars:
        n = len(t.split())
        total += n
        if t in orig_set:
            iguales += n
            continue
        cand = orig[i] if i < len(orig) else ""
        r = difflib.SequenceMatcher(None, t, cand, autojunk=False).quick_ratio()
        if r >= 0.9 and difflib.SequenceMatcher(None, t, cand, autojunk=False).ratio() >= 0.9:
            iguales += n
        else:
            cambiados.append(i)
    return 100 * iguales / max(total, 1), cambiados


def limpios_tocados(pars, ruta_original, ruta_marcado):
    """Párrafos que el primer reporte dejó sin marca y que ya no son el original."""
    marcados = {r["parrafo"] for r in json.load(open(ruta_marcado, encoding="utf-8"))}
    orig = [" ".join(p.text.split()) for p in docx.Document(ruta_original).paragraphs]
    orig_set = set(orig)
    fuera = []
    for i, t in pars:
        if t in orig_set:
            continue
        # se busca su párrafo en el original por parecido (el índice puede correrse)
        mejor, k = 0.0, None
        for j in range(max(0, i - 40), min(len(orig), i + 40)):
            if len(orig[j].split()) < 25:
                continue
            r = difflib.SequenceMatcher(None, t, orig[j], autojunk=False).quick_ratio()
            if r > mejor:
                mejor, k = r, j
        if k is not None and k not in marcados and mejor >= 0.5:
            fuera.append(i)
    return fuera


def informe(ruta, a):
    m = medir(ruta, a.desde, a.hasta)
    print(f"=== {ruta}")
    if not m:
        print("sin prosa de 25 palabras o más")
        return True
    fallas, avisos = [], []
    print(f"prosa: {m['parrafos']} párrafos, {m['palabras']} palabras")
    print(f"oración media {m['media']:.1f} · dispersión {m['dispersion']:.2f} · "
          f"cortas (≤8) {m['cortas']:.1f} % · largas (≥40) {m['largas']:.1f} % · "
          f"-ción/-miento por 100 palabras {m['nominal']:.1f} · abren con artículo {m['art']:.0f} %")
    if m["cortas"] < U["cortas_min"]:
        avisos.append(f"pocas oraciones cortas ({m['cortas']:.1f} % < {U['cortas_min']} %): ritmo parejo")
    if m["dispersion"] < U["dispersion_min"]:
        avisos.append(f"dispersión {m['dispersion']:.2f} < {U['dispersion_min']}: oraciones de largo parecido")
    if m["media"] > U["media_max"]:
        avisos.append(f"oración media {m['media']:.1f} > {U['media_max']}")
    for a_, c in m["arranques"]:
        avisos.append(f"«{a_}…» abre {c} párrafos (máximo {U['arranque_max']})")
    mayor, donde = m["plantilla"]
    if mayor > U["plantilla_max"]:
        avisos.append(f"plantilla: {mayor} párrafos seguidos abren igual (termina en ¶{donde})")
    if a.original:
        pct, _ = original_conservado(m["pars"], a.original)
        print(f"texto del autor conservado: {pct:.0f} %")
        if pct < U["original_min"]:
            avisos.append(f"solo {pct:.0f} % del cuerpo sigue siendo del autor (dato: los 0 % tienen 14-23 %)")
        if a.marcado:
            fuera = limpios_tocados(m["pars"], a.original, a.marcado)
            if len(fuera) > 2:
                fallas.append(f"{len(fuera)} párrafos limpios en el primer reporte fueron reescritos: "
                              f"{', '.join('¶' + str(x) for x in fuera[:30])}"
                              f"{' …' if len(fuera) > 30 else ''} → devolver al original (cirugía mínima si hay similitud)")
    for f in avisos:
        print(f"  aviso: {f}")
    if fallas:
        print("NO ENTREGAR:")
        for f in fallas:
            print(f"  - {f}")
        return False
    if not a.marcado:
        print("SIN FRENO: falta --original y --marcado del primer reporte con marcas, que es lo único que este control puede comprobar")
    else:
        print("ENTREGAR: como mucho 2 párrafos limpios del primer reporte fueron reescritos")
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docs", nargs="+")
    ap.add_argument("--original", help="Word de actual (o 01-original / 03-turnitin de partida)")
    ap.add_argument("--marcado", help="marcado.json del PRIMER reporte de IA con marcas")
    ap.add_argument("--desde", type=int, default=0)
    ap.add_argument("--hasta", type=int)
    a = ap.parse_args()
    ok = True
    for d in a.docs:
        ok = informe(d, a) and ok
        print()
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
