"""Devuelve la negrita al rótulo cuando una reescritura de párrafo entero la extendió a todo
el párrafo («Primera (Objetivo específico 1): …», «Estadística descriptiva: …»).

    python restaurar_rotulos.py <word de antes> <word reescrito>   (corrige el segundo en su sitio)
"""
import copy
import io
import sys

from docx import Document
from docx.text.run import Run

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def negrita(p):
    total = sum(len(r.text) for r in p.runs) or 1
    return sum(len(r.text) for r in p.runs if r.bold) / total


def main():
    antes, despues = Document(sys.argv[1]), Document(sys.argv[2])
    hechos = 0
    for i, (x, y) in enumerate(zip(antes.paragraphs, despues.paragraphs)):
        if negrita(y) - negrita(x) <= 0.15:
            continue
        rotulo = ''
        for r in x.runs:
            if not r.bold:
                break
            rotulo += r.text
        rotulo = rotulo.rstrip()
        texto = y.text
        if not rotulo or not texto.startswith(rotulo):
            print(f'¶{i} SIN ARREGLAR: el rótulo «{rotulo[:40]}» ya no abre el párrafo')
            continue
        r0 = y.runs[0]
        for r in y.runs[1:]:
            r._r.getparent().remove(r._r)
        r1 = copy.deepcopy(r0._r)
        r0._r.addnext(r1)
        r0.text, r0.bold = rotulo, True
        resto = Run(r1, y)
        resto.text, resto.bold = texto[len(rotulo):], None
        rpr = r1.find(W + 'rPr')
        if rpr is not None:
            for t in ('b', 'bCs'):
                for e in rpr.findall(W + t):
                    rpr.remove(e)
        hechos += 1
        print(f'¶{i} rótulo en negrita restaurado: {rotulo[:50]}')
    despues.save(sys.argv[2])
    print(f'Arreglados: {hechos}')


if __name__ == '__main__':
    main()
