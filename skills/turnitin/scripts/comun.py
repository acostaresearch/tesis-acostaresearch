"""Piezas compartidas por los scripts de la skill turnitin."""
import re
import sys
from pathlib import Path

import pymupdf
from docx import Document

from rutas import ASESORAR as RAIZ_ASESORAR, TURNITIN as RAIZ

for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8")


def norm(s):
    return re.sub(r"\W+", " ", s.lower()).strip()


def carpeta(arg):
    """Acepta «02-articulo», «02», una ruta completa o «asesorar\\Luna» para las de tesis-asesorar."""
    p = Path(arg)
    if p.is_dir():
        return p
    m = re.match(r"(?:tesis-)?asesorar[\\/](.+)", arg, re.I)
    if m:
        if (RAIZ_ASESORAR / m.group(1)).is_dir():
            return RAIZ_ASESORAR / m.group(1)
        hay = [d.name for d in RAIZ_ASESORAR.iterdir() if d.is_dir()]
        sys.exit(f"No encuentro la carpeta «{m.group(1)}» en {RAIZ_ASESORAR}. Hay: {', '.join(hay)}")
    if (RAIZ / arg).is_dir():
        return RAIZ / arg
    todas = [d for d in RAIZ.iterdir() if d.is_dir()]
    hay = [d for d in todas if d.name.lower().startswith(arg.lower())]
    if len(hay) == 1:
        return hay[0]
    sys.exit(f"No encuentro la carpeta «{arg}» en {RAIZ}. Hay: {', '.join(d.name for d in todas)}")


SALIDA = re.compile(r" - (humanizado|similitud|para turnitin)( \d+)?$", re.I)


def es_asesoria(base):
    """Carpeta con la organización de asesoría (01-original … 05-pasadoporturnitin), que se conserva.

    Lo son todas las de tesis-asesorar, y las de tesis-turnitin que llegaron así y no tienen «actual».
    """
    base = Path(base)
    if RAIZ_ASESORAR.resolve() in base.resolve().parents:
        return True
    return sub(base, "03-turnitin") is not None and sub(base, "actual") is None


def sitios(base):
    """(nombres, carpeta del Word analizado, carpeta de salida, carpeta de reportes) según la raíz.

    En tesis-asesorar el Word y la salida comparten 03-turnitin y los reportes van juntos en
    05-pasadoporturnitin; no se crea ni se mueve nada.
    """
    if es_asesoria(base):
        t = sub(base, "03-turnitin")
        return ("03-turnitin", "03-turnitin", "05-pasadoporturnitin"), t, t, sub(base, "05-pasadoporturnitin")
    return ("actual", "Listo", "reportes"), sub(base, "actual"), sub(base, "listo"), sub(base, "reportes", "reporte")


SUBIR = "Subir a Turnitin"


def subir(base):
    """Carpeta con la copia recortada (sin carátula, índices ni anexos) que se sube a Turnitin.

    Listo guarda el Word completo, que es el que se envía al tesista. Las de asesoría no la tienen.
    """
    return None if es_asesoria(base) else sub(base, SUBIR.lower())


def sub(base, *nombres):
    """Subcarpeta sin importar mayúsculas ni singular/plural (reporte, reportes, Listo, listo)."""
    for d in base.iterdir():
        if d.is_dir() and d.name.lower() in nombres:
            return d
    return None


def docx_de(d):
    if d is None:
        return []
    return sorted((f for f in d.glob("*.docx") if not f.name.startswith("~")), key=lambda f: f.stat().st_mtime)


def tipo_de_reporte(texto):
    """'ia', 'similitud' o None según lo que dice el PDF: no se fía del nombre ni de la carpeta."""
    t = texto.lower()
    if "detectado como ia" in t or "detected as ai" in t:
        return "ia"
    if any(x in t for x in ("similitud general", "overall similarity", "fuentes principales",
                            "indice de similitud", "índice de similitud", "informe de originalidad")):
        return "similitud"
    return None


def dato(texto, etiqueta):
    # El reporte nuevo pone el valor en la línea siguiente; el clásico, tras dos puntos.
    m = re.search(re.escape(etiqueta) + r"\s*(?::[ \t]*|\n\s*)(.+)", texto)
    return m.group(1).strip() if m else "?"


def parrafos(docx):
    return [p.text for p in Document(docx).paragraphs]


def parecido(pdf_texto_norm, docx):
    """Parte de las ventanas de 6 palabras del Word que aparecen en el reporte."""
    w = norm(" ".join(parrafos(docx))).split()
    v = [" ".join(w[i:i + 6]) for i in range(0, max(1, len(w) - 5), 7)]
    return sum(x in pdf_texto_norm for x in v) / len(v) if v else 0.0


def oraciones(texto):
    """Parte en oraciones sin cortar dentro de un paréntesis de cita."""
    out, ini, prof = [], 0, 0
    for m in re.finditer(r"[().;]", texto):
        c = m.group()
        if c == "(":
            prof += 1
        elif c == ")":
            prof = max(0, prof - 1)
        elif len(texto[ini:m.end()].split()) < 4:
            continue  # iniciales y abreviaturas («P. M.», «et al.») no cierran oración
        elif prof == 0 and re.match(r"\s+[A-ZÁÉÍÓÚÑ¿«\"]", texto[m.end():m.end() + 3]):
            out.append(texto[ini:m.end()])
            ini = m.end()
    out.append(texto[ini:])
    return [o for o in out if o.strip()]


def mismo_nombre(nombre_en_reporte, docx):
    """El reporte clásico trae las hojas como foto: solo queda comparar el nombre del archivo analizado."""
    limpio = lambda x: norm(re.sub(r"\.docx.*$", "", x).replace("_", " "))
    a, b = limpio(nombre_en_reporte), limpio(docx.name)
    # El clásico recorta por delante los nombres largos («me_final_-_Version…»): vale que sea el final.
    return a == b or (len(a) >= 25 and b.endswith(a))


def cruzar_tramos(tramos, docx, salida):
    """Cruza los tramos resaltados de un reporte de similitud con las oraciones del Word."""
    import json
    G = set()
    total_pdf = 0
    for t in tramos:
        w = norm(t).split()
        total_pdf += len(w)
        G.update(" ".join(w[k:k + 3]) for k in range(len(w) - 2))
    if not G:
        sys.exit("No encontré resaltado con texto. Si es el reporte clásico, hay que mirar las hojas como imagen.")

    res = []
    for i, p in enumerate(Document(docx).paragraphs):
        t = p.text
        if len(t.split()) < 6:
            continue
        ors, tot = [], 0
        for o in oraciones(t):
            crudas = o.split()
            w = [norm(x) for x in crudas]
            limpio = [x for x in w if x]
            idx = [k for k, x in enumerate(w) if x]
            cub = [False] * len(limpio)
            for k in range(len(limpio) - 2):
                if " ".join(limpio[k:k + 3]) in G:
                    cub[k] = cub[k + 1] = cub[k + 2] = True
            n = sum(cub)
            tot += n
            vis, dentro = [], False
            marca = {idx[k]: c for k, c in enumerate(cub)}
            for k, x in enumerate(crudas):
                m = marca.get(k, dentro)
                if m and not dentro:
                    x = "[" + x
                if not m and dentro:
                    vis[-1] += "]"
                dentro = m
                vis.append(x)
            if dentro:
                vis[-1] += "]"
            ors.append({"texto": o, "visto": " ".join(vis), "coinciden": n, "palabras": len(limpio)})
        if tot >= 5:
            res.append({"parrafo": i, "estilo": p.style.name, "coinciden": tot,
                        "palabras": len(t.split()), "oraciones": ors})

    json.dump(res, open(salida, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    suma = sum(r["coinciden"] for r in res)
    print(f"Palabras resaltadas en el reporte: {total_pdf} · encontradas en el Word: {suma} · "
          f"párrafos con coincidencia: {len(res)} · detalle en {salida}\n")
    for r in res:
        print(f"¶{r['parrafo']} · {r['estilo']} · {r['coinciden']} de {r['palabras']} palabras")
        for o in r["oraciones"]:
            if o["coinciden"]:
                print(f"   ({o['coinciden']}/{o['palabras']}) {o['visto'].strip()}")
        print()
